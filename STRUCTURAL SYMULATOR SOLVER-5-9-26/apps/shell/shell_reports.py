"""shell_reports.py -- the Shell tab's written report and its Excel round
trip. 2026-09-14. No Tkinter.

The report is written for a reader checking the calculation by hand: every
number is labelled with what it is, where it came from, and -- where the
app departs from a textbook value -- why. The membrane-theory section puts
the finite-element answer beside the classical formula with their ratio,
the same habit the Perforated Beam tab uses to reconcile with hand reports.
"""
import datetime
import json

import numpy as np

from apps.shell import shell_model as sm
from apps.shell import shell_design as sd
from apps.shell import shell_wind as wind

KN = 1e3


def _f(v, d=2):
    return f'{v:.{d}f}'


def membrane_comparison(res):
    """Rows comparing the FE interior shear with membrane theory for a hypar
    (z = p + q x + r y + k x y): Nxy = q_plan / (2k), per load case. Empty
    if the surface is not a hypar."""
    k = res.hypar_fit()
    if k is None:
        return None
    fem = res.fem
    c = fem['mesh']['centroids']
    x0, x1, y0, y1 = fem['mesh']['plan']
    xm, ym = (x0 + x1) / 2, (y0 + y1) / 2
    inner = (np.abs(c[:, 0] - xm) < 0.3 * (x1 - x0)) & (np.abs(c[:, 1] - ym) < 0.3 * (y1 - y0))
    inner &= ~fem.get('in_head', np.zeros(len(c), bool))
    plan_area = (x1 - x0) * (y1 - y0)
    rows = []
    for j, case in enumerate(res.cases):
        # load on the SHELL only -- the edge beams' own weight never passes
        # through the shell, so it must not be averaged into q
        Fz = -float(fem['Fe_cases'][j][:, 2::6].sum())
        if abs(Fz) < 1e-6:
            continue
        q = Fz / plan_area                                  # N/m^2 of plan, average
        theory = q / (2 * k)
        fe_mean = float(res.per_case[j]['Nxy'][inner].mean()) if inner.any() else float('nan')
        rows.append({'case': case, 'q_plan': q, 'theory': theory, 'fe': fe_mean,
                     'ratio': fe_mean / theory if theory else float('nan')})
    return {'k': k, 'rows': rows, 'n_inner': int(inner.sum())}


def report_lines(model, res, des, thicken_log=None):
    d = model.data
    code = model.code
    L = []
    add = L.append
    add('SHELL CALCULATION REPORT')
    add(f'Generated {datetime.datetime.now():%Y-%m-%d %H:%M} by the Shell tab '
        '(STRUCTURAL SYMULATOR SOLVER). Units: SI (kN, m, MPa, cm for sections, '
        'cm2/m for steel).')
    add('')
    add('1. GEOMETRY')
    add('  Definitions:')
    for ln in model.ws.lines:
        add(f'    {ln}')
    x0, x1, y0, y1 = res.fem['mesh']['plan']
    m = res.fem['mesh']
    add(f'  Surface z = {d["surface"]}(x, y); thickness t = {d["thickness"]}(x, y)')
    add(f'  Plan x {x0:g} .. {x1:g} m, y {y0:g} .. {y1:g} m; mesh {m["nx"]} x {m["ny"]} '
        f'four-node MITC4 shell elements (element diagonal {m["h_el"]:.2f} m)')
    t = res.fem['shells'].a
    add(f'  Thickness in the model: {t.min()*100:.1f} .. {t.max()*100:.1f} cm')
    for zn in d.get('zones', []):
        add(f'    thickening zone: where {zn["rule"]}  ->  t >= {zn["t"]*100:.1f} cm')
    layer = d.get('auto_layer')
    if layer:
        add(f'    automatic thickening layer: {len(layer["points"])} elements raised')
    add('')
    add('2. MATERIAL, SUPPORTS, EDGE BEAMS, COLUMNS')
    mat = d['material']
    add(f"  Concrete f'c = {mat['fc']:g} MPa, Ec = {res.fem['Ec_mpa']:.0f} MPa "
        f"({code.clauses.get('Ec', '')}), nu = {mat['nu']:g}, unit weight "
        f"{mat['gamma']:g} kN/m3; steel fy = {mat['fy']:g} MPa")
    for s in d.get('supports', []):
        blk = s.get('block', sm.DEFAULT_BLOCK) if s.get('at') in sm.POINT_SUPPORTS else 0
        add(f'  support: {sm.SUPPORT_AT.get(s["at"], s["at"])} [{s.get("which", "")}] '
            f'{sm.SUPPORT_LABELS.get(s["type"], s["type"])}'
            + (f', solid block {blk:g} m' if blk else ''))
    for b in d.get('beams', []):
        add(f'  edge/rib beam on {b["line"]}: {b["b"]:g} x {b["h"]:g} cm, {b.get("offset", "below")}')
    for c in d.get('columns', []):
        add(f'  column at ({c["x"]}, {c["y"]}): {c["b"]:g} x {c["h"]:g} cm, height '
            f'{c["height"]:g} m, base {c.get("base", "fixed")}, head '
            f'{c.get("capital") or max(c["b"], c["h"]) / 100:g} m')
    add('')
    add('3. LOADS')
    for c in d['cases']:
        add(f'  case {c["name"]}: {c["kind"]}')
    if d.get('self_weight'):
        add(f'  self-weight of shell, beams and columns in case {d["self_weight_case"]}')
    for ld in d.get('loads', []):
        if ld['type'] == 'point':
            add(f'  [{ld["case"]}] point load at ({ld.get("x")}, {ld.get("y")}): '
                f'Px {ld.get("Px", 0)} Py {ld.get("Py", 0)} P(down) {ld.get("Pz", 0)} kN')
        else:
            add(f'  [{ld["case"]}] {sm.LOAD_TYPES[ld["type"]]}: {ld["value"]} kN/m2'
                + (f'   ({ld["note"]})' if ld.get('note') else ''))
    w = d['wind']
    if any(k == 'W' for k in res.kinds.values()):
        for ln in wind.describe(w['V'], w['z'], w['exposure'], w['category'],
                                w.get('Kzt', 1.0), w.get('Kd', wind.KD_BUILDING),
                                w.get('G', wind.G_RIGID)):
            add('  ' + ln)
        add(f'  wind speed basis: CIRSOC 102-{w.get("basis", "2005")}; W factors in the '
            f'combinations multiplied by {res.wind_scale:g}')
        add('  ' + wind.CP_NOTE)
    add('')
    add(f'4. LOAD COMBINATIONS ({code.name}; {code.clauses.get("combinations", "")})')
    for c in res.combos:
        add(f'  {c.name}')
    add(f'  {res.service.name}  (deflection)')
    add('')
    add('5. EQUILIBRIUM (applied loads + reactions must be zero)')
    for j, case in enumerate(res.cases):
        eq = res.equilibrium[j]
        add(f'  {case}: applied (x, y, z) = ({_f(eq["applied"][0]/KN)}, '
            f'{_f(eq["applied"][1]/KN)}, {_f(eq["applied"][2]/KN)}) kN; '
            f'reactions = ({_f(eq["reaction"][0]/KN)}, {_f(eq["reaction"][1]/KN)}, '
            f'{_f(eq["reaction"][2]/KN)}) kN; residual '
            f'{np.abs(eq["force"]).max():.1e} N')
    add('')
    mc = membrane_comparison(res)
    add('6. MEMBRANE THEORY BESIDE THE FINITE-ELEMENT ANSWER')
    if mc is None:
        add('  The surface is not a hyperbolic paraboloid, so there is no closed-form '
            'membrane answer to compare with.')
    else:
        add(f'  The surface is a hypar z = p + q x + r y + k x y with k = {mc["k"]:.5f} 1/m.')
        add('  Membrane theory (rigid edge members, load per plan area): Nxy = q / (2k), '
            'Nx = Ny = 0. FE value = mean over the central 60% x 60% of the plan '
            f'({mc["n_inner"]} elements).')
        for r in mc['rows']:
            add(f'   {r["case"]:>4}: q = {r["q_plan"]/KN:.3f} kN/m2 (plan average)   '
                f'theory {r["theory"]/KN:8.2f} kN/m   FE {r["fe"]/KN:8.2f} kN/m   '
                f'ratio {r["ratio"]:.3f}')
        add('  Why they differ: the formula assumes edge members that do not deform; '
            'real edge beams stretch and bend, so the shell near them bends and the '
            'shear redistributes (tests: stiffer beams bring the ratio to 1.00). '
            'Self-weight is per SURFACE area, which a plan-area formula only '
            'approximates on a steep hypar.')
    add('')
    add('7. SHELL DESIGN')
    for ln in des.summary():
        add('  ' + ln)
    for p in getattr(des, 'punching', []):
        what = 'column head' if p['kind'] == 'column' else 'support block'
        add(f'  {what} at ({p["xy"][0]:.2f}, {p["xy"][1]:.2f}): shell shear into it '
            f'{p["Vu"]/KN:.1f} kN; critical perimeter {p["b0"]:.2f} m ({p["sides"]} sides '
            f'inside the shell) at d = {p["d"]*1000:.0f} mm; vc = {p["vc"]:.2f} MPa; '
            f'utilisation {p["util"]:.2f}; thickness needed around it '
            f'{p["t_req"]*100:.0f} cm'
            + ('  -- MESH TOO COARSE for this block: refine' if p['coarse'] else ''))
    S = des.S
    for key, label in (('x_top', 'x, top'), ('y_top', 'y, top'), ('x_bot', 'x, bottom'),
                       ('y_bot', 'y, bottom'), ('x_mid', 'x, central mesh'),
                       ('y_mid', 'y, central mesh')):
        a = des.steel[key]
        if np.all(np.isnan(a)):
            continue
        i = int(np.nanargmax(a))
        db, s = sd.suggest_bars(a[i], des.t[i], S)
        bars = f'  -> {db:.0f} mm @ {s:.0f} mm' if db else '  -> does not fit'
        add(f'  steel {label:16s} max {a[i]*1e4:6.2f} cm2/m '
            f'(min {np.nanmin(a)*1e4:5.2f}){bars}')
    if des.beams:
        add('  Edge beams and columns (PRELIMINARY: envelope forces, tension steel for '
            'N + M, simplified N-M check, stirrups for shear):')
        worst = {}
        for b in des.beams:
            key = b['tag'].split(':')[0]
            cur = worst.get(key)
            if cur is None or (b['N_tension'] - b['N_compression'] + b['M_vertical']) > \
                    (cur['N_tension'] - cur['N_compression'] + cur['M_vertical']):
                worst[key] = b
        for key, b in worst.items():
            add(f'   {key}: {b["b"]*100:.0f}x{b["h"]*100:.0f} cm  N+ {b["N_tension"]/KN:.0f} '
                f'kN  N- {b["N_compression"]/KN:.0f} kN  M {b["M_vertical"]/KN:.1f} kNm  '
                f'V {b["V"]/KN:.1f} kN  -> As {b["As_long"]*1e4:.1f} cm2, '
                f'stirrups {b["Av_s"]*1e4:.2f} cm2/m, N-M {b["u_compression"]:.2f}')
    if thicken_log:
        add('  Automatic thickening:')
        for ln in thicken_log:
            add('    ' + ln)
    add('')
    add('8. CODE CLAUSES USED')
    for k, v in code.clauses.items():
        add(f'  {k:12s} {v}')
    add('  Not yet read off the code text (verify before sign-off):')
    for u in code.UNVERIFIED:
        add(f'   - {u}')
    add('')
    add('9. ASSUMPTIONS TO REVIEW')
    add(f'  - Buckling: classical local formula x {S.kb:g} (user factor; the code gives none).')
    add('  - Point supports and columns are modelled as rigid solid blocks of the '
        'stated size; element values right beside them depend on the mesh, the '
        'block shear check does not.')
    add('  - Linear elastic analysis, uncracked stiffness; long-term deflection by the '
        f'multiplier {d["design"].get("longterm", 3.0):g}.')
    return L


# ═══════════════════════════════════════════════════════════════════════════
#  Excel
# ═══════════════════════════════════════════════════════════════════════════
MODEL_SHEET = 'Model (import)'


def export_excel(path, model, res=None, des=None, thicken_log=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MODEL_SHEET
    bold = Font(bold=True)
    ws.append(['key', 'value (JSON) -- this sheet is what Import Excel reads'])
    for c in ws[1]:
        c.font = bold
    for k, v in model.to_dict().items():
        ws.append([k, json.dumps(v, ensure_ascii=False)])
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 120

    ws2 = wb.create_sheet('Definitions')
    ws2.append(['definition'])
    ws2['A1'].font = bold
    for ln in model.ws.lines:
        ws2.append([ln])
    ws2.column_dimensions['A'].width = 70

    if res is not None and des is not None:
        ws3 = wb.create_sheet('Elements')
        head = ['element', 'x (m)', 'y (m)', 'z (m)', 't (cm)', 't needed (cm)',
                'utilisation', 'governed by',
                'Nx env max (kN/m)', 'Nx env min', 'Ny max', 'Ny min', 'Nxy max', 'Nxy min',
                'Mx max (kNm/m)', 'Mx min', 'My max', 'My min', 'Mxy max', 'Mxy min',
                'Q max (kN/m)',
                'As x top (cm2/m)', 'As y top', 'As x bot', 'As y bot', 'As x mid', 'As y mid']
        ws3.append(head)
        for c in ws3[1]:
            c.font = bold
            c.fill = PatternFill('solid', fgColor='DDDDDD')
        env = envelope(res)
        cen = res.fem['mesh']['centroids']
        for e in range(len(des.t)):
            row = [e, cen[e, 0], cen[e, 1], cen[e, 2], des.t[e] * 100,
                   des.t_req[e] * 100 if np.isfinite(des.t_req[e]) else 'cannot',
                   float(des.util_max[e]), sd.CHECK_LABELS[des.governing[e]]]
            for f in ('Nx', 'Ny', 'Nxy'):
                row += [env[f][0][e] / KN, env[f][1][e] / KN]
            for f in ('Mx', 'My', 'Mxy'):
                row += [env[f][0][e] / KN, env[f][1][e] / KN]
            row.append(env['Q'][e] / KN)
            for k in ('x_top', 'y_top', 'x_bot', 'y_bot', 'x_mid', 'y_mid'):
                v = des.steel[k][e]
                row.append(None if np.isnan(v) else v * 1e4)
            ws3.append([round(v, 4) if isinstance(v, float) else v for v in row])
        ws4 = wb.create_sheet('Report')
        for ln in report_lines(model, res, des, thicken_log):
            ws4.append([ln])
        ws4.column_dimensions['A'].width = 140
    wb.save(path)
    return path


def import_excel_model(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if MODEL_SHEET not in wb.sheetnames:
        raise ValueError(f'"{path}" has no "{MODEL_SHEET}" sheet; it was not exported '
                         'by the Shell tab.')
    ws = wb[MODEL_SHEET]
    data = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not row or row[0] is None:
            continue
        data[str(row[0])] = json.loads(row[1]) if row[1] is not None else None
    wb.close()
    return sm.ShellModel.from_dict(data)


def envelope(res):
    """Per element: (max, min) over the combinations of each force field,
    and the largest transverse shear magnitude."""
    out = {}
    fs = [res.combo_forces(c) for c in res.combos] or [res.per_case[0]]
    for f in ('Nx', 'Ny', 'Nxy', 'Mx', 'My', 'Mxy'):
        stack = np.stack([F[f] for F in fs])
        out[f] = (stack.max(axis=0), stack.min(axis=0))
    out['Q'] = np.max(np.stack([np.hypot(F['Qx'], F['Qy']) for F in fs]), axis=0)
    return out
