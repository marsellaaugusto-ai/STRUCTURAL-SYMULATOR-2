"""Piece marks where the app meets them: the group list, the panel, the
tolerance setting, and the workbook.

The comparison itself is tested in tests/test_stereo_marks.py. This file is
about the wiring: that the mark shown beside a group is the one the engine
computed, that switching the display off really does stop the work, that
the tolerance someone typed is the one used and the one exported, and that
the schedule reaches the workbook without any of it being read back.
"""
import math
import os
import sys
import time

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import tkinter as tk

from apps.stereo import stereo_marks as sm
from apps.stereo.stereo_app_marks import META_TOL_KEY
from common import _ensure_openpyxl


@pytest.fixture(autouse=True)
def dialogs(monkeypatch):
    seen = []
    for kind in ('showinfo', 'showerror', 'showwarning'):
        monkeypatch.setattr('apps.stereo.stereo_app.messagebox.%s' % kind,
                            lambda *a, _k=kind, **kw: seen.append((_k,) + a))
    return seen


@pytest.fixture(scope='module')
def tk_root():
    last = None
    for attempt in range(6):
        try:
            root = tk.Tk()
            break
        except tk.TclError as exc:
            last = exc
            time.sleep(0.5 * (attempt + 1))
    else:
        pytest.skip('no Tk display after 6 attempts: %s' % last)
    root.geometry('1400x900')
    yield root
    try:
        root.destroy()
    except tk.TclError:
        pass


@pytest.fixture
def app(tk_root):
    from apps.stereo.stereo_app import StereoApp
    tab = tk.Frame(tk_root)
    a = StereoApp(tab)
    a._generate(push_undo=False)
    tab.pack(fill='both', expand=True)
    tk_root.update_idletasks()
    tk_root.update()
    yield a
    tab.destroy()


# ── a model with three trusses, the third a mirror of the other two ───────

CHIRAL = [(0, 0, 0), (4, 0, 0), (1.7, 3.1, 0), (1.3, 0.9, 2.8),
          (3.4, 2.2, 4.1)]
BARS = [(0, 1), (1, 2), (2, 0), (0, 3), (1, 3), (2, 3), (3, 4), (1, 4)]


def three_trusses(app, third_mirrored):
    nodes, members, groups = [], [], []
    for k, mirror in enumerate((False, False, third_mirrored)):
        at = len(nodes)
        for p in CHIRAL:
            nodes.append((p[0] + 40 * k, p[1], -p[2] if mirror else p[2]))
        first = len(members)
        for a, b in BARS:
            members.append({'a': a + at, 'b': b + at, 'profile': 'IPE 200',
                            'A': 28.5, 'I': 1940.0, 'conn': 'pin',
                            'E': 200.0, 'Fy': 235.0, 'Fu': 360.0, 'K': 1.0})
        groups.append({'id': k + 1, 'name': 'Truss %s' % 'ABC'[k],
                       'parent': None,
                       'members': set(range(first, len(members)))})
    app.nodes, app.members, app.groups = nodes, members, groups
    app.loads, app.supports = [], []
    app.results = None
    app.member_checks = None
    app._refresh_group_list()
    return groups


# ── the panel and the list ────────────────────────────────────────────────

def test_the_panel_exists_and_starts_switched_off(app):
    """Off by default: a comparison nobody asked to see should not be paid
    for on every refresh of a list."""
    assert hasattr(app, 'marks_on')
    assert app.marks_on.get() is False
    assert app._mark_of_group() == {}


def test_switched_off_the_list_says_nothing_about_marks(app):
    three_trusses(app, third_mirrored=False)
    app.marks_on.set(False)
    labels = [label for label, _gid in app._group_rows()]
    assert not any('T1' in s for s in labels)


def test_switched_on_identical_groups_show_one_mark_and_a_count(app):
    three_trusses(app, third_mirrored=False)
    app.marks_on.set(True)
    app._refresh_marks()
    labels = [label for label, _gid in app._group_rows()]
    assert sum('T1 ×3' in s for s in labels) == 3


def test_a_mirrored_truss_is_marked_apart_on_the_list(app):
    three_trusses(app, third_mirrored=True)
    app.marks_on.set(True)
    app._refresh_marks()
    rows = dict((gid, label) for label, gid in app._group_rows())
    assert 'T1 ×2' in rows[1] and 'T1 ×2' in rows[2]
    assert 'T1/m' in rows[3]
    assert '×2' not in rows[3].split('T1/m')[1]


def test_the_mark_on_the_list_is_the_one_the_engine_computed(app):
    three_trusses(app, third_mirrored=True)
    app.marks_on.set(True)
    by_gid, _rows = sm.assembly_marks(app.nodes, app.members, app.groups,
                                      app._mark_tol_mm())
    assert app._mark_of_group() == by_gid


def test_the_note_counts_the_parts_and_the_repeats(app):
    three_trusses(app, third_mirrored=True)
    app.marks_on.set(True)
    app._refresh_marks_note()
    text = app.marks_note.cget('text')
    assert 'part' in text
    assert '1 built more than once' in text, text


def test_switched_off_the_note_does_no_work_at_all(app):
    """_refresh_all runs this on every edit and every undo. Comparing a
    large model there would be a stutter on every keystroke, for an answer
    nobody asked for."""
    three_trusses(app, third_mirrored=True)
    app.marks_on.set(False)
    called = []
    app._marks_now = lambda *a, **k: called.append(1) or {}
    app._refresh_marks_note()
    assert called == []
    assert 'Tick' in app.marks_note.cget('text')


def test_a_model_with_no_groups_still_has_a_note(app):
    app.groups = []
    app.marks_on.set(True)
    app._refresh_marks_note()
    assert 'part' in app.marks_note.cget('text')


def test_an_empty_model_says_so_rather_than_dividing_by_zero(app):
    app._clear_model()
    app.marks_on.set(True)
    app._refresh_marks_note()
    assert app.marks_note.cget('text') == 'No rods yet.'
    assert app._mark_of_group() == {}


# ── the tolerance ─────────────────────────────────────────────────────────

def test_the_tolerance_starts_at_the_default(app):
    assert app._mark_tol_mm() == sm.DEFAULT_TOL_MM


def test_a_typed_tolerance_is_used(app):
    app.mark_tol.set('25')
    assert app._mark_tol_mm() == 25.0


def test_nonsense_in_the_box_falls_back_and_says_what_it_understood(app):
    """A half-typed number must not stop the panel comparing anything."""
    app.mark_tol.set('abc')
    assert app._mark_tol_mm() == sm.DEFAULT_TOL_MM
    app._on_mark_tol_changed()
    assert app.mark_tol.get() == '%g' % sm.DEFAULT_TOL_MM


def test_a_coarser_tolerance_can_merge_two_nearly_equal_trusses(app):
    """The setting earns its place: 20 mm out is a different part at 1 mm
    and the same part at 100 mm."""
    three_trusses(app, third_mirrored=False)
    moved = sorted(app.groups[2]['members'])
    node = app.members[moved[0]]['b']
    x, y, z = app.nodes[node]
    app.nodes[node] = (x + 0.02, y, z)          # 20 mm
    app.marks_on.set(True)

    app.mark_tol.set('1')
    fine = app._mark_of_group()
    assert fine[1] == fine[2] != fine[3]

    app.mark_tol.set('100')
    coarse = app._mark_of_group()
    assert coarse[1] == coarse[2] == coarse[3]


# ── the schedule window ───────────────────────────────────────────────────

def test_the_schedule_window_opens_and_lists_both_levels(app):
    three_trusses(app, third_mirrored=True)
    win = app._marks_dialog()
    app.root.update_idletasks()
    text = win.children[[k for k in win.children if 'text' in k][0]].get(
        '1.0', 'end')
    assert 'ASSEMBLIES' in text and 'PARTS' in text
    assert 'T1' in text and 'T1/m' in text
    assert 'TOTAL' in text
    win.destroy()


def test_the_schedule_text_is_read_only(app):
    three_trusses(app, third_mirrored=False)
    win = app._marks_dialog()
    txt = win.children[[k for k in win.children if 'text' in k][0]]
    assert txt.cget('state') == 'disabled'
    win.destroy()


def test_the_roles_panel_follows_the_model(app):
    """It was built before the tab had a model and never refreshed with
    one, so it said "No rods yet" over a full structure."""
    three_trusses(app, third_mirrored=False)
    for m in app.members:
        m['role'] = 'diagonal'
    app._refresh_all()
    labels = [w.cget('text')
              for line in app._roles_rows.winfo_children()
              for w in line.winfo_children() if hasattr(w, 'cget')
              and 'text' in w.keys()]
    assert any('Diagonal' in s for s in labels), labels


def test_the_report_totals_match_the_model(app):
    three_trusses(app, third_mirrored=True)
    data = app._marks_now()
    lines = app._marks_report(data)
    total = [ln for ln in lines if ln.startswith('TOTAL')][0]
    assert str(len(app.members)) in total
    assert sum(r['qty'] for r in data['parts']) == len(app.members)


def test_a_model_with_no_assemblies_says_so_instead_of_an_empty_table(app):
    app.groups = []
    lines = app._marks_report()
    assert any('nothing' in ln for ln in lines)
    assert any(ln.startswith('B1') for ln in lines)


# ── the workbook ──────────────────────────────────────────────────────────

@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_the_tolerance_travels_in_the_workbook(app, tmp_path, monkeypatch):
    from apps.stereo import stereo_reports as sr
    three_trusses(app, third_mirrored=True)
    app.mark_tol.set('7')
    path = str(tmp_path / 'm.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    assert os.path.exists(path)
    assert float(sr.read_excel_meta(path)[META_TOL_KEY]) == 7.0

    app.mark_tol.set('1')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.askopenfilename',
                        lambda **kw: path)
    app._import_excel()
    assert app._mark_tol_mm() == 7.0, 'the setting the file was written with'


@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_the_workbook_carries_the_schedule(app, tmp_path, monkeypatch):
    import openpyxl
    from apps.stereo.stereo_marks_excel import SHEET
    three_trusses(app, third_mirrored=True)
    path = str(tmp_path / 'm.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    wb = openpyxl.load_workbook(path, data_only=True)
    assert SHEET in wb.sheetnames
    seen = [[c for c in row] for row in wb[SHEET].iter_rows(values_only=True)]
    flat = [str(c) for row in seen for c in row if c is not None]
    assert 'T1' in flat and 'T1/m' in flat
    assert 'B1' in flat
    assert any('ASSEMBLIES' in s for s in flat)
    assert any('PARTS' in s for s in flat)


@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_the_schedule_sheet_is_never_read_back(app, tmp_path, monkeypatch):
    """A mark is derived. Editing the sheet and re-importing must change
    nothing about the model -- otherwise a typed mark would become a fact."""
    import openpyxl
    from apps.stereo.stereo_marks_excel import SHEET
    three_trusses(app, third_mirrored=False)
    path = str(tmp_path / 'm.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()

    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'T1':
                cell.value = 'NONSENSE'
    wb.save(path)

    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.askopenfilename',
                        lambda **kw: path)
    app._import_excel()
    app.marks_on.set(True)
    assert set(app._mark_of_group().values()) == {'T1'}
    assert len(app.groups) == 3


@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_a_workbook_written_without_marks_has_no_such_sheet(tmp_path):
    """Every workbook exported before this existed, and the Export 3D path,
    which writes geometry for SketchUp rather than a schedule."""
    import openpyxl
    from apps.stereo import stereo_reports as sr
    from apps.stereo.stereo_marks_excel import SHEET
    nodes = [(0, 0, 0), (1, 0, 0)]
    members = [{'a': 0, 'b': 1, 'A': 20.0, 'I': 400.0, 'J': 400.0,
                'E': 200.0, 'Fy': 235.0, 'Fu': 360.0, 'K': 1.0,
                'r_gyr': 4.0, 'conn': 'pin'}]
    path = str(tmp_path / 'plain.xlsx')
    sr.export_excel(nodes, members, [], [], None, path)
    assert SHEET not in openpyxl.load_workbook(path).sheetnames


# ── issuing the numbers, so they survive the next revision ───────────────

@pytest.fixture
def issued(app, monkeypatch):
    """Issue the marks of whatever model is in front of us."""
    def go(answer=True):
        monkeypatch.setattr('apps.stereo.stereo_app_marks.messagebox'
                            '.askyesno', lambda *a, **kw: answer)
        return app._marks_issue()
    return go


def by_length(app):
    _by, rows = sm.part_marks(app.nodes, app.members, app._mark_tol_mm(),
                              app.mark_register)
    return {round(r['length_m'], 3): r['mark'] for r in rows}


def add_rods(app, length, count, profile='IPE 200'):
    for _ in range(count):
        at = len(app.nodes)
        app.nodes.extend([(0, len(app.nodes), 0),
                          (length, len(app.nodes), 0)])
        app.members.append({'a': at, 'b': at + 1, 'profile': profile,
                            'A': 28.5, 'I': 1940.0, 'conn': 'pin',
                            'E': 200.0, 'Fy': 235.0, 'Fu': 360.0, 'K': 1.0})


def only_rods(app, lengths):
    app.nodes, app.members, app.groups = [], [], []
    for length, count in lengths:
        add_rods(app, length, count)
    app.loads, app.supports, app.results, app.member_checks = [], [], None, None


def test_the_numbers_start_free(app):
    assert app.mark_register is None
    assert 'Issue' in app.mark_issue_btn.cget('text')


def test_issuing_holds_each_part_to_the_number_it_carries(app, issued):
    only_rods(app, [(3.0, 5), (4.0, 3)])
    before = by_length(app)
    issued()
    assert app.mark_register is not None
    add_rods(app, 6.0, 9)                  # would otherwise take B1
    after = by_length(app)
    assert after[3.0] == before[3.0] == 'B1'
    assert after[6.0] == 'B3'


def test_declining_the_question_issues_nothing(app, issued):
    only_rods(app, [(3.0, 2)])
    assert issued(answer=False) is None
    assert app.mark_register is None


def test_issuing_is_undoable(app, issued):
    only_rods(app, [(3.0, 2)])
    issued()
    assert app.mark_register is not None
    app._undo()
    assert app.mark_register is None


def test_the_button_says_which_way_it_goes(app, issued):
    only_rods(app, [(3.0, 2)])
    issued()
    assert 'Release' in app.mark_issue_btn.cget('text')
    app._undo()
    app._refresh_issue_button()
    assert 'Issue' in app.mark_issue_btn.cget('text')


def test_releasing_lets_the_numbers_move_again(app, issued, monkeypatch):
    only_rods(app, [(3.0, 5), (4.0, 3)])
    issued()
    add_rods(app, 6.0, 9)
    assert by_length(app)[3.0] == 'B1'
    monkeypatch.setattr('apps.stereo.stereo_app_marks.messagebox.askyesno',
                        lambda *a, **kw: True)
    app._marks_release()
    assert app.mark_register is None
    assert by_length(app)[6.0] == 'B1', 'free again, so the most used wins'


def test_an_empty_model_has_nothing_to_issue(app, issued, dialogs):
    app._clear_model()
    assert issued() is None
    assert app.mark_register is None


def test_the_note_says_the_numbers_are_held(app, issued):
    only_rods(app, [(3.0, 2)])
    app.marks_on.set(True)
    issued()
    app._refresh_marks_note()
    assert 'issued' in app.marks_note.cget('text').lower()


def test_the_note_names_what_was_withdrawn(app, issued):
    only_rods(app, [(3.0, 5), (4.0, 3)])
    app.marks_on.set(True)
    issued()
    app.members[:] = [m for m in app.members
                      if abs(sm.rod_length(app.nodes, m) - 4.0) > 1e-9]
    app._refresh_marks_note()
    text = app.marks_note.cget('text')
    assert 'no longer built' in text and 'B2' in text


def test_the_note_warns_when_the_tolerance_moved_under_the_register(app,
                                                                    issued):
    """Its signatures were computed at the issued tolerance, so at another
    one none of them match and everything would renumber silently."""
    only_rods(app, [(3.0, 2)])
    app.marks_on.set(True)
    app.mark_tol.set('1')
    issued()
    app.mark_tol.set('25')
    app._refresh_marks_note()
    assert 'none of the held numbers apply' in app.marks_note.cget('text')


# ── the register in the workbook ─────────────────────────────────────────

@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_the_register_travels_in_the_workbook(app, issued, tmp_path,
                                              monkeypatch):
    from apps.stereo.stereo_marks_excel import REGISTER_SHEET
    import openpyxl
    only_rods(app, [(3.0, 5), (4.0, 3)])
    issued()
    held = dict(app.mark_register[sm.PART_PREFIX])
    path = str(tmp_path / 'r.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    assert REGISTER_SHEET in openpyxl.load_workbook(path).sheetnames

    app.mark_register = None
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.askopenfilename',
                        lambda **kw: path)
    app._import_excel()
    assert app.mark_register is not None
    assert app.mark_register[sm.PART_PREFIX] == held


@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_the_numbers_hold_across_a_save_and_a_revision(app, issued, tmp_path,
                                                        monkeypatch):
    """The whole point: a model exported, edited and exported again must
    not renumber the parts that did not change."""
    only_rods(app, [(3.0, 5), (4.0, 3)])
    issued()
    before = by_length(app)
    path = str(tmp_path / 'rev1.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.askopenfilename',
                        lambda **kw: path)
    app._import_excel()
    add_rods(app, 6.0, 9)
    after = by_length(app)
    assert after[3.0] == before[3.0]
    assert after[4.0] == before[4.0]
    assert after[6.0] not in before.values()


@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_a_model_whose_numbers_are_free_writes_no_register(app, tmp_path,
                                                            monkeypatch):
    import openpyxl
    from apps.stereo.stereo_marks_excel import REGISTER_SHEET
    only_rods(app, [(3.0, 2)])
    path = str(tmp_path / 'free.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    assert REGISTER_SHEET not in openpyxl.load_workbook(path).sheetnames


@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_importing_a_file_with_no_register_frees_the_numbers(app, issued,
                                                              tmp_path,
                                                              monkeypatch):
    """The register belongs to the model, like its groups do."""
    only_rods(app, [(3.0, 2)])
    path = str(tmp_path / 'free.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    issued()
    assert app.mark_register is not None
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.askopenfilename',
                        lambda **kw: path)
    app._import_excel()
    assert app.mark_register is None


@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_the_register_sheet_says_what_each_number_stood_for(app, issued,
                                                             tmp_path,
                                                             monkeypatch):
    """A signature is 64 characters of hex; a reader deserves better."""
    import openpyxl
    from apps.stereo.stereo_marks_excel import REGISTER_SHEET
    only_rods(app, [(3.0, 5)])
    issued()
    path = str(tmp_path / 'r.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    ws = openpyxl.load_workbook(path, data_only=True)[REGISTER_SHEET]
    flat = [str(c) for row in ws.iter_rows(values_only=True) for c in row
            if c is not None]
    assert any('IPE 200' in s and '3.000' in s for s in flat), flat


@pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable')
def test_a_damaged_register_sheet_does_not_stop_the_import(app, issued,
                                                            tmp_path,
                                                            monkeypatch):
    """A model whose numbers are free is how every model used to be."""
    import openpyxl
    from apps.stereo.stereo_marks_excel import REGISTER_SHEET
    only_rods(app, [(3.0, 2)])
    issued()
    path = str(tmp_path / 'r.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    wb = openpyxl.load_workbook(path)
    ws = wb[REGISTER_SHEET]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'level':
                cell.value = 'nope'
    wb.save(path)
    rods = len(app.members)
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.askopenfilename',
                        lambda **kw: path)
    app._import_excel()
    assert len(app.members) == rods
    assert app.mark_register is None
