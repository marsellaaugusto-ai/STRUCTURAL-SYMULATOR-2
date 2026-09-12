"""
arch_app.py — Arch tab.

2D frame (beam-column) element solver following an arbitrary y=f(x) shape
(ArchModel), per-element N/V/M + extreme-fibre stress recovery (ArchResult),
and the Arch schematic + 2x2 diagram UI (ArchApp), plus this tab's Excel
report/import. Three-hinged arches get an independent rotational DOF at the
hinge node rather than a shared one -- global equilibrium of that free DOF
enforces zero moment there with no extra condensation math needed.
"""
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import math, os, sys, subprocess

from common import (
    _ensure_openpyxl,
    PANEL_W, INIT_CW, INIT_CH, INIT_DH,
    _beam_gauss_solve, _GAUSS5_NODES, _GAUSS5_WEIGHTS,
    _nice_ticks, _find_diagram_maxima, make_shape_fn,
)

def _adaptive_integral(fn, a, b, tol=1e-7, max_depth=20):
    """Recursive, error-controlled Gauss-Legendre quadrature for a scalar
    integrand. Compares a 5-point Gauss estimate over [a,b] against the sum
    of two independent 5-point estimates over its two halves; if they don't
    agree to `tol` (relative to the estimate's own scale), it recurses into
    each half. This concentrates panels exactly where the integrand actually
    varies quickly — including a steep-but-integrable edge singularity like
    1/sqrt(...) — rather than relying on a fixed, uniform panel count that
    may under-resolve such a spike (a flat 10-panel scheme was found to
    under-integrate a 1/sqrt-type edge singularity by ~16%). Every sample
    point stays strictly interior to whatever sub-panel it belongs to, at
    every recursion depth, so a function singular exactly at a or b is
    still never evaluated there.

    A function that is integrable but genuinely unbounded at an edge (e.g.
    1/sqrt(x) as x->0) never fully satisfies the error test as panels keep
    shrinking toward that edge — recursion is capped both by `max_depth` and
    by an absolute minimum panel width, so it terminates cleanly (accepting
    the small remaining truncation error right at the edge) instead of
    recursing until floating-point rounding lands a sample exactly on the
    singularity."""
    min_width = max(1e-12, (b - a) * 1e-10) if b > a else 1e-12

    def gauss5(lo, hi):
        h = (hi - lo) / 2
        mid = (lo + hi) / 2
        s = 0.0
        for node, wt in zip(_GAUSS5_NODES, _GAUSS5_WEIGHTS):
            try:
                s += wt * fn(mid + node*h)
            except (ZeroDivisionError, ValueError, OverflowError):
                pass   # skip an unevaluable sample; the other 4 nodes still contribute
        return s * h

    def rec(lo, hi, whole, depth):
        if depth >= max_depth or (hi - lo) <= min_width:
            return whole
        mid = (lo + hi) / 2
        left = gauss5(lo, mid)
        right = gauss5(mid, hi)
        if abs((left + right) - whole) <= tol * max(1.0, abs(whole)):
            return left + right
        return rec(lo, mid, left, depth+1) + rec(mid, hi, right, depth+1)

    if b <= a:
        return 0.0
    return rec(a, b, gauss5(a, b), 0)

def _integrate_load(fn, a, b, n_panels=None):
    """Integral of fn over [a,b] via adaptive Gauss-Legendre quadrature (see
    `_adaptive_integral`). Used to lump an arbitrary (possibly non-uniform)
    distributed load q(x) into an equivalent nodal force over a node's
    tributary sub-interval — accurate even when q is singular exactly at
    a or b, or varies sharply near either edge. `n_panels` is accepted for
    backward compatibility but no longer used (kept as a no-op parameter)."""
    return _adaptive_integral(fn, a, b)

class ArchModel:
    """
    Arch centerline y=f(x) discretized into straight prismatic 2D frame
    elements (axial + Euler bending, 3 DOF/node: u, v, theta), solved via
    the direct-stiffness method — the same family of method used for the
    truss and beam tabs.

    Supports at the two ends: 'pin' (restrains u,v; free rotation) or
    'fixed' (restrains u,v,theta). An optional internal hinge (for the
    three-hinged arch) is modelled by giving that node an *independent*
    rotational DOF for the element on each side, instead of one shared
    DOF — global equilibrium of that free, unloaded DOF then enforces
    zero moment on that element end automatically, with no special
    condensation math needed. Validated against the exact closed-form
    thrust H = wL^2/(8f) for a 3-hinged parabolic arch under a uniform
    load, and against classic cantilever / fixed-fixed beam formulas.

    Units: internal SI (m, N, N*m, Pa). The App layer converts to/from
    the kN, kN/m, GPa, cm units used elsewhere in this program.
    Sign convention: N > 0 = tension. M uses the sagging-positive
    (concave-up) convention, consistent along the arc.
    """
    def __init__(self, shape_fn, span, n_elem=60):
        self.shape_fn = shape_fn
        self.span = span
        self.n_elem = n_elem
        self.support_a = 'pin'    # at x = 0
        self.support_b = 'pin'    # at x = span
        self.hinge_x = None       # x-coordinate of internal hinge, or None
        self.E = 200e9            # Pa
        self.A = 1e-2             # m^2
        self.I = 1e-4             # m^4
        self.c_top = 0.1          # m, extrados fiber distance
        self.c_bot = 0.1          # m, intrados fiber distance
        self.w_weight = 0.0       # N/m, horizontal-projected, +down
        self.w_wind = 0.0         # N/m, horizontal-projected, +x
        self.point_loads = []     # [{'x':, 'fx':, 'fy':}] N
        self.distributed_loads = []  # [{'fn': q(x)->N/m, 'x1':, 'x2':, 'direction': 'vertical'|'horizontal'}]

    def nodes(self):
        xs = [self.span * i / self.n_elem for i in range(self.n_elem + 1)]
        ys = [self.shape_fn(x) for x in xs]
        return xs, ys

    def solve(self):
        xs, ys = self.nodes()
        n = self.n_elem
        nn = n + 1

        hinge_node = None
        if self.hinge_x is not None:
            best_i, best_d = None, None
            for i in range(1, n):
                d_ = abs(xs[i] - self.hinge_x)
                if best_d is None or d_ < best_d:
                    best_d, best_i = d_, i
            hinge_node = best_i

        dof = 2 * nn
        theta_id = {}
        for i in range(nn):
            if i == hinge_node:
                theta_id[i] = {'L': dof, 'R': dof + 1}
                dof += 2
            else:
                theta_id[i] = {'N': dof}
                dof += 1

        K = [[0.0] * dof for _ in range(dof)]
        F = [0.0] * dof
        elements = []

        for e in range(n):
            xa, ya = xs[e], ys[e]
            xb, yb = xs[e + 1], ys[e + 1]
            dx, dy = xb - xa, yb - ya
            L = math.hypot(dx, dy)
            if L < 1e-9:
                continue
            c, s = dx / L, dy / L
            EA_L = self.E * self.A / L
            EI = self.E * self.I
            kb = [
                [12*EI/L**3, 6*EI/L**2, -12*EI/L**3, 6*EI/L**2],
                [6*EI/L**2, 4*EI/L, -6*EI/L**2, 2*EI/L],
                [-12*EI/L**3, -6*EI/L**2, 12*EI/L**3, -6*EI/L**2],
                [6*EI/L**2, 2*EI/L, -6*EI/L**2, 4*EI/L],
            ]
            kloc = [[0.0]*6 for _ in range(6)]
            kloc[0][0] = EA_L; kloc[0][3] = -EA_L
            kloc[3][0] = -EA_L; kloc[3][3] = EA_L
            bidx = [1, 2, 4, 5]
            for i in range(4):
                for j in range(4):
                    kloc[bidx[i]][bidx[j]] += kb[i][j]

            T = [[c, s, 0, 0, 0, 0],
                 [-s, c, 0, 0, 0, 0],
                 [0, 0, 1, 0, 0, 0],
                 [0, 0, 0, c, s, 0],
                 [0, 0, 0, -s, c, 0],
                 [0, 0, 0, 0, 0, 1]]
            Tt = list(map(list, zip(*T)))
            tmp = [[sum(Tt[i][k]*kloc[k][j] for k in range(6)) for j in range(6)] for i in range(6)]
            kgl = [[sum(tmp[i][k]*T[k][j] for k in range(6)) for j in range(6)] for i in range(6)]

            a_theta = theta_id[e]['R'] if e == hinge_node else theta_id[e]['N']
            b_theta = theta_id[e+1]['L'] if (e+1) == hinge_node else theta_id[e+1]['N']
            gidx = [2*e, 2*e+1, a_theta, 2*(e+1), 2*(e+1)+1, b_theta]

            for i in range(6):
                for j in range(6):
                    K[gidx[i]][gidx[j]] += kgl[i][j]

            elements.append({'e': e, 'a': e, 'b': e+1, 'L': L, 'c': c, 's': s,
                              'T': T, 'kloc': kloc, 'gidx': gidx})

        trib = [0.0] * nn
        node_span = [(0.0, 0.0)] * nn   # (ta, tb) tributary domain per node
        for i in range(nn):
            left = xs[i] - xs[i-1] if i > 0 else 0.0
            right = xs[i+1] - xs[i] if i < n else 0.0
            trib[i] = (left + right) / 2
            ta = xs[i] - left/2
            tb = xs[i] + right/2
            node_span[i] = (ta, tb)

        for i in range(nn):
            F[2*i+1] += -self.w_weight * trib[i]
            F[2*i]   += self.w_wind * trib[i]

        for patch in self.distributed_loads:
            fn = patch['fn']
            px1, px2 = patch['x1'], patch['x2']
            if px2 < px1:
                px1, px2 = px2, px1
            direction = patch.get('direction', 'vertical')
            for i in range(nn):
                ta, tb = node_span[i]
                a_ = max(ta, px1); b_ = min(tb, px2)
                if b_ - a_ < 1e-12:
                    continue
                q_int = _integrate_load(fn, a_, b_)
                if direction == 'horizontal':
                    F[2*i] += q_int
                else:
                    F[2*i+1] += -q_int

        for p in self.point_loads:
            best_i, best_d = None, None
            for i in range(nn):
                d_ = abs(xs[i] - p['x'])
                if best_d is None or d_ < best_d:
                    best_d, best_i = d_, i
            F[2*best_i]   += p.get('fx', 0.0)
            F[2*best_i+1] += p.get('fy', 0.0)

        constrained = set()

        def apply_support(node, typ):
            if typ == 'free':
                return
            constrained.add(2*node); constrained.add(2*node+1)
            if typ == 'fixed':
                tid = theta_id[node]
                constrained.add(tid.get('N', tid.get('L')))

        apply_support(0, self.support_a)
        apply_support(n, self.support_b)

        free = [i for i in range(dof) if i not in constrained]
        if not free:
            raise ValueError("Arch is fully constrained; nothing to solve.")
        Kff = [[K[i][j] for j in free] for i in free]
        Ff = [F[i] for i in free]
        dfree = _beam_gauss_solve(Kff, Ff)
        if dfree is None:
            raise ValueError(
                "Singular stiffness matrix — arch is a mechanism "
                "(unstable / check supports and hinge placement).")
        d = [0.0] * dof
        for k, i in enumerate(free):
            d[i] = dfree[k]

        Kd = [sum(K[i][j]*d[j] for j in range(dof)) for i in range(dof)]
        R = [Kd[i] - F[i] for i in range(dof)]

        return ArchResult(self, xs, ys, elements, d, R, theta_id)

class ArchResult:
    def __init__(self, model, xs, ys, elements, d, R, theta_id):
        self.model = model
        self.xs, self.ys = xs, ys
        self.elements = elements
        self.d, self.R = d, R
        self.theta_id = theta_id

    def reaction(self, node):
        """Returns (Rx, Ry, M_or_None) at node 0 or node n."""
        tid = self.theta_id[node]
        key = tid.get('N', tid.get('L'))
        M = self.R[key]
        return self.R[2*node], self.R[2*node+1], M

    def internal_forces(self):
        """Per-element local end forces. N>0 tension; Ma, Mb sagging-positive
        bending moment at each element's node-a / node-b, N*m."""
        out = []
        for el in self.elements:
            gidx = el['gidx']; T = el['T']; kloc = el['kloc']
            dgl = [self.d[i] for i in gidx]
            dloc = [sum(T[i][j]*dgl[j] for j in range(6)) for i in range(6)]
            floc = [sum(kloc[i][j]*dloc[j] for j in range(6)) for i in range(6)]
            out.append({'N': floc[3], 'V': floc[1], 'Ma': -floc[2], 'Mb': floc[5],
                        'L': el['L'], 'a': el['a'], 'b': el['b']})
        return out

    def sample(self, n_per_elem=6):
        """Stations along the arc for plotting/stress-checking: arc length s,
        (x,y), a smoothed unit tangent (tc,ts) for the diagram's visual
        offset direction, the true per-element unit tangent (ec,es) that
        N/V/M at that station are actually defined in, N (const per
        element), V (const per element), M (linear interp Ma->Mb per
        element).

        Each node is emitted exactly once. (tc,ts) is averaged across the
        two elements meeting at that node (or the single adjacent element at
        the springs) purely so the diagram's perpendicular offset doesn't
        flip abruptly at every node. (ec,es) is NOT averaged — it's each
        station's owning element's own orientation, matching the exact
        local frame that N and V were computed in — use this whenever you
        need to resolve N/V into global Fx/Fy; using the smoothed (tc,ts)
        for that instead subtly mismatches the frame at every node and
        produces a spurious wavy-looking Fx/Fy curve even though N, V, M
        are each perfectly smooth."""
        forces = self.internal_forces()
        n_elem = len(self.elements)

        node_tc = [0.0] * (n_elem + 1)
        node_ts = [0.0] * (n_elem + 1)
        for i in range(n_elem + 1):
            if i == 0:
                c, s_ = self.elements[0]['c'], self.elements[0]['s']
            elif i == n_elem:
                c, s_ = self.elements[-1]['c'], self.elements[-1]['s']
            else:
                c1, s1 = self.elements[i-1]['c'], self.elements[i-1]['s']
                c2, s2 = self.elements[i]['c'], self.elements[i]['s']
                c, s_ = c1 + c2, s1 + s2
                norm = math.hypot(c, s_)
                if norm > 1e-12:
                    c, s_ = c/norm, s_/norm
                else:
                    c, s_ = c2, s2
            node_tc[i], node_ts[i] = c, s_

        s_out, x_out, y_out = [], [], []
        tc_out, ts_out, ec_out, es_out = [], [], [], []
        N_out, V_out, M_out = [], [], []
        s = 0.0
        for ei, (el, f) in enumerate(zip(self.elements, forces)):
            xa, ya = self.xs[el['a']], self.ys[el['a']]
            k_start = 0 if ei == 0 else 1   # skip k=0: identical to previous element's last point
            for k in range(k_start, n_per_elem + 1):
                t = k / n_per_elem
                x = xa + el['c'] * el['L'] * t
                y = ya + el['s'] * el['L'] * t
                M = f['Ma'] + (f['Mb'] - f['Ma']) * t
                if k == 0:
                    tc, ts = node_tc[el['a']], node_ts[el['a']]
                elif k == n_per_elem:
                    tc, ts = node_tc[el['b']], node_ts[el['b']]
                else:
                    tc, ts = el['c'], el['s']
                s_out.append(s + el['L']*t); x_out.append(x); y_out.append(y)
                tc_out.append(tc); ts_out.append(ts)
                ec_out.append(el['c']); es_out.append(el['s'])
                N_out.append(f['N']); V_out.append(f['V']); M_out.append(M)
            s += el['L']
        return {'s': s_out, 'x': x_out, 'y': y_out, 'tc': tc_out, 'ts': ts_out,
                'ec': ec_out, 'es': es_out, 'N': N_out, 'V': V_out, 'M': M_out}

    def fiber_stresses(self):
        """Extreme-fiber stresses (Pa) at every sampled station: returns
        (s, x, y, sigma_top, sigma_bot) with tension positive."""
        A, I = self.model.A, self.model.I
        ct, cb = self.model.c_top, self.model.c_bot
        samp = self.sample()
        sig_top, sig_bot = [], []
        for N, M in zip(samp['N'], samp['M']):
            axial = N / A
            sig_top.append(axial - M * ct / I)
            sig_bot.append(axial + M * cb / I)
        return samp['s'], samp['x'], samp['y'], sig_top, sig_bot

def export_arch_excel(state, path, result=None, model=None):
    """
    Writes a workbook with:
      • 'Model'   — every input needed to rebuild this exact arch (geometry,
                    supports, loads, section) — paired with import_arch_excel().
      • 'Results' — reactions, extremes, and a station-by-station table of
                    N, M, and extreme-fibre stresses (only if `result` given).
    `state` is a plain dict with keys: span, rise, n_elem, shape_expr,
    arch_type, hinge_frac, w_weight, w_wind, point_loads, distributed_loads,
    profile.
    """
    if not _ensure_openpyxl():
        raise RuntimeError('openpyxl is not available.')
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('Model')
    row = 1
    ws.cell(row=row, column=1, value='ARCH MODEL DATA — for Import from Excel (do not reorder columns)')
    ws.cell(row=row, column=1).font = Font(bold=True, size=11, color='1F4E79')
    row += 2

    ws.cell(row=row, column=1, value='[GEOMETRY]'); row += 1
    for col, lbl in enumerate(['span_m', 'rise_m', 'n_elem', 'shape_expr'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    ws.cell(row=row, column=1, value=state['span'])
    ws.cell(row=row, column=2, value=state['rise'])
    ws.cell(row=row, column=3, value=state['n_elem'])
    ws.cell(row=row, column=4, value=state['shape_expr'])
    row += 2

    ws.cell(row=row, column=1, value='[SUPPORTS]'); row += 1
    for col, lbl in enumerate(['arch_type', 'hinge_frac'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    ws.cell(row=row, column=1, value=state['arch_type'])
    ws.cell(row=row, column=2, value=state['hinge_frac'])
    row += 2

    ws.cell(row=row, column=1, value='[UNIFORM_LOADS]'); row += 1
    for col, lbl in enumerate(['w_weight_kNm', 'w_wind_kNm'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    ws.cell(row=row, column=1, value=state['w_weight'])
    ws.cell(row=row, column=2, value=state['w_wind'])
    row += 2

    ws.cell(row=row, column=1, value='[DISTRIBUTED_LOADS]'); row += 1
    for col, lbl in enumerate(['expr', 'x1_m', 'x2_m', 'direction'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for p in state['distributed_loads']:
        ws.cell(row=row, column=1, value=p['expr'])
        ws.cell(row=row, column=2, value=p['x1'])
        ws.cell(row=row, column=3, value=p['x2'])
        ws.cell(row=row, column=4, value=p['direction'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[POINT_LOADS]'); row += 1
    for col, lbl in enumerate(['x_m', 'fx_kN', 'fy_kN'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for p in state['point_loads']:
        ws.cell(row=row, column=1, value=p['x'])
        ws.cell(row=row, column=2, value=p['fx'])
        ws.cell(row=row, column=3, value=p['fy'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[SECTION]'); row += 1
    keys = ['E', 'A', 'I', 'c_top', 'c_bot', 'allow_tension', 'allow_compression']
    labels = ['E_GPa', 'A_cm2', 'I_cm4', 'c_top_cm', 'c_bot_cm',
              'allow_tension_kNcm2', 'allow_compression_kNcm2']
    for col, lbl in enumerate(labels, 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for col, key in enumerate(keys, 1):
        ws.cell(row=row, column=col, value=state['profile'][key])
    row += 1

    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 14

    if result is not None and model is not None:
        wr = wb.create_sheet('Results')
        Rx0, Ry0, M0 = result.reaction(0)
        Rxn, Ryn, Mn = result.reaction(model.n_elem)
        forces = result.internal_forces()
        s, x, y, sig_top, sig_bot = result.fiber_stresses()

        wr.cell(row=1, column=1, value='ARCH RESULTS').font = Font(bold=True, size=11, color='1F4E79')
        wr.cell(row=3, column=1, value='Reaction, left (x=0)')
        wr.cell(row=3, column=2, value='Rx_kN'); wr.cell(row=3, column=3, value=Rx0/1e3)
        wr.cell(row=3, column=4, value='Ry_kN'); wr.cell(row=3, column=5, value=Ry0/1e3)
        wr.cell(row=3, column=6, value='M_kNm'); wr.cell(row=3, column=7, value=M0/1e3)
        wr.cell(row=4, column=1, value=f'Reaction, right (x={model.span:.3f})')
        wr.cell(row=4, column=2, value='Rx_kN'); wr.cell(row=4, column=3, value=Rxn/1e3)
        wr.cell(row=4, column=4, value='Ry_kN'); wr.cell(row=4, column=5, value=Ryn/1e3)
        wr.cell(row=4, column=6, value='M_kNm'); wr.cell(row=4, column=7, value=Mn/1e3)

        Nmax_t = max((f['N'] for f in forces), default=0.0)
        Nmax_c = min((f['N'] for f in forces), default=0.0)
        Mmax = max([abs(f['Ma']) for f in forces] + [abs(f['Mb']) for f in forces], default=0.0)
        wr.cell(row=6, column=1, value='Max N tension (kN)'); wr.cell(row=6, column=2, value=Nmax_t/1e3)
        wr.cell(row=7, column=1, value='Max N compression (kN)'); wr.cell(row=7, column=2, value=Nmax_c/1e3)
        wr.cell(row=8, column=1, value='Max |M| (kN·m)'); wr.cell(row=8, column=2, value=Mmax/1e3)
        wr.cell(row=9, column=1, value='Max fibre tension (kN/cm²)'); wr.cell(row=9, column=2, value=max(sig_top+sig_bot)*1e-7)
        wr.cell(row=10, column=1, value='Max fibre compression (kN/cm²)'); wr.cell(row=10, column=2, value=min(sig_top+sig_bot)*1e-7)

        hdr_row = 12
        for col, lbl in enumerate(['s_m', 'x_m', 'y_m', 'N_kN', 'M_kNm', 'sigma_top_kNcm2', 'sigma_bot_kNcm2'], 1):
            wr.cell(row=hdr_row, column=col, value=lbl)
        samp = result.sample()
        for i in range(len(s)):
            r = hdr_row + 1 + i
            wr.cell(row=r, column=1, value=s[i])
            wr.cell(row=r, column=2, value=x[i])
            wr.cell(row=r, column=3, value=y[i])
            wr.cell(row=r, column=4, value=samp['N'][i]/1e3)
            wr.cell(row=r, column=5, value=samp['M'][i]/1e3)
            wr.cell(row=r, column=6, value=sig_top[i]*1e-7)
            wr.cell(row=r, column=7, value=sig_bot[i]*1e-7)
        for col in range(1, 8):
            wr.column_dimensions[get_column_letter(col)].width = 13

    wb.save(path)


def import_arch_excel(path):
    """
    Reads a 'Model' sheet written by export_arch_excel() and returns a plain
    state dict (same shape as the `state` argument passed into
    export_arch_excel), or raises ValueError with a human-readable message.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if 'Model' not in wb.sheetnames:
        raise ValueError('This workbook has no "Model" sheet — it was not '
                          'exported by the Arch tab.')
    ws = wb['Model']
    rows = list(ws.iter_rows(values_only=True))

    def find_section(name):
        for i, r in enumerate(rows):
            if r and r[0] == name:
                return i
        return -1

    def read_table(start_idx):
        hdr_i = start_idx + 1
        if hdr_i >= len(rows) or not rows[hdr_i] or rows[hdr_i][0] is None:
            return []
        headers = [h for h in rows[hdr_i] if h is not None]
        out = []
        i = hdr_i + 1
        while i < len(rows) and rows[i] and rows[i][0] is not None and not str(rows[i][0]).startswith('['):
            vals = rows[i]
            out.append({headers[j]: vals[j] for j in range(len(headers))})
            i += 1
        return out

    gi = find_section('[GEOMETRY]')
    if gi < 0:
        raise ValueError('Missing [GEOMETRY] section.')
    grow = read_table(gi)
    if not grow:
        raise ValueError('Empty [GEOMETRY] section.')
    g = grow[0]
    span = float(g['span_m']); rise = float(g['rise_m'])
    n_elem = int(g['n_elem']); shape_expr = str(g['shape_expr'])

    si = find_section('[SUPPORTS]')
    arch_type, hinge_frac = 'Two-hinged', 0.5
    if si >= 0:
        srow = read_table(si)
        if srow:
            arch_type = str(srow[0]['arch_type'])
            hinge_frac = float(srow[0]['hinge_frac']) if srow[0]['hinge_frac'] is not None else 0.5

    ui = find_section('[UNIFORM_LOADS]')
    w_weight, w_wind = 0.0, 0.0
    if ui >= 0:
        urow = read_table(ui)
        if urow:
            w_weight = float(urow[0]['w_weight_kNm'] or 0.0)
            w_wind = float(urow[0]['w_wind_kNm'] or 0.0)

    distributed_loads = []
    di = find_section('[DISTRIBUTED_LOADS]')
    if di >= 0:
        for r in read_table(di):
            distributed_loads.append({'expr': str(r['expr']), 'x1': float(r['x1_m']),
                                       'x2': float(r['x2_m']), 'direction': str(r['direction'])})

    point_loads = []
    pi = find_section('[POINT_LOADS]')
    if pi >= 0:
        for r in read_table(pi):
            point_loads.append({'x': float(r['x_m']), 'fx': float(r['fx_kN']), 'fy': float(r['fy_kN'])})

    profile = {'E': 200.0, 'A': 100.0, 'I': 20000.0, 'c_top': 20.0, 'c_bot': 20.0,
               'allow_tension': 16.0, 'allow_compression': 16.0}
    ci = find_section('[SECTION]')
    if ci >= 0:
        crow = read_table(ci)
        if crow:
            c = crow[0]
            key_map = [('E', 'E_GPa'), ('A', 'A_cm2'), ('I', 'I_cm4'),
                       ('c_top', 'c_top_cm'), ('c_bot', 'c_bot_cm'),
                       ('allow_tension', 'allow_tension_kNcm2'),
                       ('allow_compression', 'allow_compression_kNcm2')]
            for pk, xk in key_map:
                if c.get(xk) is not None:
                    profile[pk] = float(c[xk])

    return {'span': span, 'rise': rise, 'n_elem': n_elem, 'shape_expr': shape_expr,
            'arch_type': arch_type, 'hinge_frac': hinge_frac,
            'w_weight': w_weight, 'w_wind': w_wind,
            'point_loads': point_loads, 'distributed_loads': distributed_loads,
            'profile': profile}

class ArchApp(tk.Frame):
    """
    Arch tab — two-hinged, three-hinged, and fixed-base arches from a smooth
    y=f(x) shape, via the ArchModel direct-stiffness solver above.
    Units: lengths in m, forces in kN, moments in kN*m, distributed loads in
    kN/m (horizontal-projected), E in GPa, section I in cm^4, A in cm^2,
    c_top/c_bot (extreme fibre distances) in cm. Stresses in kN/cm^2, with
    tension positive (matches the sign convention used elsewhere in the app).
    """
    CARCH, CSUP, CLOAD, CHINGE = '#333333', '#555555', '#D85A30', '#c0392b'
    CN_, CM_, CT, CC = '#7F77DD', '#D85A30', '#e24b4a', '#378add'
    CGRID = '#e8e8e8'
    CTHH, CTHV, CTHR = '#2ecc71', '#8e44ad', '#1a6bbd'   # thrust vector: horiz / vert / resultant

    ARCH_TYPES = ['Two-hinged', 'Three-hinged', 'Fixed']
    SHAPE_PRESETS = ['Parabolic', 'Circular', 'Catenary-like', 'Custom']

    def __init__(self, master, **kw):
        super().__init__(master, bg='#f5f5f3', **kw)
        self.span = 10.0
        self.rise = 2.5
        self.n_elem = 50
        self.arch_type = 'Two-hinged'
        self.hinge_frac = 0.5
        self.shape_preset = 'Parabolic'
        self.shape_expr = '4*rise*x*(L-x)/L**2'
        self.w_weight = 10.0   # kN/m, horizontal-projected, +down
        self.w_wind = 5.0      # kN/m, horizontal-projected, +x
        self.point_loads = []  # [{'x','fx','fy'}] kN
        self.distributed_loads = []  # [{'expr','x1','x2','direction'}] kN/m over [x1,x2] m
        self._probe_point = None
        self._band_transforms = []
        self._last_samp = None
        self._tooltip = None
        self.profile = {'E': 200.0, 'A': 100.0, 'I': 20000.0,
                         'c_top': 20.0, 'c_bot': 20.0,
                         'allow_tension': 16.0, 'allow_compression': 16.0}
        self.result = None
        self.model = None
        self._build_ui()
        self._draw_schematic()

    # ── UI ────────────────────────────────────────────────────────────────────
    def _build_ui(self):
        tb = tk.Frame(self, bg='#ebebea')
        tb.pack(fill='x', padx=6, pady=(6, 0))
        tk.Label(tb, text='Span (m):', bg='#ebebea', font=('Helvetica', 11)).pack(side='left', padx=(4, 2))
        self.span_var = tk.DoubleVar(value=self.span)
        tk.Entry(tb, textvariable=self.span_var, width=6, font=('Helvetica', 11)).pack(side='left')
        tk.Label(tb, text='Rise (m):', bg='#ebebea', font=('Helvetica', 11)).pack(side='left', padx=(8, 2))
        self.rise_var = tk.DoubleVar(value=self.rise)
        tk.Entry(tb, textvariable=self.rise_var, width=6, font=('Helvetica', 11)).pack(side='left')
        tk.Button(tb, text='Set geometry', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._set_geometry).pack(side='left', padx=4)
        tk.Frame(tb, width=1, bg='#ccc').pack(side='left', fill='y', padx=6, pady=3)
        tk.Button(tb, text='Example: two-hinged', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._load_example_two_hinged).pack(side='left', padx=2)
        tk.Button(tb, text='Example: three-hinged', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._load_example_three_hinged).pack(side='left', padx=2)
        tk.Button(tb, text='Example: fixed', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._load_example_fixed).pack(side='left', padx=2)
        tk.Button(tb, text='Clear', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._clear_all).pack(side='left', padx=2)
        tk.Frame(tb, width=1, bg='#ccc').pack(side='left', fill='y', padx=6, pady=3)
        tk.Button(tb, text='Export Excel', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), fg='#1a6bbd', command=self._export_excel).pack(side='left', padx=2)
        tk.Button(tb, text='Import Excel', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), fg='#1a6bbd', command=self._import_excel).pack(side='left', padx=2)
        tk.Frame(tb, width=1, bg='#ccc').pack(side='left', fill='y', padx=6, pady=3)
        tk.Button(tb, text='▶ Analyze', relief='flat', bd=0, padx=10, pady=4,
                  bg='#1a6bbd', fg='white', font=('Helvetica', 11, 'bold'),
                  command=self._analyze).pack(side='left', padx=2)

        main = tk.Frame(self, bg='#f5f5f3')
        main.pack(fill='both', expand=True, padx=6, pady=6)

        SCHEM_SIZE = 400
        schem_outer = tk.Frame(main, bg='#f5f5f3', width=SCHEM_SIZE)
        schem_outer.pack(side='left', fill='y', padx=(0, 8))
        schem_outer.pack_propagate(False)

        tk.Label(schem_outer, text='Arch schematic', bg='#f5f5f3',
                 font=('Helvetica', 9, 'bold'), fg='#777').pack(anchor='w')
        schem_sq = tk.Frame(schem_outer, bg='#f5f5f3', width=SCHEM_SIZE, height=SCHEM_SIZE)
        schem_sq.pack(pady=(0, 8))
        schem_sq.pack_propagate(False)
        self.schem = tk.Canvas(schem_sq, bg='white', bd=1, relief='solid', highlightthickness=0)
        self.schem.pack(fill='both', expand=True)
        self.schem.bind('<Configure>', lambda e: self._draw_schematic())

        probe_hdr = tk.Frame(schem_outer, bg='#f5f5f3')
        probe_hdr.pack(fill='x', pady=(2, 0))
        self.probe_enabled_var = tk.BooleanVar(value=False)
        tk.Checkbutton(probe_hdr, text='Force probe', variable=self.probe_enabled_var,
                       bg='#f5f5f3', font=('Helvetica', 9, 'bold'),
                       command=self._on_probe_toggle).pack(side='left')

        probe_row = tk.Frame(schem_outer, bg='#f5f5f3')
        probe_row.pack(fill='x', pady=(2, 2))
        tk.Label(probe_row, text='s from support A (m):', bg='#f5f5f3',
                 font=('Helvetica', 9)).pack(side='left')
        self.probe_s_var = tk.DoubleVar(value=0.0)
        probe_entry = tk.Entry(probe_row, textvariable=self.probe_s_var, width=7, font=('Helvetica', 9))
        probe_entry.pack(side='left', padx=4)
        probe_entry.bind('<Return>', lambda e: self._update_probe())
        tk.Button(probe_row, text='Show', command=self._update_probe).pack(side='left', padx=2)

        self.probe_result_label = tk.Label(schem_outer, text='', bg='#f5f5f3',
                                           font=('Courier', 9), justify='left', anchor='w')
        self.probe_result_label.pack(fill='x', pady=(2, 4), padx=2)

        thrust_row = tk.Frame(schem_outer, bg='#f5f5f3')
        thrust_row.pack(fill='x', pady=(2, 0))
        self.thrust_enabled_var = tk.BooleanVar(value=False)
        tk.Checkbutton(thrust_row, text='Show thrust vectors', variable=self.thrust_enabled_var,
                       bg='#f5f5f3', font=('Helvetica', 9, 'bold'),
                       command=self._draw_schematic).pack(side='left')

        thrust_row2 = tk.Frame(schem_outer, bg='#f5f5f3')
        thrust_row2.pack(fill='x', pady=(2, 2))
        tk.Label(thrust_row2, text='stations:', bg='#f5f5f3', font=('Helvetica', 9)).pack(side='left')
        self.thrust_count_var = tk.IntVar(value=10)
        tk.Scale(thrust_row2, from_=3, to=24, orient='horizontal', variable=self.thrust_count_var,
                 length=110, showvalue=True, bg='#f5f5f3', bd=0, highlightthickness=0,
                 relief='flat', font=('Helvetica', 8),
                 command=lambda _v: self._draw_schematic()).pack(side='left', padx=4)

        thrust_legend = tk.Frame(schem_outer, bg='#f5f5f3')
        thrust_legend.pack(fill='x', pady=(0, 8))
        tk.Label(thrust_legend, text='━', fg=self.CTHH, bg='#f5f5f3', font=('Helvetica', 10, 'bold')).pack(side='left')
        tk.Label(thrust_legend, text='horizontal   ', bg='#f5f5f3', font=('Helvetica', 8)).pack(side='left')
        tk.Label(thrust_legend, text='━', fg=self.CTHV, bg='#f5f5f3', font=('Helvetica', 10, 'bold')).pack(side='left')
        tk.Label(thrust_legend, text='vertical   ', bg='#f5f5f3', font=('Helvetica', 8)).pack(side='left')
        tk.Label(thrust_legend, text='━', fg=self.CTHR, bg='#f5f5f3', font=('Helvetica', 10, 'bold')).pack(side='left')
        tk.Label(thrust_legend, text='resultant', bg='#f5f5f3', font=('Helvetica', 8)).pack(side='left')

        mid = tk.Frame(main, bg='#f5f5f3')
        mid.pack(side='left', fill='both', expand=True)

        tk.Label(mid, text='Diagrams (2×2): axial N · thrust Fx/Fy · bending M · fibre stress'
                            ' (run ▶ Analyze; hover over any arch outline for a live readout)',
                 bg='#f5f5f3', font=('Helvetica', 9, 'bold'), fg='#777').pack(anchor='w')

        scale_bar = tk.Frame(mid, bg='#f5f5f3')
        scale_bar.pack(fill='x', pady=(2, 0))
        scale_bar2 = tk.Frame(mid, bg='#f5f5f3')
        scale_bar2.pack(fill='x', pady=(0, 2))

        def _scale_slider(parent, label, var):
            f = tk.Frame(parent, bg='#f5f5f3')
            f.pack(side='left', padx=(0, 12))
            tk.Label(f, text=label, bg='#f5f5f3', font=('Helvetica', 9)).pack(side='left', padx=(0, 4))
            tk.Scale(f, from_=1, to=300, orient='horizontal', variable=var, length=80,
                     showvalue=False, bg='#f5f5f3', bd=0, highlightthickness=0,
                     relief='flat', font=('Helvetica', 8),
                     command=lambda v, var=var: self._on_scale_changed(var, v)).pack(side='left')
            entry = tk.Entry(f, textvariable=var, width=5, font=('Helvetica', 8))
            entry.pack(side='left', padx=(3, 0))
            entry.bind('<Return>', lambda e, var=var: self._on_scale_changed(var, var.get()))
            tk.Label(f, text='%', bg='#f5f5f3', font=('Helvetica', 8)).pack(side='left', padx=(1, 0))

        self.n_scale_var = tk.IntVar(value=100)
        self.thrust_scale_var = tk.IntVar(value=100)
        self.m_scale_var = tk.IntVar(value=100)
        self.sigma_scale_var = tk.IntVar(value=100)
        self._all_scale_vars = (self.n_scale_var, self.thrust_scale_var,
                                 self.m_scale_var, self.sigma_scale_var)
        self.scale_lock_var = tk.BooleanVar(value=False)

        _scale_slider(scale_bar, 'N scale %:', self.n_scale_var)
        _scale_slider(scale_bar, 'Fx/Fy scale %:', self.thrust_scale_var)

        self.thrust_view_var = tk.StringVar(value='Both')
        thrust_view_f = tk.Frame(scale_bar, bg='#f5f5f3')
        thrust_view_f.pack(side='left', padx=(0, 14))
        tk.Label(thrust_view_f, text='view:', bg='#f5f5f3', font=('Helvetica', 9)).pack(side='left', padx=(0, 4))
        for opt in ('Horizontal', 'Vertical', 'Both'):
            tk.Radiobutton(thrust_view_f, text=opt, variable=self.thrust_view_var, value=opt,
                          bg='#f5f5f3', font=('Helvetica', 8),
                          command=self._draw_diagrams).pack(side='left')

        tk.Checkbutton(scale_bar, text='🔒 Lock scales together', variable=self.scale_lock_var,
                       bg='#f5f5f3', font=('Helvetica', 8, 'bold')).pack(side='left', padx=(4, 0))

        _scale_slider(scale_bar2, 'M scale %:', self.m_scale_var)
        _scale_slider(scale_bar2, 'σ scale %:', self.sigma_scale_var)

        self.diag_canvas = tk.Canvas(mid, bg='#fafaf8', bd=1, relief='solid', highlightthickness=0)
        self.diag_canvas.pack(fill='both', expand=True)
        self.diag_canvas.bind('<Configure>', lambda e: self._draw_diagrams())
        self.diag_canvas.bind('<Motion>', self._on_diagram_motion)
        self.diag_canvas.bind('<Leave>', lambda e: self._hide_tooltip())

        panel_outer = tk.Frame(main, width=PANEL_W + 105, bg='#f0f0ee', bd=1, relief='solid')
        panel_outer.pack(side='right', fill='y', padx=(6, 0))
        panel_outer.pack_propagate(False)
        panel_canvas = tk.Canvas(panel_outer, bg='#f0f0ee', highlightthickness=0, width=PANEL_W + 105)
        panel_sb = tk.Scrollbar(panel_outer, orient='vertical', command=panel_canvas.yview)
        panel_canvas.configure(yscrollcommand=panel_sb.set)
        panel_sb.pack(side='right', fill='y')
        panel_canvas.pack(side='left', fill='both', expand=True)
        panel = tk.Frame(panel_canvas, width=PANEL_W + 105, bg='#f0f0ee')
        panel_canvas.create_window((0, 0), window=panel, anchor='nw', width=PANEL_W + 105)

        def _on_cfg(event):
            panel_canvas.configure(scrollregion=panel_canvas.bbox('all'))
        panel.bind('<Configure>', _on_cfg)

        def _wheel(event):
            panel_canvas.yview_scroll(int(-1 * (event.delta / 120)), 'units')
        panel_canvas.bind('<Enter>', lambda e: panel_canvas.bind_all('<MouseWheel>', _wheel))
        panel_canvas.bind('<Leave>', lambda e: panel_canvas.unbind_all('<MouseWheel>'))

        self._build_panel(panel)

    def _build_panel(self, panel):
        pad = dict(padx=8, pady=(8, 2))

        tk.Label(panel, text='SHAPE  y = f(x)', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        shp = tk.Frame(panel, bg='#f0f0ee'); shp.pack(fill='x', padx=8)
        self.shape_var = tk.StringVar(value=self.shape_preset)
        cb = ttk.Combobox(shp, textvariable=self.shape_var, values=self.SHAPE_PRESETS,
                           width=14, state='readonly')
        cb.grid(row=0, column=0, columnspan=2, sticky='w', pady=2)
        cb.bind('<<ComboboxSelected>>', lambda e: self._apply_shape_preset())
        tk.Label(shp, text='expression (vars: x, L, rise, R, k; ^ or ** = power):', bg='#f0f0ee',
                 font=('Helvetica', 8)).grid(row=1, column=0, columnspan=2, sticky='w', pady=(4, 0))
        self.expr_var = tk.StringVar(value=self.shape_expr)
        self.expr_entry = tk.Entry(shp, textvariable=self.expr_var, width=26, font=('Helvetica', 9))
        self.expr_entry.grid(row=2, column=0, columnspan=2, sticky='we', pady=2)
        tk.Label(shp, text='Elements (n):', bg='#f0f0ee', font=('Helvetica', 9)).grid(row=3, column=0, sticky='w', pady=(4, 0))
        self.nelem_var = tk.IntVar(value=self.n_elem)
        tk.Entry(shp, textvariable=self.nelem_var, width=6, font=('Helvetica', 9)).grid(row=3, column=1, sticky='w', pady=(4, 0))

        tk.Label(panel, text='SUPPORTS / HINGE', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        sup = tk.Frame(panel, bg='#f0f0ee'); sup.pack(fill='x', padx=8)
        self.arch_type_var = tk.StringVar(value=self.arch_type)
        cb2 = ttk.Combobox(sup, textvariable=self.arch_type_var, values=self.ARCH_TYPES,
                            width=14, state='readonly')
        cb2.grid(row=0, column=0, columnspan=2, sticky='w', pady=2)
        cb2.bind('<<ComboboxSelected>>', lambda e: self._on_arch_type_change())
        tk.Label(sup, text='Hinge location (fraction of span, 0–1):', bg='#f0f0ee',
                 font=('Helvetica', 8)).grid(row=1, column=0, columnspan=2, sticky='w', pady=(4, 0))
        self.hinge_var = tk.DoubleVar(value=self.hinge_frac)
        self.hinge_entry = tk.Entry(sup, textvariable=self.hinge_var, width=8, font=('Helvetica', 9))
        self.hinge_entry.grid(row=2, column=0, sticky='w', pady=2)
        self._on_arch_type_change()

        tk.Label(panel, text='LOADS (horizontal-projected)', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        ld = tk.Frame(panel, bg='#f0f0ee'); ld.pack(fill='x', padx=8)
        tk.Label(ld, text='Weight w (kN/m, +down):', bg='#f0f0ee', font=('Helvetica', 9)).grid(row=0, column=0, sticky='w', pady=1)
        self.wweight_var = tk.DoubleVar(value=self.w_weight)
        tk.Entry(ld, textvariable=self.wweight_var, width=8, font=('Helvetica', 9)).grid(row=0, column=1, pady=1, padx=4)
        tk.Label(ld, text='Wind w (kN/m, +x):', bg='#f0f0ee', font=('Helvetica', 9)).grid(row=1, column=0, sticky='w', pady=1)
        self.wwind_var = tk.DoubleVar(value=self.w_wind)
        tk.Entry(ld, textvariable=self.wwind_var, width=8, font=('Helvetica', 9)).grid(row=1, column=1, pady=1, padx=4)

        tk.Label(panel, text='NON-UNIFORM DISTRIBUTED LOADS', bg='#f0f0ee',
                 font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        tk.Label(panel, text='q(x) in kN/m, over a sub-domain [x₁,x₂] (m)', bg='#f0f0ee',
                 font=('Helvetica', 8), fg='#777').pack(anchor='w', padx=8)
        self.dl_tree = ttk.Treeview(panel, columns=('expr', 'x1', 'x2', 'dir'),
                                    show='headings', height=3)
        for c, w, lbl in [('expr', 100, 'q(x)'), ('x1', 45, 'x₁'), ('x2', 45, 'x₂'), ('dir', 55, 'dir')]:
            self.dl_tree.heading(c, text=lbl); self.dl_tree.column(c, width=w)
        self.dl_tree.pack(fill='x', padx=8)
        dlf = tk.Frame(panel, bg='#f0f0ee'); dlf.pack(fill='x', padx=8, pady=(2, 8))
        tk.Button(dlf, text='Add', command=self._add_distributed_load).pack(side='left', padx=2)
        tk.Button(dlf, text='Delete',
                  command=lambda: self._del_row(self.dl_tree, self.distributed_loads)).pack(side='left', padx=2)

        tk.Label(panel, text='POINT LOADS (kN)', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        self.pl_tree = ttk.Treeview(panel, columns=('x', 'fx', 'fy'), show='headings', height=3)
        for c, w in [('x', 60), ('fx', 60), ('fy', 60)]:
            self.pl_tree.heading(c, text=c); self.pl_tree.column(c, width=w)
        self.pl_tree.pack(fill='x', padx=8)
        pf = tk.Frame(panel, bg='#f0f0ee'); pf.pack(fill='x', padx=8, pady=(2, 8))
        tk.Button(pf, text='Add', command=self._add_pointload).pack(side='left', padx=2)
        tk.Button(pf, text='Delete', command=lambda: self._del_row(self.pl_tree, self.point_loads)).pack(side='left', padx=2)

        tk.Label(panel, text='CROSS-SECTION / MATERIAL', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        sec = tk.Frame(panel, bg='#f0f0ee'); sec.pack(fill='x', padx=8)
        self.sec_vars = {}
        fields = [('E', 'E (GPa)'), ('A', 'Area A (cm²)'), ('I', 'I (cm⁴)'),
                  ('c_top', 'c extrados (cm)'), ('c_bot', 'c intrados (cm)'),
                  ('allow_tension', 'Allow. tension σ (kN/cm²)'),
                  ('allow_compression', 'Allow. compression σ (kN/cm²)')]
        for i, (key, label) in enumerate(fields):
            tk.Label(sec, text=label, bg='#f0f0ee', font=('Helvetica', 9)).grid(row=i, column=0, sticky='w', pady=1)
            v = tk.DoubleVar(value=self.profile[key])
            self.sec_vars[key] = v
            tk.Entry(sec, textvariable=v, width=8, font=('Helvetica', 9)).grid(row=i, column=1, pady=1, padx=4)

        tk.Label(panel, text='RESULTS', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        self.res_text = tk.Text(panel, height=18, bg='white', font=('Courier', 9), relief='solid', bd=1)
        self.res_text.pack(fill='both', expand=True, padx=8, pady=(2, 10))

    # ── combobox-driven helpers ─────────────────────────────────────────────
    def _apply_shape_preset(self):
        preset = self.shape_var.get()
        L = self.span_var.get() or self.span
        rise = self.rise_var.get() or self.rise
        if preset == 'Parabolic':
            self.expr_var.set('4*rise*x*(L-x)/L**2')
            self.expr_entry.config(state='disabled')
        elif preset == 'Circular':
            self.expr_var.set('(rise - R) + sqrt(R**2 - (x-L/2)**2)')
            self.expr_entry.config(state='disabled')
        elif preset == 'Catenary-like':
            self.expr_var.set('rise*(1-(cosh(k*(x-L/2))-1)/(cosh(k*L/2)-1))')
            self.expr_entry.config(state='disabled')
        else:  # Custom
            self.expr_entry.config(state='normal')
        self._draw_schematic()

    def _on_arch_type_change(self):
        at = self.arch_type_var.get()
        if at == 'Three-hinged':
            self.hinge_entry.config(state='normal')
        else:
            self.hinge_entry.config(state='disabled')
        self._draw_schematic()

    # ── row management ──────────────────────────────────────────────────────
    def _del_row(self, tree, data_list):
        sel = tree.selection()
        if not sel:
            return
        idx = tree.index(sel[0])
        del data_list[idx]
        self._refresh_tables()

    def _refresh_tables(self):
        self.pl_tree.delete(*self.pl_tree.get_children())
        for r in self.point_loads:
            self.pl_tree.insert('', 'end', values=(r['x'], r['fx'], r['fy']))
        self.dl_tree.delete(*self.dl_tree.get_children())
        for r in self.distributed_loads:
            self.dl_tree.insert('', 'end', values=(r['expr'], r['x1'], r['x2'], r['direction']))
        self._draw_schematic()

    def _ask(self, title, fields):
        win = tk.Toplevel(self); win.title(title); win.grab_set()
        win.configure(bg='#f0f0ee')
        vars_ = {}
        for i, (key, label, default) in enumerate(fields):
            tk.Label(win, text=label, bg='#f0f0ee').grid(row=i, column=0, sticky='w', padx=8, pady=4)
            v = tk.DoubleVar(value=default)
            vars_[key] = v
            tk.Entry(win, textvariable=v, width=10).grid(row=i, column=1, padx=8, pady=4)
        result = {}

        def ok():
            for k, v in vars_.items():
                result[k] = v.get()
            win.destroy()
        tk.Button(win, text='OK', command=ok, bg='#1a6bbd', fg='white').grid(
            row=len(fields), column=0, columnspan=2, pady=8)
        win.wait_window()
        return result if result else None

    def _add_pointload(self):
        r = self._ask('Add point load', [('x', 'x (m)', self.span / 2),
                                          ('fx', 'Fx (kN, +→)', 0.0),
                                          ('fy', 'Fy (kN, +down)', 10.0)])
        if r: self.point_loads.append(r); self._refresh_tables()

    def _add_distributed_load(self):
        win = tk.Toplevel(self); win.title('Add non-uniform distributed load'); win.grab_set()
        win.configure(bg='#f0f0ee')
        tk.Label(win, text='q(x) in kN/m  (vars: x, L, rise, R, k; ^ or ** = power):', bg='#f0f0ee',
                 font=('Helvetica', 9)).grid(row=0, column=0, columnspan=2, sticky='w', padx=8, pady=(8, 2))
        expr_var = tk.StringVar(value='10*sin(pi*x/L)')
        tk.Entry(win, textvariable=expr_var, width=28, font=('Helvetica', 9)).grid(
            row=1, column=0, columnspan=2, sticky='we', padx=8, pady=2)

        tk.Label(win, text='x₁ (m):', bg='#f0f0ee', font=('Helvetica', 9)).grid(
            row=2, column=0, sticky='w', padx=8, pady=4)
        x1_var = tk.DoubleVar(value=0.0)
        tk.Entry(win, textvariable=x1_var, width=10).grid(row=2, column=1, padx=8, pady=4)

        tk.Label(win, text='x₂ (m):', bg='#f0f0ee', font=('Helvetica', 9)).grid(
            row=3, column=0, sticky='w', padx=8, pady=4)
        x2_var = tk.DoubleVar(value=self.span)
        tk.Entry(win, textvariable=x2_var, width=10).grid(row=3, column=1, padx=8, pady=4)

        tk.Label(win, text='Direction:', bg='#f0f0ee', font=('Helvetica', 9)).grid(
            row=4, column=0, sticky='w', padx=8, pady=4)
        dir_var = tk.StringVar(value='vertical')
        ttk.Combobox(win, textvariable=dir_var, values=['vertical', 'horizontal'],
                     width=10, state='readonly').grid(row=4, column=1, padx=8, pady=4)

        result = {}

        def ok():
            expr = expr_var.get().strip()
            try:
                ctx = {'L': self.span_var.get(), 'rise': self.rise_var.get(),
                       'R': (self.span_var.get()**2)/(8*self.rise_var.get()) + self.rise_var.get()/2
                            if self.rise_var.get() else 1e9,
                       'k': 3.0/self.span_var.get() if self.span_var.get() else 1.0}
                x1_, x2_ = x1_var.get(), x2_var.get()
                x_mid = (min(x1_, x2_) + max(x1_, x2_)) / 2
                make_shape_fn(expr, ctx)(x_mid)   # validate it compiles & evaluates
            except Exception as e:
                messagebox.showerror('Invalid expression', str(e)); return
            result['expr'] = expr
            result['x1'] = x1_var.get()
            result['x2'] = x2_var.get()
            result['direction'] = dir_var.get()
            win.destroy()

        tk.Button(win, text='OK', command=ok, bg='#1a6bbd', fg='white').grid(
            row=5, column=0, columnspan=2, pady=8)
        win.wait_window()
        if result:
            self.distributed_loads.append(result)
            self._refresh_tables()

    def _set_geometry(self):
        self.span = self.span_var.get()
        self.rise = self.rise_var.get()
        if self.shape_var.get() != 'Custom':
            self._apply_shape_preset()
        self._draw_schematic()

    def _clear_all(self):
        self.point_loads = []
        self.distributed_loads = []
        self.result = None; self.model = None
        self._probe_point = None
        self._band_transforms = []
        self._last_samp = None
        self._hide_tooltip()
        self._refresh_tables()
        self.res_text.delete('1.0', 'end')
        self.diag_canvas.delete('all')
        if hasattr(self, 'probe_result_label'):
            self.probe_result_label.config(text='')

    # ── Excel export / import ────────────────────────────────────────────────
    def _current_state(self):
        return {
            'span': self.span_var.get(), 'rise': self.rise_var.get(),
            'n_elem': int(self.nelem_var.get()), 'shape_expr': self.expr_var.get(),
            'arch_type': self.arch_type_var.get(), 'hinge_frac': self.hinge_var.get(),
            'w_weight': self.wweight_var.get(), 'w_wind': self.wwind_var.get(),
            'point_loads': self.point_loads, 'distributed_loads': self.distributed_loads,
            'profile': {k: v.get() for k, v in self.sec_vars.items()},
        }

    def _export_excel(self):
        if not _ensure_openpyxl():
            messagebox.showerror(
                'Missing library',
                'Could not install openpyxl automatically.\n\n'
                'Please open a terminal and run:\n'
                '    pip install openpyxl\n'
                'then try again.')
            return
        path = filedialog.asksaveasfilename(
            defaultextension='.xlsx',
            filetypes=[('Excel workbook', '*.xlsx')],
            initialfile='arch_report.xlsx',
            title='Save Excel report')
        if not path:
            return
        try:
            export_arch_excel(self._current_state(), path,
                               result=self.result, model=self.model)
            messagebox.showinfo(
                'Exported', f'Report saved to:\n{path}\n\n'
                'Includes a "Model" sheet — use "Import" to rebuild this '
                'exact arch from the file later.' +
                ('' if self.result else '\n\n(Run ▶ Analyze first to also '
                                         'include a "Results" sheet next time.)'))
        except Exception as e:
            messagebox.showerror('Export failed', str(e))

    def _import_excel(self):
        if not _ensure_openpyxl():
            messagebox.showerror(
                'Missing library',
                'Could not install openpyxl automatically.\n\n'
                'Please open a terminal and run:\n'
                '    pip install openpyxl\n'
                'then try again.')
            return
        path = filedialog.askopenfilename(
            filetypes=[('Excel workbook', '*.xlsx')],
            title='Import arch from Excel')
        if not path:
            return
        try:
            st = import_arch_excel(path)
        except Exception as e:
            messagebox.showerror('Import failed', str(e)); return

        self._clear_all()
        self.span = st['span']; self.rise = st['rise']
        self.span_var.set(st['span']); self.rise_var.set(st['rise'])
        self.nelem_var.set(st['n_elem'])
        self.shape_var.set('Custom')
        self.expr_entry.config(state='normal')
        self.expr_var.set(st['shape_expr'])
        self.arch_type_var.set(st['arch_type'])
        self._on_arch_type_change()
        self.hinge_var.set(st['hinge_frac'])
        self.wweight_var.set(st['w_weight'])
        self.wwind_var.set(st['w_wind'])
        self.point_loads = st['point_loads']
        self.distributed_loads = st['distributed_loads']
        for k, v in st['profile'].items():
            if k in self.sec_vars:
                self.sec_vars[k].set(v)
                self.profile[k] = v
        self._refresh_tables()
        self._draw_schematic()

    # ── examples ─────────────────────────────────────────────────────────────
    def _load_example_two_hinged(self):
        self._clear_all()
        self.span = 20.0; self.rise = 5.0
        self.span_var.set(20.0); self.rise_var.set(5.0)
        self.shape_var.set('Parabolic'); self._apply_shape_preset()
        self.arch_type_var.set('Two-hinged'); self._on_arch_type_change()
        self.wweight_var.set(10.0); self.wwind_var.set(3.0)
        self._refresh_tables()

    def _load_example_three_hinged(self):
        self._clear_all()
        self.span = 20.0; self.rise = 5.0
        self.span_var.set(20.0); self.rise_var.set(5.0)
        self.shape_var.set('Parabolic'); self._apply_shape_preset()
        self.arch_type_var.set('Three-hinged'); self.hinge_var.set(0.5)
        self._on_arch_type_change()
        self.wweight_var.set(10.0); self.wwind_var.set(3.0)
        self._refresh_tables()

    def _load_example_fixed(self):
        self._clear_all()
        self.span = 16.0; self.rise = 3.0
        self.span_var.set(16.0); self.rise_var.set(3.0)
        self.shape_var.set('Circular'); self._apply_shape_preset()
        self.arch_type_var.set('Fixed'); self._on_arch_type_change()
        self.wweight_var.set(12.0); self.wwind_var.set(4.0)
        self._refresh_tables()

    # ── analysis ─────────────────────────────────────────────────────────────
    def _analyze(self):
        try:
            self.span = self.span_var.get()
            self.rise = self.rise_var.get()
            self.n_elem = max(4, int(self.nelem_var.get()))
            ctx = {'L': self.span, 'rise': self.rise,
                   'R': (self.span**2)/(8*self.rise) + self.rise/2 if self.rise else 1e9,
                   'k': 3.0/self.span if self.span else 1.0}
            expr = self.expr_var.get()
            shape_fn = make_shape_fn(expr, ctx)

            m = ArchModel(shape_fn, self.span, n_elem=self.n_elem)
            at = self.arch_type_var.get()
            if at == 'Two-hinged':
                m.support_a = 'pin'; m.support_b = 'pin'; m.hinge_x = None
            elif at == 'Three-hinged':
                m.support_a = 'pin'; m.support_b = 'pin'
                m.hinge_x = max(0.0, min(1.0, self.hinge_var.get())) * self.span
            else:
                m.support_a = 'fixed'; m.support_b = 'fixed'; m.hinge_x = None

            for k in self.sec_vars:
                self.profile[k] = self.sec_vars[k].get()
            m.E = self.profile['E'] * 1e9
            m.A = self.profile['A'] * 1e-4
            m.I = self.profile['I'] * 1e-8
            m.c_top = self.profile['c_top'] * 1e-2
            m.c_bot = self.profile['c_bot'] * 1e-2

            self.w_weight = self.wweight_var.get()
            self.w_wind = self.wwind_var.get()
            m.w_weight = self.w_weight * 1e3
            m.w_wind = self.w_wind * 1e3
            m.point_loads = [{'x': p['x'], 'fx': p['fx']*1e3, 'fy': -p['fy']*1e3}
                              for p in self.point_loads]

            dloads = []
            for p in self.distributed_loads:
                qfn = make_shape_fn(p['expr'], ctx)
                x1 = max(0.0, min(self.span, p['x1']))
                x2 = max(0.0, min(self.span, p['x2']))
                if x2 < x1:
                    x1, x2 = x2, x1

                def scaled_fn(x, _qfn=qfn):
                    return _qfn(x) * 1e3   # kN/m -> N/m

                dloads.append({'fn': scaled_fn, 'x1': x1, 'x2': x2,
                                'direction': p['direction']})
            m.distributed_loads = dloads

            self.result = m.solve()
            self.model = m
            self._draw_diagrams()
            self._show_results()
            self._draw_schematic()
        except Exception as e:
            messagebox.showerror('Analysis failed', str(e))

    def _show_results(self):
        r = self.result; m = self.model
        Rx0, Ry0, M0 = r.reaction(0)
        Rxn, Ryn, Mn = r.reaction(m.n_elem)
        s, x, y, sig_top, sig_bot = r.fiber_stresses()
        forces = r.internal_forces()
        Nmax_c = min(f['N'] for f in forces)   # most negative = max compression
        Nmax_t = max(f['N'] for f in forces)
        Mmax = max((abs(f['Ma']) for f in forces), default=0.0)
        Mmax = max(Mmax, max((abs(f['Mb']) for f in forces), default=0.0))

        sig_all = sig_top + sig_bot
        max_tension = max(sig_all)
        max_compression = min(sig_all)

        allow_t = self.profile['allow_tension']
        allow_c = self.profile['allow_compression']
        t_kncm2 = max_tension * 1e-7
        c_kncm2 = max_compression * 1e-7
        t_ratio = t_kncm2 / allow_t if allow_t else 0
        c_ratio = abs(c_kncm2) / allow_c if allow_c else 0

        lines = ['REACTIONS']
        lines.append(f"  Left  (x=0)       : Rx={Rx0/1e3:+8.2f} kN  Ry={Ry0/1e3:+8.2f} kN"
                      + (f"  M={M0/1e3:+8.2f} kN·m" if m.support_a == 'fixed' else ''))
        lines.append(f"  Right (x={m.span:.2f}) : Rx={Rxn/1e3:+8.2f} kN  Ry={Ryn/1e3:+8.2f} kN"
                      + (f"  M={Mn/1e3:+8.2f} kN·m" if m.support_b == 'fixed' else ''))
        lines += ['', f'Max axial (tension)     N = {Nmax_t/1e3:8.2f} kN',
                  f'Max axial (compression) N = {Nmax_c/1e3:8.2f} kN',
                  f'Max |bending moment|    M = {Mmax/1e3:8.2f} kN·m', '',
                  'STRESS CHECK (extreme fibres, tension +)']
        lines.append(f'  Max tension     σ = {t_kncm2:7.3f} kN/cm²   allow = {allow_t:.3f} '
                      f'({"OK" if t_ratio <= 1.0 else "FAIL"}, ratio {t_ratio:.2f})')
        lines.append(f'  Max compression σ = {c_kncm2:7.3f} kN/cm²   allow = {allow_c:.3f} '
                      f'({"OK" if c_ratio <= 1.0 else "FAIL"}, ratio {c_ratio:.2f})')

        self.res_text.delete('1.0', 'end')
        self.res_text.insert('1.0', '\n'.join(lines))

    # ── drawing ──────────────────────────────────────────────────────────────
    def _current_shape_fn(self):
        try:
            L = self.span_var.get(); rise = self.rise_var.get()
            ctx = {'L': L, 'rise': rise,
                   'R': (L**2)/(8*rise) + rise/2 if rise else 1e9,
                   'k': 3.0/L if L else 1.0}
            return make_shape_fn(self.expr_var.get(), ctx), L
        except Exception:
            return None, self.span

    def _draw_schematic(self):
        c = self.schem
        c.delete('all')
        w = c.winfo_width() or 400
        h = c.winfo_height() or 400
        fn, L = self._current_shape_fn()
        L = max(L, 1e-6)
        if fn is None:
            return
        n_s = 100
        xs_s = [L*i/n_s for i in range(n_s+1)]
        try:
            ys_s = [fn(x) for x in xs_s]
        except Exception:
            return

        minx, maxx = 0.0, L
        miny, maxy = min(ys_s + [0.0]), max(ys_s + [0.0])
        span_x = max(maxx - minx, 1e-6)
        span_y = max(maxy - miny, 1e-6)

        pad_side = 30
        pad_top = 26
        pad_bottom = 55   # room for support symbols + axis labels
        avail_w = max(w - 2*pad_side, 10)
        avail_h = max(h - pad_top - pad_bottom, 10)
        scale = min(avail_w / span_x, avail_h / span_y)   # uniform: fits the arch, no distortion

        offx = pad_side + (avail_w - span_x*scale)/2 - minx*scale
        offy_top = pad_top + (avail_h - span_y*scale)/2

        def X(x): return offx + x*scale
        def Y(y): return offy_top + (maxy - y)*scale

        pts = [(X(x), Y(y)) for x, y in zip(xs_s, ys_s)]
        flat = [v for p in pts for v in p]
        c.create_line(*flat, fill=self.CARCH, width=3, smooth=True)

        at = self.arch_type_var.get()
        ends = [(xs_s[0], ys_s[0]), (xs_s[-1], ys_s[-1])]
        for x0, y0 in ends:
            typ = 'fixed' if at == 'Fixed' else 'pin'
            sx, sy = X(x0), Y(y0)
            if typ == 'fixed':
                c.create_line(sx-14, sy+18, sx+14, sy+18, width=3, fill=self.CSUP)
                for k in range(-2, 3):
                    c.create_line(sx+k*6, sy+18, sx+k*6-6, sy+26, fill=self.CSUP)
            else:
                c.create_polygon(sx-10, sy+18, sx+10, sy+18, sx, sy, fill='', outline=self.CSUP, width=2)

        if at == 'Three-hinged':
            hf = self.hinge_var.get() if hasattr(self, 'hinge_var') else 0.5
            hx = max(0.0, min(1.0, hf)) * L
            try:
                hy = fn(hx)
            except Exception:
                hy = maxy
            c.create_oval(X(hx)-6, Y(hy)-6, X(hx)+6, Y(hy)+6, fill=self.CHINGE, outline='')

        c.create_text(X(0), h-14, text='0', font=('Helvetica', 8), fill='#555')
        c.create_text(X(L), h-14, text=f'{L:.1f} m', font=('Helvetica', 8), fill='#555')

        if self.probe_enabled_var.get() and self._probe_point is not None:
            px, py = self._probe_point
            sx, sy = X(px), Y(py)
            c.create_line(sx, pad_top, sx, sy, fill='#1a6bbd', dash=(3, 2))
            c.create_oval(sx-6, sy-6, sx+6, sy+6, fill='#1a6bbd', outline='white', width=1.5)

        if self.thrust_enabled_var.get() and self.result is not None:
            self._draw_thrust_vectors(c, X, Y, L)

    def _draw_thrust_vectors(self, c, X, Y, L):
        """Overlays the internal thrust vector — decomposed into its global
        horizontal and vertical components, plus the resultant — at evenly
        spaced stations along the arch. Vector lengths share one common
        force→length scale (set by the single largest resultant among the
        displayed stations), so relative magnitudes stay comparable at a
        glance; drawn in the arch's own model coordinates through the same
        X()/Y() transform as the arch itself, so the overlay always lines up
        and rescales together with the schematic."""
        samp = self.result.sample()
        s_max = samp['s'][-1]
        n_vec = max(3, self.thrust_count_var.get())

        stations = []
        for k in range(n_vec):
            s_target = (k + 0.5) / n_vec * s_max
            best_i, best_d = 0, None
            for i, sv in enumerate(samp['s']):
                d = abs(sv - s_target)
                if best_d is None or d < best_d:
                    best_d, best_i = d, i
            N = samp['N'][best_i]; V = samp['V'][best_i]
            tc, ts = samp['ec'][best_i], samp['es'][best_i]
            Fx = -N*tc - V*ts
            Fy = V*tc - N*ts
            stations.append((samp['x'][best_i], samp['y'][best_i], Fx, Fy))

        max_R = max((math.hypot(fx, fy) for _, _, fx, fy in stations), default=0.0)
        if max_R < 1e-9:
            return
        force_scale = (0.22 * L) / max_R   # metres of arrow length per newton

        for x0, y0, Fx, Fy in stations:
            fx_m, fy_m = Fx * force_scale, Fy * force_scale
            ox, oy = X(x0), Y(y0)

            # horizontal component
            hx, hy = X(x0 + fx_m), Y(y0)
            if abs(fx_m) > 1e-6:
                c.create_line(ox, oy, hx, hy, fill=self.CTHH, width=2,
                              arrow='last', arrowshape=(8, 10, 3))
            # vertical component
            vx, vy = X(x0), Y(y0 + fy_m)
            if abs(fy_m) > 1e-6:
                c.create_line(ox, oy, vx, vy, fill=self.CTHV, width=2,
                              arrow='last', arrowshape=(8, 10, 3))
            # resultant
            rx, ry = X(x0 + fx_m), Y(y0 + fy_m)
            c.create_line(ox, oy, rx, ry, fill=self.CTHR, width=2,
                          arrow='last', arrowshape=(9, 11, 4))
            c.create_oval(ox-3, oy-3, ox+3, oy+3, fill='#333', outline='')

    # ── force probe (schematic) ─────────────────────────────────────────────
    def _on_scale_changed(self, var, value):
        """Handles both slider-drag and Enter-in-entry updates for the
        diagram scale controls. When 'Lock scales together' is checked, the
        new value is applied to all four scales at once (N, Fx/Fy, M, σ);
        otherwise only the slider/entry that changed is updated. Typed
        values are not clamped to the slider's own 1-300 visual range —
        the slider just shows pinned at its end while the actual amplitude
        used for drawing follows the typed number exactly."""
        try:
            v = int(round(float(value)))
        except (ValueError, TypeError):
            v = var.get()
        v = max(1, v)
        if self.scale_lock_var.get():
            for ov in self._all_scale_vars:
                ov.set(v)
        else:
            var.set(v)
        self._draw_diagrams()

    def _on_probe_toggle(self):
        if not self.probe_enabled_var.get():
            self.probe_result_label.config(text='')
            self._probe_point = None
            self._draw_schematic()
        else:
            self._update_probe()

    def _update_probe(self):
        if not self.probe_enabled_var.get():
            return
        if not self.result:
            self.probe_result_label.config(text='Run ▶ Analyze first.')
            return
        samp = self.result.sample()
        s_max = samp['s'][-1]
        s_target = max(0.0, min(s_max, self.probe_s_var.get()))
        best_i, best_d = 0, None
        for i, sv in enumerate(samp['s']):
            d = abs(sv - s_target)
            if best_d is None or d < best_d:
                best_d, best_i = d, i

        N = samp['N'][best_i]; V = samp['V'][best_i]
        tc, ts = samp['ec'][best_i], samp['es'][best_i]
        Fx = -N*tc - V*ts
        Fy = V*tc - N*ts
        R = math.hypot(Fx, Fy)

        self._probe_point = (samp['x'][best_i], samp['y'][best_i])
        self.probe_result_label.config(text=(
            f"s = {samp['s'][best_i]:.2f} m  (x={samp['x'][best_i]:.2f}, y={samp['y'][best_i]:.2f})\n"
            f"Fx = {Fx/1e3:+8.2f} kN\n"
            f"Fy = {Fy/1e3:+8.2f} kN\n"
            f"R  = {R/1e3:8.2f} kN"))
        self._draw_schematic()

    # ── hover tooltip (diagrams) ─────────────────────────────────────────────
    def _show_tooltip(self, x_root, y_root, text):
        if self._tooltip is None:
            self._tooltip = tk.Toplevel(self)
            self._tooltip.overrideredirect(True)
            try:
                self._tooltip.attributes('-topmost', True)
            except Exception:
                pass
            self._tt_label = tk.Label(self._tooltip, text=text, bg='#ffffe0', fg='#222',
                                      font=('Helvetica', 8), justify='left',
                                      bd=1, relief='solid', padx=6, pady=4)
            self._tt_label.pack()
        else:
            self._tt_label.config(text=text)
        self._tooltip.geometry(f'+{x_root+16}+{y_root+12}')

    def _hide_tooltip(self):
        if self._tooltip is not None:
            self._tooltip.destroy()
            self._tooltip = None

    def _on_diagram_motion(self, event):
        if not self.result or not self._band_transforms:
            self._hide_tooltip()
            return
        mx, my = event.x, event.y
        band = None
        for bt in self._band_transforms:
            if bt['top'] <= my <= bt['bot'] and bt['left'] <= mx <= bt['right']:
                band = bt
                break
        if band is None:
            self._hide_tooltip()
            return

        samp = self._last_samp
        offx, offy_top, scale, maxy = band['offx'], band['offy_top'], band['scale'], band['maxy']
        best_i, best_d2 = 0, None
        for i in range(len(samp['x'])):
            px = offx + samp['x'][i]*scale
            py = offy_top + (maxy - samp['y'][i])*scale
            d2 = (px-mx)**2 + (py-my)**2
            if best_d2 is None or d2 < best_d2:
                best_d2, best_i = d2, i
        if best_d2 is None or best_d2 > 18**2:
            self._hide_tooltip()
            return

        N = samp['N'][best_i]; V = samp['V'][best_i]
        M = samp['M'][best_i]
        tc, ts = samp['ec'][best_i], samp['es'][best_i]
        Fx = -N*tc - V*ts
        Fy = V*tc - N*ts
        R = math.hypot(Fx, Fy)

        lines = [f"s = {samp['s'][best_i]:.2f} m",
                 f"Fx = {Fx/1e3:+.2f} kN   Fy = {Fy/1e3:+.2f} kN",
                 f"R  = {R/1e3:.2f} kN"]
        if band['kind'] == 'N':
            lines.append(f"N = {N/1e3:+.2f} kN")
        elif band['kind'] == 'M':
            lines.append(f"M = {M/1e3:+.2f} kN·m")
        elif band['kind'] == 'sigma':
            A, I = self.model.A, self.model.I
            ct, cb = self.model.c_top, self.model.c_bot
            sig_top = N/A - M*ct/I
            sig_bot = N/A + M*cb/I
            lines.append(f"σ_top = {sig_top*1e-7:+.3f} kN/cm²")
            lines.append(f"σ_bot = {sig_bot*1e-7:+.3f} kN/cm²")
        # 'Fx' / 'Fy' bands: already shown in the shared Fx/Fy line above

        self._show_tooltip(event.x_root, event.y_root, '\n'.join(lines))

    def _draw_diagrams(self):
        c = self.diag_canvas
        c.delete('all')
        if not self.result:
            self._band_transforms = []
            self._last_samp = None
            self._hide_tooltip()
            return
        w = c.winfo_width() or 600
        h = c.winfo_height() or 420
        samp = self.result.sample()
        A, I = self.model.A, self.model.I
        ct, cb = self.model.c_top, self.model.c_bot
        sig_top = [(N/A - M*ct/I) for N, M in zip(samp['N'], samp['M'])]
        sig_bot = [(N/A + M*cb/I) for N, M in zip(samp['N'], samp['M'])]
        Fx_list = [-N*ec - V*es for N, V, ec, es in zip(samp['N'], samp['V'], samp['ec'], samp['es'])]
        Fy_list = [V*ec - N*es for N, V, ec, es in zip(samp['N'], samp['V'], samp['ec'], samp['es'])]

        self._last_samp = samp
        self._band_transforms = []

        col_w = w / 2
        row_h = h / 2
        c.create_line(col_w, 0, col_w, h, fill=self.CGRID)
        c.create_line(0, row_h, w, row_h, fill=self.CGRID)

        thrust_view = self.thrust_view_var.get()
        thrust_curves = []
        if thrust_view in ('Horizontal', 'Both'):
            thrust_curves.append({'vals': [v/1e3 for v in Fx_list], 'color': self.CTHH,
                                   'sign_color': False, 'dash': False, 'label': 'Fx'})
        if thrust_view in ('Vertical', 'Both'):
            thrust_curves.append({'vals': [v/1e3 for v in Fy_list], 'color': self.CTHV,
                                   'sign_color': False, 'dash': (thrust_view == 'Both'), 'label': 'Fy'})

        quads = [
            ('AXIAL FORCE N  (kN) — red = tension, blue = compression',
             [{'vals': [v/1e3 for v in samp['N']], 'color': None, 'sign_color': True, 'dash': False}],
             (0, 0, col_w, row_h), self.n_scale_var.get()/100.0, 'N'),
            (f'THRUST  Fx / Fy  (kN, global components) — showing: {thrust_view}',
             thrust_curves, (col_w, 0, w, row_h), self.thrust_scale_var.get()/100.0, 'thrust'),
            ('BENDING MOMENT M  (kN·m)',
             [{'vals': [v/1e3 for v in samp['M']], 'color': self.CM_, 'sign_color': False, 'dash': False}],
             (0, row_h, col_w, h), self.m_scale_var.get()/100.0, 'M'),
            ('FIBRE STRESS  (kN/cm²) — red = tension, blue = compression, solid = extrados, dashed = intrados',
             [{'vals': [v*1e-7 for v in sig_top], 'color': None, 'sign_color': True, 'dash': False},
              {'vals': [v*1e-7 for v in sig_bot], 'color': None, 'sign_color': True, 'dash': True}],
             (col_w, row_h, w, h), self.sigma_scale_var.get()/100.0, 'sigma'),
        ]
        for label, curves, box, user_scale, kind in quads:
            left, top, right, bot = box
            c.create_text(left + 6, top + 10, text=label, anchor='w',
                          font=('Helvetica', 8, 'bold'), fill='#555', width=right-left-16)
            if curves:
                self._draw_curve_diagram(c, samp, curves, (left, top+16, right, bot), user_scale, kind)
            if kind == 'thrust':
                Rx0, Ry0, _ = self.result.reaction(0)
                Rxn, Ryn, _ = self.result.reaction(self.model.n_elem)
                yy = bot - 8
                c.create_text(left + 6, yy, anchor='w', font=('Helvetica', 8, 'bold'), fill='#c0392b',
                              text=f"◄ R_left: Rx={Rx0/1e3:+.2f} kN, Ry={Ry0/1e3:+.2f} kN ▲")
                c.create_text(right - 6, yy, anchor='e', font=('Helvetica', 8, 'bold'), fill='#c0392b',
                              text=f"► R_right: Rx={Rxn/1e3:+.2f} kN, Ry={Ryn/1e3:+.2f} kN ▲")

    def _draw_curve_diagram(self, c, samp, curves, box, user_scale=1.0, kind=None):
        """Draws the arch centerline plus one or more value-curves offset
        perpendicular to the local tangent at every station — so each curve
        literally follows the arch's own path instead of a flat axis. All
        curves in `curves` share one amplitude normalization (their combined
        max) and one geometric scale, so multiple quantities (e.g. Fx and Fy
        overlaid) stay directly comparable in the same panel.
        The arch's own geometric scale (pixels per metre) is computed from
        the arch centerline alone and stays fixed; only the offset amplitude
        responds to `user_scale`, so zooming the diagram in/out never
        changes how big the arch itself is drawn."""
        left, top, right, bot = box
        xs_, ys_, tc, ts = samp['x'], samp['y'], samp['tc'], samp['ts']
        n = len(xs_)
        if n < 2 or bot <= top or right <= left or not curves:
            return
        allvals = [v for curve in curves for v in curve['vals']]
        maxabs = max(1e-9, max(abs(v) for v in allvals))
        span_geo = max(max(xs_) - min(xs_), max(ys_) - min(ys_), 1e-6)
        offset_scale = 0.35 * span_geo * user_scale

        minx, maxx = min(xs_), max(xs_)
        miny, maxy = min(ys_), max(ys_)
        span_x = max(maxx - minx, 1e-6); span_y = max(maxy - miny, 1e-6)
        margin, pad = 8, 14
        avail_w = max((right-left) - 2*margin - 2*pad, 1)
        avail_h = max((bot-top) - 2*pad, 1)
        scale = min(avail_w / span_x, avail_h / span_y)   # arch-only fit, independent of diagram amplitude
        offx = left + margin + pad + (avail_w - span_x*scale)/2 - minx*scale
        offy_top = top + pad + (avail_h - span_y*scale)/2

        self._band_transforms.append({'top': top, 'bot': bot, 'left': left, 'right': right,
                                       'offx': offx, 'offy_top': offy_top, 'scale': scale,
                                       'maxy': maxy, 'kind': kind})

        def X(px): return offx + px*scale
        def Y(py): return offy_top + (maxy - py)*scale

        def offset_pts(values):
            pts = []
            for x, y, cc, ss, v in zip(xs_, ys_, tc, ts, values):
                nx_, ny_ = -ss, cc
                off = (v / maxabs) * offset_scale
                pts.append((x + nx_*off, y + ny_*off))
            return pts

        # value gridlines that follow the arch's own shape: each is the arch
        # centerline offset by a constant "nice" tick value, using the same
        # perpendicular-offset transform as the real data curves, so they
        # read consistently with the plotted quantity even though the arch
        # itself is curved.
        v_ticks = _nice_ticks(-maxabs, maxabs, 5)
        for vt in v_ticks:
            gpts = [(X(px), Y(py)) for px, py in offset_pts([vt] * n)]
            flat_g = [v for pt in gpts for v in pt]
            gcolor = '#bbbbbb' if abs(vt) > 1e-9 else '#999999'
            c.create_line(*flat_g, fill=gcolor, width=1, dash=(2, 2))
            c.create_text(gpts[0][0] - 4, gpts[0][1], text=f'{vt:g}', anchor='e',
                          font=('Helvetica', 6), fill='#999')

        arch_pts = [(X(x), Y(y)) for x, y in zip(xs_, ys_)]
        for i in range(n - 1):
            c.create_line(*arch_pts[i], *arch_pts[i+1], fill='#999', width=1)

        def draw_offset(off_pts, values, color, sign_color, dashed):
            pts = [(X(px), Y(py)) for px, py in off_pts]
            dash = (4, 3) if dashed else None
            for i in range(n - 1):
                v_avg = (values[i] + values[i+1]) / 2
                if sign_color:
                    seg_color = self.CT if v_avg >= 0 else self.CC
                else:
                    seg_color = color
                quad = [arch_pts[i], arch_pts[i+1], pts[i+1], pts[i]]
                flat = [v for pt in quad for v in pt]
                c.create_polygon(flat, fill=seg_color, outline='', stipple='gray50')
                if dash:
                    c.create_line(*pts[i], *pts[i+1], fill=seg_color, width=2, dash=dash)
                else:
                    c.create_line(*pts[i], *pts[i+1], fill=seg_color, width=2)

        for curve in curves:
            off_pts = offset_pts(curve['vals'])
            draw_offset(off_pts, curve['vals'], curve.get('color'),
                        curve.get('sign_color', False), curve.get('dash', False))
            locs, _ = _find_diagram_maxima(samp['s'], curve['vals'])
            for sv in locs:
                idx = min(range(n), key=lambda i: abs(samp['s'][i] - sv))
                mx, my = off_pts[idx]
                sx, sy = X(mx), Y(my)
                c.create_oval(sx - 4, sy - 4, sx + 4, sy + 4, fill='#222222', outline='white', width=1)
                c.create_text(sx, sy - 10, text=f"s={sv:.2f}", font=('Helvetica', 7), fill='#222222')

        c.create_text(right - 8, top - 6, text=f"max ±{maxabs:.2f}", anchor='e',
                      font=('Helvetica', 8), fill='#777')
