"""The piece-mark schedule as a workbook sheet.

The counterpart to stereo_groups_excel: that sheet is EDITABLE, and
importing it writes back onto the model. This one is not. A mark is derived
from the geometry every time (see stereo_marks), so a mark typed into a
spreadsheet would be a number with nothing behind it -- and the one way a
wrong mark reaches the shop floor is by surviving the edit that made it
wrong. The sheet says so at the top, and nothing here reads it back.
"""
from apps.stereo import stereo_marks as sm

SHEET = 'Piece Marks'

ASSEMBLY_HEADERS = ('mark', 'qty', 'rods each', 'length each (m)',
                    'total length (m)', 'mass each (kg)', 'total mass (kg)',
                    'mirrored', 'groups')
PART_HEADERS = ('mark', 'qty', 'profile', 'conn', 'length each (m)',
                'total length (m)', 'mass each (kg)', 'total mass (kg)',
                'role(s)')


def write_marks_sheet(wb, nodes, members, groups=(), tol_mm=None,
                      register=None):
    """Add the Piece Marks sheet. Returns (sheet, the schedule dict)."""
    from openpyxl.styles import Font, PatternFill, Alignment
    data = sm.schedule(nodes, members, groups or (), tol_mm, register)
    ws = wb.create_sheet(SHEET)
    hdr_fill = PatternFill('solid', fgColor='404040')

    def band(row, title):
        c = ws.cell(row=row, column=1, value=title)
        c.font = Font(bold=True, size=11, color='1F4E79')
        return row + 1

    def head(row, headers):
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = Font(bold=True, color='FFFFFF', size=10)
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal='center')
        return row + 1

    def note(row, text):
        ws.cell(row=row, column=1, value=text).font = Font(
            italic=True, size=9, color='5F6368')
        return row + 1

    r = band(1, 'PIECE MARKS — identical pieces, and how many of each')
    r = note(r, 'Read-only. Marks are worked out from the geometry every '
                'time this is exported, never stored, so they always agree '
                'with the steel. Editing this sheet changes nothing on the '
                'way back in.')
    r = note(r, 'Compared to the nearest %g mm. Two pieces that differ by '
                'less than that count as the same piece; the setting is in '
                'the Groups panel.' % data['tol_mm'])
    r = note(r, 'A mark ending "%s" is the MIRROR image of the one it is '
                'numbered with -- the same shape built the other way round. '
                'It is a different piece of steel and cannot be installed '
                'in its place.' % sm.MIRROR_SUFFIX)
    if data['issued']:
        r = note(r, 'The numbers are ISSUED: each part keeps the number it '
                    'went out under, whatever else changes around it. The '
                    'Mark Register sheet is what remembers that.')
        if data['withdrawn']:
            r = note(r, 'Issued and no longer built: %s. Those numbers stay '
                        'reserved -- the list has holes in it rather than '
                        'giving a retired number to a different part.'
                        % ', '.join(data['withdrawn']))
    else:
        r = note(r, 'These marks describe THIS model. They are not '
                    'comparable with a schedule exported before the model '
                    'changed -- issue them to hold the numbers still.')
    r += 1

    r = band(r, 'ASSEMBLIES — groups fabricated and shipped as one piece')
    if data['assemblies']:
        r = head(r, ASSEMBLY_HEADERS)
        for row in data['assemblies']:
            for col, v in enumerate((
                    row['mark'], row['qty'], row['n_rods'],
                    row['length_m'], row['total_length_m'],
                    row['mass_kg'], row['total_mass_kg'],
                    'yes' if row['mirrored'] else None,
                    ', '.join(row['names'])), 1):
                c = ws.cell(row=r, column=col, value=v)
                if isinstance(v, float):
                    c.number_format = '0.000'
            r += 1
    else:
        r = note(r, 'No groups with rods of their own, so there is nothing '
                    'to fabricate as an assembly yet.')
    r += 1

    r = band(r, 'PARTS — single rods, the list steel is ordered from')
    r = head(r, PART_HEADERS)
    for row in data['parts']:
        for col, v in enumerate((
                row['mark'], row['qty'], row['profile'], row['conn'],
                row['length_m'], row['total_length_m'],
                row['mass_kg'], row['total_mass_kg'],
                ', '.join(row['roles'])), 1):
            c = ws.cell(row=r, column=col, value=v)
            if isinstance(v, float):
                c.number_format = '0.000'
        r += 1

    r += 1
    ws.cell(row=r, column=1, value='TOTAL').font = Font(bold=True)
    ws.cell(row=r, column=2,
            value=sum(x['qty'] for x in data['parts'])).font = Font(bold=True)
    for col, total in ((6, sum(x['total_length_m'] for x in data['parts'])),
                       (8, sum(x['total_mass_kg'] for x in data['parts']))):
        c = ws.cell(row=r, column=col, value=total)
        c.font = Font(bold=True)
        c.number_format = '0.000'
    return ws, data


# ── the register: which number meant what when the drawings went out ──────

REGISTER_SHEET = 'Mark Register'
REGISTER_HEADERS = ('level', 'number', 'signature', 'info: what it was')


def write_register_sheet(wb, register):
    """Add the Mark Register sheet. Nothing to write means no sheet, so a
    model whose numbers were never issued does not grow one."""
    from openpyxl.styles import Font, PatternFill, Alignment
    if not register:
        return None
    rows = []
    for prefix in (sm.ASSEMBLY_PREFIX, sm.PART_PREFIX):
        for sig, n in sorted((register.get(prefix) or {}).items(),
                             key=lambda kv: kv[1]):
            rows.append((prefix, n, sig,
                         (register.get('was') or {}).get(sig, '')))
    if not rows:
        return None

    ws = wb.create_sheet(REGISTER_SHEET)
    ws.cell(row=1, column=1,
            value='MARK REGISTER — which number meant which part when the '
                  'drawings went out').font = Font(bold=True, size=12,
                                                   color='1F4E79')
    notes = [
        'This is what keeps a piece mark the same between revisions. A '
        'part listed here keeps its number however the model changes '
        'around it.',
        'A number here is reserved for good, even once the part it named '
        'has left the model: two different parts called %s1 in two '
        'revisions is the failure this exists to prevent, so the numbering '
        'has holes in it rather than reusing one.' % sm.ASSEMBLY_PREFIX,
        'signature: what the part IS -- its shape, section and the '
        'tolerance it was compared under. It is matched, never read; '
        'editing it only loses the number it was keeping.',
        'Compared to the nearest %g mm. The same register under a '
        'different tolerance matches nothing, and everything would be '
        'numbered again from scratch.'
        % sm.clamp_tol(register.get('tol_mm')),
        'Delete this sheet to let the numbers be worked out freely again.',
    ]
    for k, text in enumerate(notes):
        ws.cell(row=2 + k, column=1, value=text).font = Font(
            italic=True, size=9, color='5F6368')
    hdr_row = 2 + len(notes) + 1
    fill = PatternFill('solid', fgColor='404040')
    info = PatternFill('solid', fgColor='8C8C8C')
    for col, h in enumerate(REGISTER_HEADERS, 1):
        c = ws.cell(row=hdr_row, column=col, value=h)
        c.font = Font(bold=True, color='FFFFFF', size=10)
        c.fill = info if h.startswith('info:') else fill
        c.alignment = Alignment(horizontal='center')
    r = hdr_row + 1
    for level, n, sig, was in rows:
        for col, v in enumerate((level, n, sig, was), 1):
            ws.cell(row=r, column=col, value=v)
        r += 1
    ws.cell(row=hdr_row, column=5, value='tol_mm').font = Font(
        bold=True, size=9, color='5F6368')
    ws.cell(row=hdr_row + 1, column=5,
            value=sm.clamp_tol(register.get('tol_mm')))
    return ws


def read_register_sheet(wb):
    """The register a workbook carries, or None if it carries none."""
    if REGISTER_SHEET not in wb.sheetnames:
        return None
    rows = list(wb[REGISTER_SHEET].iter_rows(values_only=True))
    hdr_i = next((i for i, r in enumerate(rows)
                  if r and str(r[0]).strip() == 'level'), None)
    if hdr_i is None:
        raise ValueError('The Mark Register sheet has no header row (a row '
                         'starting with "level").')
    out = {sm.PART_PREFIX: {}, sm.ASSEMBLY_PREFIX: {}, 'was': {},
           'tol_mm': None}
    head = rows[hdr_i]
    if len(head) > 4 and str(head[4]).strip() == 'tol_mm' \
            and len(rows) > hdr_i + 1 and len(rows[hdr_i + 1]) > 4:
        try:
            out['tol_mm'] = float(rows[hdr_i + 1][4])
        except (TypeError, ValueError):
            out['tol_mm'] = None
    for r in rows[hdr_i + 1:]:
        if not r or all(v in (None, '') for v in r[:4]):
            continue
        level = str(r[0] or '').strip()
        if level not in (sm.PART_PREFIX, sm.ASSEMBLY_PREFIX):
            continue
        try:
            n = int(float(r[1]))
        except (TypeError, ValueError):
            continue
        sig = str(r[2] or '').strip()
        if not sig:
            continue
        out[level][sig] = n
        if len(r) > 3 and r[3]:
            out['was'][sig] = str(r[3])
    if not out[sm.PART_PREFIX] and not out[sm.ASSEMBLY_PREFIX]:
        return None
    return out
