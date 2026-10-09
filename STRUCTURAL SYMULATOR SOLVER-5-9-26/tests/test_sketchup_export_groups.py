"""The SketchUp extension's export, run for real against a stub SketchUp.

tests/test_sketchup_plugin_compat.py checks that the xlsx LAYOUT the plugin
writes is one the Stereo tab can read, by rebuilding that layout in Python.
This file is the other half: it runs the plugin's own Ruby -- the shipped
files under sketchup_plugin/, not a copy of them -- against a stand-in for
the SketchUp API (tests/ruby/stub), and then reads the workbook it really
wrote with the app's real importer.

WHAT IT IS GUARDING. Both export paths used to look only at the edges lying
loose in the current context (`entities.grep(Sketchup::Edge)`), so geometry
inside a group or a component was invisible to them. Measured against the
shipped code, that meant:

  * a grouped truss selected on its own  -> "No edges found in the selection"
  * a group selected WITH a few loose edges -> "Exported 3 nodes and 3
    members", having dropped the group's three and said nothing
  * picking while inside a group -> "Exported 3 nodes and 0 members", a
    model with no connectivity at all

The last two are the dangerous ones, because the export looked like it
worked. Every scenario below would fail on that code.
"""
import json
import os
import shutil
import subprocess
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import pytest

from common import _ensure_openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HARNESS = os.path.join(ROOT, 'tests', 'ruby', 'export_harness.rb')
STUB = os.path.join(ROOT, 'tests', 'ruby', 'stub')
RUBY = shutil.which('ruby')

pytestmark = [
    pytest.mark.skipif(not _ensure_openpyxl(), reason='openpyxl unavailable'),
    pytest.mark.skipif(RUBY is None,
                       reason='no ruby to run the plugin sources with'),
]

IN_TO_M = 0.0254


def _run(how, scenario, tmp_path):
    """Run the plugin's export and return (diagnostics, workbook path)."""
    out = str(tmp_path / ('%s_%s.xlsx' % (how, scenario)))
    proc = subprocess.run([RUBY, '-I', STUB, HARNESS, how, scenario, out],
                          cwd=ROOT, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
    return json.loads(proc.stdout.strip().splitlines()[-1]), out


def _read(path):
    """(nodes, members, groups) through the app's own importers."""
    from apps.stereo import stereo_reports as sr
    from apps.stereo import stereo_groups_excel as sge
    nodes, members, _loads, _supports, profiles = sr.import_excel_model(path)
    groups, _report = sge.import_groups(path, members, profiles)
    return nodes, members, groups


def _by_name(groups):
    out = {}
    for g in groups or []:
        out.setdefault(g['name'], []).append(g)
    return out


# ── the geometry survives, container or not ──────────────────────────────

@pytest.mark.parametrize('how', ['select', 'pick'])
@pytest.mark.parametrize('scenario,n_nodes,n_members', [
    # Nothing grouped: the case that already worked, so following
    # containers must not have disturbed it.
    ('flat', 3, 3),
    # The whole truss in one group, 1000in from the origin. Used to export
    # no members at all.
    ('grouped', 3, 3),
    # A group inside a group, each with its own offset.
    ('nested', 3, 3),
    # Two instances of one definition: the same local geometry, twice.
    ('components', 6, 6),
    # An instance turned a quarter circle, not merely moved.
    ('rotated', 3, 3),
    # A group selected alongside loose edges: 6, not the 3 it used to be.
    ('mixed', 6, 6),
    # A face in the selection contributes its boundary edges.
    ('face', 3, 3),
    # Editing inside a group, so the context transform is not the identity.
    ('inside_group', 3, 3),
])
def test_every_rod_comes_across(how, scenario, n_nodes, n_members, tmp_path):
    _diag, path = _run(how, scenario, tmp_path)
    nodes, members, _groups = _read(path)
    assert len(nodes) == n_nodes
    assert len(members) == n_members
    # Every member joins two nodes that exist and are not the same node --
    # the shape of the failure a mis-applied transform produces.
    for m in members:
        assert m['a'] != m['b']
        assert 0 <= m['a'] < len(nodes)
        assert 0 <= m['b'] < len(nodes)


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_a_grouped_truss_keeps_its_size(how, tmp_path):
    """The group's own 1000in offset is not in the exported coordinates --
    the origin is the model's own corner -- but its 100in legs are."""
    _diag, path = _run(how, 'grouped', tmp_path)
    nodes, _members, _groups = _read(path)
    xs = sorted(n[0] for n in nodes)
    ys = sorted(n[1] for n in nodes)
    assert xs[0] == pytest.approx(0.0)
    assert xs[-1] == pytest.approx(100 * IN_TO_M)
    assert ys[-1] == pytest.approx(100 * IN_TO_M)


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_a_group_moved_away_from_loose_edges_stays_away(how, tmp_path):
    """`mixed` is a loose triangle at the origin and a grouped one 500in
    along y. Lose the group's transformation and the two coincide, the node
    merge folds them into one, and the export quietly halves the model."""
    _diag, path = _run(how, 'mixed', tmp_path)
    nodes, _members, _groups = _read(path)
    ys = sorted(n[1] for n in nodes)
    assert ys[0] == pytest.approx(0.0)
    assert ys[-1] == pytest.approx(600 * IN_TO_M)


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_a_turned_instance_is_turned(how, tmp_path):
    """A quarter turn about z plus 500in along x sends (x, y, z) to
    (500 - y, x, z). A walk that carried only translations would put the
    corners somewhere else, and the picked points would miss the edges."""
    diag, path = _run(how, 'rotated', tmp_path)
    assert {tuple(p) for seg in diag['world'] for p in seg} == {
        (500.0, 0.0, 0.0), (500.0, 100.0, 0.0), (400.0, 0.0, 0.0)}
    nodes, members, _groups = _read(path)
    assert len(members) == 3
    assert len(nodes) == 3


# ── and so do the groups themselves ──────────────────────────────────────

@pytest.mark.parametrize('how', ['select', 'pick'])
def test_a_model_with_no_containers_writes_no_groups_sheet(how, tmp_path):
    """No Groups sheet at all, not an empty one: the Stereo tab tells "this
    file says there are no groups" from "this file does not mention
    groups", and an import must not drop the groups already in the model
    for a workbook that never had any."""
    _diag, path = _run(how, 'flat', tmp_path)
    _nodes, _members, groups = _read(path)
    assert groups is None


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_one_group_holds_every_rod(how, tmp_path):
    _diag, path = _run(how, 'grouped', tmp_path)
    _nodes, members, groups = _read(path)
    assert len(groups) == 1
    g = groups[0]
    assert g['name'] == 'Roof'
    assert g['parent'] is None
    assert g['members'] == set(range(len(members)))


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_nesting_survives_the_trip(how, tmp_path):
    """Roof holds Bay, Bay holds the rods. The outer group keeps a row of
    its own even though no rod belongs to it directly -- it is a real
    branch of the tree, and flattening it would lose the nesting."""
    _diag, path = _run(how, 'nested', tmp_path)
    _nodes, members, groups = _read(path)
    named = _by_name(groups)
    assert sorted(named) == ['Bay', 'Roof']
    roof, bay = named['Roof'][0], named['Bay'][0]
    assert roof['parent'] is None
    assert bay['parent'] == roof['id']
    assert bay['members'] == set(range(len(members)))
    assert roof['members'] == set()


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_two_instances_of_one_component_are_two_groups(how, tmp_path):
    """They share a definition, so they share a name; they are still two
    separate things in the model and two separate rows, because a rod
    belongs to exactly one group and these are different rods."""
    _diag, path = _run(how, 'components', tmp_path)
    _nodes, members, groups = _read(path)
    assert len(groups) == 2
    assert [g['name'] for g in groups] == ['Truss A', 'Truss A']
    assert all(g['parent'] is None for g in groups)
    assert [len(g['members']) for g in groups] == [3, 3]
    claimed = set()
    for g in groups:
        assert not (g['members'] & claimed), 'a rod in two groups'
        claimed |= g['members']
    assert claimed == set(range(len(members)))


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_loose_rods_stay_ungrouped(how, tmp_path):
    """`mixed` has three rods in a group and three lying loose. The loose
    ones belong to no group -- which is the Stereo tab's implicit Ungrouped
    set, not a group named for the context they came from."""
    _diag, path = _run(how, 'mixed', tmp_path)
    _nodes, members, groups = _read(path)
    assert len(groups) == 1
    assert groups[0]['name'] == 'Roof'
    assert len(groups[0]['members']) == 3
    assert len(members) == 6


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_the_instances_own_name_wins_over_its_definitions(how, tmp_path):
    """`rotated` is an instance called "Turned" of a definition called
    "Truss A". The outliner shows the instance's name, so that is the one
    the group carries."""
    _diag, path = _run(how, 'rotated', tmp_path)
    _nodes, _members, groups = _read(path)
    assert [g['name'] for g in groups] == ['Turned']


def test_an_empty_group_is_refused_not_exported(tmp_path):
    """Nothing to export is said out loud, and no file is written."""
    diag, path = _run('select', 'empty_group', tmp_path)
    assert diag['segments'] == 0
    assert not diag['wrote']
    assert not os.path.exists(path)
    assert 'No edges found' in diag['messages'][-1]


# ── the sheet itself ─────────────────────────────────────────────────────

def test_the_groups_sheet_sets_no_section_values(tmp_path):
    """Only the columns that say WHERE a rod belongs are written -- which
    part it is a copy of included. SketchUp knows nothing about steel, and
    a column of plausible defaults would overwrite the real sections of a
    model being re-imported."""
    import openpyxl
    from apps.stereo import stereo_groups_excel as sge
    _diag, path = _run('select', 'nested', tmp_path)
    wb = openpyxl.load_workbook(path, data_only=True)
    assert 'Groups' in wb.sheetnames
    rows = sge.read_groups_sheet(wb)
    assert sorted(rows[0]) == ['component', 'id', 'name', 'parent', 'rods']
    for field, _key, _kind in sge.FIELDS:
        assert field not in rows[0]


def test_importing_the_groups_sheet_changes_no_rod(tmp_path):
    """The report is what the sheet DID to the model's rods. A sheet that
    only says where things belong must do nothing to them."""
    from apps.stereo import stereo_reports as sr
    from apps.stereo import stereo_groups_excel as sge
    _diag, path = _run('select', 'nested', tmp_path)
    _nodes, members, _loads, _sup, profiles = sr.import_excel_model(path)
    before = [dict(m) for m in members]
    _groups, report = sge.import_groups(path, members, profiles)
    assert report == []
    assert members == before


def test_the_rod_lists_are_written_as_ranges(tmp_path):
    """'0-5' rather than '0, 1, 2, 3, 4, 5' -- the spelling the Stereo tab
    writes itself, so a round trip through Excel looks the same in both
    directions."""
    import openpyxl
    _diag, path = _run('select', 'grouped', tmp_path)
    wb = openpyxl.load_workbook(path, data_only=True)
    rods = [r[3] for r in wb['Groups'].iter_rows(values_only=True)
            if r and r[0] == 1]
    assert rods == ['0-2']


# ── components arrive as shared parts ────────────────────────────────────

def _parts(groups):
    """{component name: [group names]} for the groups that are copies."""
    from apps.stereo import stereo_components as scp
    out = {}
    for g in groups or ():
        name = scp.name_of(g)
        if name:
            out.setdefault(name, []).append(g['name'])
    return out


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_three_instances_of_one_component_are_one_part(how, tmp_path):
    """The case the whole thing is for. SketchUp's component is a shared
    definition -- one drawing, placed three times -- and it arrives as the
    Stereo tab's component, so the three get one section and one mark
    rather than being sized apart."""
    _diag, path = _run(how, 'three_copies', tmp_path)
    _nodes, members, groups = _read(path)
    assert len(members) == 9
    parts = _parts(groups)
    assert list(parts) == ['Truss A']
    assert len(parts['Truss A']) == 3


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_a_mirrored_instance_is_still_the_same_part(how, tmp_path):
    """A left-hand copy is the same drawing built the other way round.
    The Stereo tab marks the handedness itself (T1 and T1/m), so splitting
    it here would only hide that they are one part."""
    _diag, path = _run(how, 'mirrored_copies', tmp_path)
    _nodes, _members, groups = _read(path)
    parts = _parts(groups)
    assert list(parts) == ['Truss A']
    assert len(parts['Truss A']) == 2


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_the_same_definition_at_two_sizes_is_two_parts(how, tmp_path):
    """Same drawing on screen, different steel: scaling an instance
    changes every bar length in it. Calling them one part would have the
    Stereo tab size the small one from the big one's loads."""
    _diag, path = _run(how, 'scaled_copies', tmp_path)
    _nodes, _members, groups = _read(path)
    parts = _parts(groups)
    assert len(parts) == 2, parts
    assert all(len(v) == 1 for v in parts.values())
    assert all('Truss A' in k for k in parts), parts
    assert any('0.5000' in k for k in parts), 'and the size is named'


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_two_copies_of_a_group_are_not_a_shared_part(how, tmp_path):
    """A SketchUp group is unique -- editing one does not touch another --
    so two of them are two groups. Making them one part is a decision, and
    the Stereo tab's own Make component is where it gets made."""
    _diag, path = _run(how, 'copied_groups', tmp_path)
    _nodes, _members, groups = _read(path)
    assert len(groups) == 2
    assert _parts(groups) == {}


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_a_group_inside_a_definition_is_shared_with_it(how, tmp_path):
    """A definition's contents are one set of entities, shared by every
    instance rather than copied -- so the bay inside copy 1 and the bay
    inside copy 2 are the same steel, and arrive saying so."""
    _diag, path = _run(how, 'nested_in_component', tmp_path)
    _nodes, _members, groups = _read(path)
    parts = _parts(groups)
    assert list(parts) == ['Truss A / Bay']
    assert len(parts['Truss A / Bay']) == 2


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_an_instance_holding_only_subgroups_is_a_branch_not_a_part(how,
                                                                    tmp_path):
    """The Stereo tab's rule is that a component is a group's OWN rods; it
    refuses to make a part of a branch. Writing one here would only be
    refused later, having inflated the copy count in between."""
    _diag, path = _run(how, 'nested_in_component', tmp_path)
    _nodes, _members, groups = _read(path)
    from apps.stereo import stereo_components as scp
    for g in groups:
        if not g['members']:
            assert scp.name_of(g) is None, g['name']


@pytest.mark.parametrize('how', ['select', 'pick'])
def test_two_sub_groups_with_one_name_are_still_two_parts(how, tmp_path):
    """They would both read as "Truss A / Bay". They are different steel,
    so the second is numbered -- a false merge here would have one part
    sized from another part's loads."""
    _diag, path = _run(how, 'twin_bays', tmp_path)
    _nodes, _members, groups = _read(path)
    parts = _parts(groups)
    assert len(parts) == 2, parts
    assert all(len(v) == 1 for v in parts.values())


@pytest.mark.parametrize('scenario', ['flat', 'grouped', 'nested', 'mixed'])
def test_a_model_with_no_components_names_no_parts(scenario, tmp_path):
    """Nothing in these is a component instance, so nothing claims to be a
    copy of anything."""
    _diag, path = _run('select', scenario, tmp_path)
    _nodes, _members, groups = _read(path)
    assert _parts(groups) == {}


def test_the_sheet_carries_the_component_column(tmp_path):
    import openpyxl
    from apps.stereo import stereo_groups_excel as sge
    _diag, path = _run('select', 'three_copies', tmp_path)
    rows = sge.read_groups_sheet(openpyxl.load_workbook(path, data_only=True))
    assert sorted(rows[0]) == ['component', 'id', 'name', 'parent', 'rods']
    assert {r['component'] for r in rows} == {'Truss A'}


def test_the_export_says_how_many_shared_parts_it_found(tmp_path):
    diag, _path = _run('select', 'three_copies', tmp_path)
    assert '1 of them shared part(s)' in diag['messages'][-1]


def test_a_model_with_no_components_does_not_mention_parts(tmp_path):
    diag, _path = _run('select', 'grouped', tmp_path)
    assert 'shared part' not in diag['messages'][-1]
