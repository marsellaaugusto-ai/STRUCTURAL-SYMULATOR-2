"""What the Beam tab TELLS the user about its own answer.

The solver-level half of this lives in `test_beam_math.py` (R-1, R-2). This
file is the reporting half, driven through the real tab and the real workbook,
because a residual the user never sees is not a check:

* an equilibrium block in the RESULTS panel, so a wrong answer is visible
  without re-deriving it;
* the same rows in the exported `Results` sheet, where a reviewer looks;
* a note when two supports at one station were merged -- silently analysing a
  structure the user did not describe is the bug this came from (R-1), and
  merging is only the right answer if it is reported.
"""
import os
import sys
import tempfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

tk = pytest.importorskip('tkinter')

import units
from apps.beam.beam_app import BeamApp


@pytest.fixture
def beam():
    units.set_current('cirsoc')
    try:
        root = tk.Tk()
    except tk.TclError:                       # pragma: no cover - headless CI
        pytest.skip('no display')
    root.geometry('1400x900')
    app = BeamApp(root)
    app.pack(fill='both', expand=True)
    root.update()
    try:
        yield app, root
    finally:
        units.set_current('cirsoc')
        try:
            root.destroy()
        except Exception:
            pass


def _ss_udl(app):
    app.length = 6.0
    app.len_var.set(6.0)
    app.supports = [{'x': 0.0, 'type': 'pin'}, {'x': 6.0, 'type': 'roller'}]
    app.dloads = [{'x1': 0.0, 'x2': 6.0, 'w1': 10.0, 'w2': 10.0}]
    app._refresh_tables()


def _results(app, root):
    app._analyze()
    root.update()
    return app.res_text.get('1.0', 'end')


def test_the_results_panel_reports_equilibrium(beam):
    app, root = beam
    _ss_udl(app)
    text = _results(app, root)
    assert 'EQUILIBRIUM' in text
    assert 'ok' in text.lower()


def test_the_panel_states_the_total_load_the_reactions_balance(beam):
    """A reader checking a 10 kN/m x 6 m span wants to see 60 kN of load against
    60 kN of reaction, in the convention they selected -- not a bare residual."""
    app, root = beam
    _ss_udl(app)
    text = _results(app, root)
    assert '60.00' in text, text


def test_the_panel_reports_the_degree_of_indeterminacy(beam):
    app, root = beam
    _ss_udl(app)
    app.supports = [{'x': 0.0, 'type': 'fixed'}, {'x': 6.0, 'type': 'fixed'}]
    app._refresh_tables()
    text = _results(app, root)
    assert 'indeterminate' in text.lower()
    assert 'degree 2' in text.lower()


def test_a_determinate_beam_is_called_determinate(beam):
    app, root = beam
    _ss_udl(app)
    assert 'determinate' in _results(app, root).lower()


def test_the_panel_warns_when_two_supports_were_merged(beam):
    """R-1. The tab's own list still holds both rows -- the user typed them --
    so the one place this can be said is the results panel."""
    app, root = beam
    _ss_udl(app)
    app.supports = [{'x': 0.0, 'type': 'pin'}, {'x': 0.0, 'type': 'guided'},
                    {'x': 6.0, 'type': 'roller'}]
    app._refresh_tables()
    text = _results(app, root)
    assert 'NOTES' in text
    assert 'merged' in text.lower()
    assert 'fixed' in text.lower()


def test_merged_supports_report_one_reaction_each_station(beam):
    """The headline symptom: three support rows at two stations reported three
    reactions totalling 90 kN for 60 kN of load."""
    app, root = beam
    _ss_udl(app)
    app.supports = [{'x': 0.0, 'type': 'pin'}, {'x': 0.0, 'type': 'roller'},
                    {'x': 6.0, 'type': 'roller'}]
    app._refresh_tables()
    text = _results(app, root)
    reaction_lines = [ln for ln in text.splitlines() if 'Ry=' in ln]
    assert len(reaction_lines) == 2, text
    assert '90.00' not in text, text


def test_merging_never_rewrites_the_user_s_own_table(beam):
    """The merge belongs to the analysed model, not to the tab's state.

    If it edited `app.supports`, pressing Analyze would silently retype the
    user's rows -- and `_current_state()` feeds the Excel export, so a saved
    workbook would stop being the model that was entered. Re-analysing must
    also not stack up notes.
    """
    app, root = beam
    _ss_udl(app)
    typed = [{'x': 0.0, 'type': 'pin'}, {'x': 0.0, 'type': 'guided'},
             {'x': 6.0, 'type': 'roller'}]
    app.supports = [dict(s) for s in typed]
    app._refresh_tables()
    for _ in range(3):
        text = _results(app, root)
        assert app.supports == typed
        assert text.lower().count('were merged') == 1, text
        assert [s['type'] for s in app.model.supports] == ['fixed', 'roller']


def test_a_clean_analysis_has_no_notes_block(beam):
    app, root = beam
    _ss_udl(app)
    assert 'NOTES' not in _results(app, root)


@pytest.mark.parametrize('key', ['aisc', 'eurocode'])
def test_the_equilibrium_block_follows_the_unit_selector(beam, key):
    app, root = beam
    _ss_udl(app)
    app._analyze()
    units.set_current(key)
    root.update()
    text = app.res_text.get('1.0', 'end')
    assert 'EQUILIBRIUM' in text
    assert units.label('force') in text


# ------------------------------------------------------------------- workbook

def test_the_excel_results_sheet_carries_the_equilibrium_rows(beam):
    pytest.importorskip('openpyxl')
    import openpyxl
    from apps.beam.beam_app import export_beam_excel

    app, root = beam
    _ss_udl(app)
    app._analyze()
    root.update()
    assert app.result is not None

    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, 'beam.xlsx')
        export_beam_excel(app._current_state(), path,
                          result=app.result, model=app.model)
        wb = openpyxl.load_workbook(path, data_only=True)
        assert 'Results' in wb.sheetnames
        labels = [v for row in wb['Results'].iter_rows(values_only=True)
                  for v in row if isinstance(v, str)]
        joined = ' | '.join(labels)
        wb.close()
    for want in ('Total load', 'Sum of reactions', 'Residual'):
        assert want in joined, joined
