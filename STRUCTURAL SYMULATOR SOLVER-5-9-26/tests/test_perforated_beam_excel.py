"""
test_perforated_beam_excel.py -- the Perforated Beam tab's Excel round-trip.

Two layers, because they fail in different ways:

  * The FORMAT layer (perforated_beam_excel) is pure and headless. Its
    risk is silent infidelity -- an opening that comes back as a slightly
    different polygon, a support that loses its fixity, a dimension that
    becomes 0 where the user had left it blank.
  * The WIRING layer (the app's _gather_state/_apply_state) is where a
    field gets forgotten. Its risk is that the file is perfect and the
    widget it belongs to never receives it.

The load-bearing test is `test_box_girder_survives_the_round_trip`, which
does not compare state dicts at all: it ANALYSES both the original and
the imported model and demands the same reactions, the same governing
opening and the same deflection. A state comparison would happily pass a
model whose openings had quietly become 47-gons; this one cannot.
"""
import sys
import time
sys.path.insert(0, str(__import__('pathlib').Path(__file__).resolve().parents[1]))

import pytest

openpyxl = pytest.importorskip('openpyxl')

from apps.perforated_beam import perforated_beam_math as pbm
from apps.perforated_beam import welded_section_math as wsm
from apps.perforated_beam import hyperstatic_math as hym
from apps.perforated_beam import section_profile_math as secm
from apps.perforated_beam import perforated_beam_excel as pbx


# ── helpers ──────────────────────────────────────────────────────────────

H, B, T, LT, LB = 1800.0, 250.0, 10.0, 100.0, 200.0
LIPPED_C = [(0., -H/2), (B, -H/2), (B, -H/2+LB), (B-T, -H/2+LB), (B-T, -H/2+T),
            (T, -H/2+T), (T, H/2-T), (B-T, H/2-T), (B-T, H/2-LT), (B, H/2-LT),
            (B, H/2), (0., H/2)]


def sketch_through(pts):
    """A closed drawn outline -- the same object the profile designer
    hands to Apply, so its JSON is the JSON the app would write."""
    out = secm.SectionOutline(tuple(pts[0]))
    for p in list(pts[1:]) + [pts[0]]:
        out.segments.append(secm.LineSegment(tuple(p)))
    return secm.SectionSketch(out)


def box_girder_state():
    """The 50.4 m double-profile box girder: compound section from two
    drawn profiles, 26 circular openings, three supports, ten loads."""
    L = 50400.0
    st = pbx.default_state()
    st.update({
        'length': L, 'Fy': 250.0, 'E': 200000.0,
        'section_mode': 'compound',
        'compound_offsets': {'cp_a': (125.0, 0.0), 'cp_b': (375.0, 0.0)},
        'assembly': {'kind': 'closed', 'plate_t': 10.0},
        'support_mode': 'advanced',
        'support_specs': [{'x': 1800.0, 'kind': hym.PIN, 'torsion': True},
                          {'x': 16200.0, 'kind': hym.PIN, 'torsion': True},
                          {'x': 48600.0, 'kind': hym.PIN, 'torsion': True}],
        'openings': pbm.uniform_layout(pbm.opening_circle(1350.0), 26, 1800.0, 2700.0),
        'loads': ([{'type': 'Point load', 'x1': i * 7200.0, 'v1': 100000.0, 'e': 0.0, 'case': 'L'}
                   for i in range(8)]
                  + [{'type': 'Distributed load', 'x1': 0.0, 'x2': L,
                      'v1': 5.0, 'v2': 5.0, 'e': 0.0, 'case': 'L'},
                     {'type': 'Distributed load', 'x1': 18000.0, 'x2': 46800.0,
                      'v1': 1.0, 'v2': 1.0, 'e': 0.0, 'case': 'L'}]),
    })
    st['bases']['cp_a'] = dict(pbx._default_base(), kind='custom_profile',
                               sketch=pbx.sketch_to_text(sketch_through(LIPPED_C)))
    st['bases']['cp_b'] = dict(
        pbx._default_base(), kind='custom_profile',
        sketch=pbx.sketch_to_text(
            sketch_through([(-x, y) for x, y in reversed(LIPPED_C)])))
    return st


def beam_from_state(st):
    """Assemble the beam a state dict describes, by the same route
    PerforatedBeamApp._build_beam takes."""
    parts = []
    for prefix in ('cp_a', 'cp_b'):
        sec, _ = pbx.sketch_to_section(st['bases'][prefix]['sketch'])
        dx, dy = st['compound_offsets'][prefix]
        parts.append(wsm.PlacedProfile(sec, dx=dx, dy=dy))
    section = wsm.CompoundSection(
        parts, detail=wsm.AssemblyDetail(st['assembly']['kind'],
                                         plate_thickness=st['assembly']['plate_t']))
    beam = pbm.BeamConfig(
        L=st['length'], section=section,
        material=pbm.Material(Fy=st['Fy'], E=st['E']),
        openings=list(st['openings']), supports=None,
        support_specs=[hym.SupportSpec(s['x'], s['kind'], s['torsion'])
                       for s in st['support_specs']],
        Lb=st['Lb'], Cb=st['Cb'], load_on_top_flange=st['load_on_top_flange'])
    for ld in st['loads']:
        if ld['type'] == 'Point load':
            beam.point_loads.append(pbm.PointLoad(ld['x1'], ld['v1'], ld.get('e', 0.0)))
        elif ld['type'] == 'Distributed load':
            beam.dist_loads.append(
                pbm.DistLoad(ld['x1'], ld['x2'], ld['v1'], ld['v2'], ld.get('e', 0.0)))
    return beam


def fingerprint(beam):
    rep = pbm.analyze_beam(beam)
    go = rep['governing_opening']
    _xs, v = pbm.deflection_profile(beam, n=201)
    return {
        'A': beam.section.A, 'I': beam.section.I, 'd': beam.section.d,
        'reactions': tuple(round(R, 9) for _x, R, _M in rep['support_reactions']),
        'gov_opening': (go['opening'].label, go['governing'].util_top,
                        go['governing'].util_bot) if go else None,
        'gov_post': rep['governing_webpost'].util if rep['governing_webpost'] else None,
        'n_openings': len(rep['openings']),
        'defl_max': max(v),
    }


# ── the format layer ─────────────────────────────────────────────────────

@pytest.mark.parametrize('verts,expect', [
    (pbm.opening_circle(1350.0), 'Circle'),
    (pbm.opening_circle(250.0, n=96), 'Circle'),
    (pbm.opening_rectangle(400.0, 250.0), 'Rectangle'),
    (pbm.opening_hexagon(500.0, 300.0), 'Hexagon'),
    (pbm.opening_hexagon(500.0, 300.0, 120.0), 'Hexagon'),
])
def test_standard_shapes_round_trip_exactly(verts, expect):
    """Recognition is by reconstruction, so a recognised shape must
    regenerate its polygon vertex for vertex -- not merely close."""
    shape, p1, p2, p3, n, poly = pbx.describe_shape(verts)
    assert shape == expect
    assert pbx.build_shape(shape, p1, p2, p3, n, poly) == verts


def test_unrecognised_shape_falls_back_to_explicit_vertices():
    """A polygon no constructor produces must survive anyway. This one is
    deliberately irregular: were recognition to guess 'Rectangle' from
    its bounding box, the reconstructed opening would be a different hole
    in a different place, and nothing downstream would notice."""
    odd = [(-100.0, -50.0), (120.0, -40.0), (90.0, 60.0), (-80.0, 45.0)]
    shape, p1, p2, p3, n, poly = pbx.describe_shape(odd)
    assert shape == 'Custom polygon'
    assert pbx.build_shape(shape, p1, p2, p3, n, poly) == odd


def test_a_near_miss_is_not_recognised_as_a_circle():
    """One vertex moved by a hundredth of a millimetre is no longer the
    polygon `opening_circle` produces, and must not be stored as though
    it were -- otherwise import would silently return the true circle."""
    verts = pbm.opening_circle(1350.0)
    verts[7] = (verts[7][0] + 0.01, verts[7][1])
    assert pbx.describe_shape(verts)[0] == 'Custom polygon'
    shape, p1, p2, p3, n, poly = pbx.describe_shape(verts)
    rebuilt = pbx.build_shape(shape, p1, p2, p3, n, poly)
    assert rebuilt[7] == pytest.approx(verts[7])


def test_blank_dimensions_stay_blank(tmp_path):
    """An empty custom-dimension entry means 'use the catalog'. Writing it
    back as 0.0 would turn 'not given' into a real -- and invalid --
    dimension, which RolledSection would then reject at Analyze."""
    st = pbx.default_state()
    st['custom_ibeam'] = {'d': '', 'bf': '', 'tf': '', 'tw': '',
                          'rotate90': False, 'mirror': False}
    p = tmp_path / 'blank.xlsx'
    pbx.export_state(st, str(p))
    back = pbx.import_state(str(p))
    assert [back['custom_ibeam'][k] for k in pbx.IBEAM_KEYS] == ['', '', '', '']


def test_dimensions_round_trip_as_typed(tmp_path):
    st = pbx.default_state()
    st['custom_ibeam'] = {'d': '600', 'bf': '200', 'tf': '15', 'tw': '9',
                          'rotate90': True, 'mirror': False}
    p = tmp_path / 'dims.xlsx'
    pbx.export_state(st, str(p))
    back = pbx.import_state(str(p))
    assert [back['custom_ibeam'][k] for k in pbx.IBEAM_KEYS] == ['600', '200', '15', '9']
    assert back['custom_ibeam']['rotate90'] is True
    assert back['custom_ibeam']['mirror'] is False


def test_positions_are_metres_and_section_dims_are_millimetres(tmp_path):
    """The unit split the tab settled on: metres along the span, mm across
    the section. Read the cells directly -- a round-trip alone would pass
    even if both ends used the wrong unit consistently."""
    st = pbx.default_state()
    st['length'] = 50400.0
    st['support_mode'] = 'advanced'
    st['support_specs'] = [{'x': 16200.0, 'kind': hym.PIN, 'torsion': True}]
    st['openings'] = pbm.uniform_layout(pbm.opening_circle(1350.0), 1, 1800.0, 2700.0)
    st['assembly'] = {'kind': 'closed', 'plate_t': 10.0}
    p = tmp_path / 'units.xlsx'
    pbx.export_state(st, str(p))

    rows = list(openpyxl.load_workbook(str(p))['Model'].iter_rows(values_only=True))
    rd = pbx._Reader(rows)
    assert rd.one('GEOMETRY')['length_m'] == pytest.approx(50.4)
    assert rd.block('SUPPORTS')[0]['x_m'] == pytest.approx(16.2)
    assert rd.block('OPENINGS')[0]['x_center_m'] == pytest.approx(2.7)
    assert rd.block('OPENINGS')[0]['p1_mm'] == pytest.approx(1350.0)
    assert rd.one('ASSEMBLY')['plate_t_mm'] == pytest.approx(10.0)


def test_supports_and_loads_round_trip(tmp_path):
    st = pbx.default_state()
    st['support_mode'] = 'advanced'
    st['support_specs'] = [
        {'x': 0.0, 'kind': hym.FIXED, 'torsion': True},
        {'x': 8000.0, 'kind': hym.PIN, 'torsion': False}]
    st['loads'] = [
        {'type': 'Point load', 'x1': 2000.0, 'v1': 5000.0, 'e': 40.0, 'case': 'L'},
        {'type': 'Point moment', 'x1': 3000.0, 'v1': 1.2e6, 'case': 'L'},
        {'type': 'Point torque', 'x1': 4000.0, 'v1': 3.0e5, 'case': 'L'},
        {'type': 'Distributed load', 'x1': 0.0, 'x2': 8000.0,
         'v1': 15.0, 'v2': 25.0, 'e': 0.0, 'case': 'L'},
        {'type': 'Distributed torque', 'x1': 1000.0, 'x2': 7000.0, 'v1': 10.0, 'v2': 10.0, 'case': 'L'},
    ]
    p = tmp_path / 'sl.xlsx'
    pbx.export_state(st, str(p))
    back = pbx.import_state(str(p))
    assert back['support_specs'] == st['support_specs']
    assert back['loads'] == st['loads']
    # a fixed end is the whole reason the advanced path exists; losing its
    # kind would silently turn a built-in support into a pin
    assert back['support_specs'][0]['kind'] == hym.FIXED
    assert back['support_specs'][1]['torsion'] is False


def test_compound_welds_round_trip(tmp_path):
    st = pbx.default_state()
    st['compound_welds'] = [
        wsm.WeldLine((0.0, 900.0), (500.0, 900.0), leg=6.0, n_lines=2,
                     label='CW1', t_thicker=10.0, t_thinner=8.0, flange_to_web=True),
        wsm.WeldLine((0.0, -900.0), (500.0, -900.0), leg=0.0, n_lines=1, label='CW2'),
    ]
    p = tmp_path / 'welds.xlsx'
    pbx.export_state(st, str(p))
    back = pbx.import_state(str(p))
    assert len(back['compound_welds']) == 2
    a, b = back['compound_welds']
    assert (a.p1, a.p2, a.leg, a.n_lines, a.label) == ((0.0, 900.0), (500.0, 900.0), 6.0, 2, 'CW1')
    assert (a.t_thicker, a.t_thinner, a.flange_to_web) == (10.0, 8.0, True)
    assert (b.leg, b.n_lines, b.flange_to_web) == (0.0, 1, False)


def test_drawn_profile_survives_as_a_section(tmp_path):
    """The sketch is the expensive thing to re-enter, so it must come back
    as the same SECTION, not merely as the same JSON."""
    st = pbx.default_state()
    st['section_mode'] = 'custom_profile'
    st['main_sketch'] = pbx.sketch_to_text(sketch_through(LIPPED_C))
    p = tmp_path / 'sketch.xlsx'
    pbx.export_state(st, str(p))
    back = pbx.import_state(str(p))

    a, _ = pbx.sketch_to_section(st['main_sketch'])
    b, _ = pbx.sketch_to_section(back['main_sketch'])
    assert b.A == pytest.approx(a.A, rel=1e-12)
    assert b.I == pytest.approx(a.I, rel=1e-12)
    assert b.d == pytest.approx(a.d, rel=1e-12)


def test_a_missing_block_falls_back_to_its_default(tmp_path):
    """A hand-written sheet holding only a span and a load set is a valid
    input: blocks are found by name, so absence means 'default', not
    'corrupt'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Model'
    ws['A1'] = '[GEOMETRY]'
    ws['A2'], ws['B2'] = 'length_m', 'Fy_MPa'
    ws['A3'], ws['B3'] = 12.0, 345.0
    ws['A5'] = '[LOADS]'
    for col, label in enumerate(['type', 'x1_m', 'x2_m', 'v1'], 1):
        ws.cell(row=6, column=col, value=label)
    ws['A7'], ws['B7'], ws['C7'], ws['D7'] = 'Distributed load', 0.0, 12.0, 20.0
    p = tmp_path / 'sparse.xlsx'
    wb.save(str(p))

    st = pbx.import_state(str(p))
    assert st['length'] == pytest.approx(12000.0)
    assert st['Fy'] == pytest.approx(345.0)
    assert st['E'] == pytest.approx(200000.0)          # defaulted
    assert st['section_mode'] == 'catalog'             # defaulted
    assert st['loads'] == [{'type': 'Distributed load', 'x1': 0.0, 'x2': 12000.0,
                            'v1': 20.0, 'v2': 20.0, 'e': 0.0, 'case': 'L'}]


def test_reordered_columns_still_read(tmp_path):
    """Columns are located by header. Someone will move one."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Model'
    ws['A1'] = '[GEOMETRY]'
    ws['A2'], ws['B2'], ws['C2'] = 'Fy_MPa', 'length_m', 'E_MPa'
    ws['A3'], ws['B3'], ws['C3'] = 275.0, 9.0, 210000.0
    p = tmp_path / 'reordered.xlsx'
    wb.save(str(p))
    st = pbx.import_state(str(p))
    assert (st['length'], st['Fy'], st['E']) == (9000.0, 275.0, 210000.0)


def test_a_foreign_workbook_is_refused_by_name(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = 'Sheet1'
    p = tmp_path / 'foreign.xlsx'
    wb.save(str(p))
    with pytest.raises(ValueError, match='no "Model" sheet'):
        pbx.import_state(str(p))


@pytest.mark.parametrize('block,header,bad,match', [
    ('LOADS', ['type', 'x1_m', 'v1'], ['Wind', 0.0, 5.0], 'unknown type'),
    ('SUPPORTS', ['x_m', 'kind'], [1.8, 'springy'], 'kind must be'),
    ('OPENINGS', ['x_center_m', 'shape'], [2.7, 'Blob'], 'Unknown opening shape'),
    ('ASSEMBLY', ['kind'], ['welded-ish'], 'none, open or closed'),
])
def test_bad_values_raise_a_message_meant_for_a_dialog(tmp_path, block, header, bad, match):
    """Every failure the user can cause by typing must arrive as a
    ValueError they can act on, not a KeyError from three frames down."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Model'
    ws['A1'] = '[GEOMETRY]'
    ws['A2'] = 'length_m'
    ws['A3'] = 8.0
    ws['A5'] = f'[{block}]'
    for col, label in enumerate(header, 1):
        ws.cell(row=6, column=col, value=label)
    for col, value in enumerate(bad, 1):
        ws.cell(row=7, column=col, value=value)
    p = tmp_path / 'bad.xlsx'
    wb.save(str(p))
    with pytest.raises(ValueError, match=match):
        pbx.import_state(str(p))


def test_a_load_list_missing_x2_is_refused(tmp_path):
    """A distributed load with no end station has no meaning; accepting it
    with x2 defaulted to 0 would silently reverse it."""
    st = pbx.default_state()
    p = tmp_path / 'nox2.xlsx'
    pbx.export_state(st, str(p))
    wb = openpyxl.load_workbook(str(p))
    ws = wb['Model']
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows, 1)
               if r and isinstance(r[0], str) and r[0].strip() == '[LOADS]')
    ws.cell(row=hdr + 2, column=1, value='Distributed load')
    ws.cell(row=hdr + 2, column=2, value=0.0)
    ws.cell(row=hdr + 2, column=4, value=15.0)
    wb.save(str(p))
    with pytest.raises(ValueError, match='needs x2_m'):
        pbx.import_state(str(p))


def test_an_unreadable_sketch_is_reported_at_import(tmp_path):
    """Better to fail on the sheet that carries the damage than at Analyze,
    three steps later, with no clue which profile is at fault."""
    st = pbx.default_state()
    st['main_sketch'] = '{ this is not json'
    p = tmp_path / 'badsketch.xlsx'
    pbx.export_state(st, str(p))
    with pytest.raises(ValueError, match="drawing for 'main'"):
        pbx.import_state(str(p))


def test_support_mode_follows_the_support_list(tmp_path):
    """The list is what routes the beam through the stiffness solver, so a
    sheet listing three supports must not import as 'simple' -- that would
    quietly analyse a different structure than the one described."""
    st = pbx.default_state()
    st['support_mode'] = 'simple'
    st['support_specs'] = [{'x': 0.0, 'kind': hym.PIN, 'torsion': True},
                           {'x': 4000.0, 'kind': hym.PIN, 'torsion': True},
                           {'x': 8000.0, 'kind': hym.PIN, 'torsion': True}]
    p = tmp_path / 'mode.xlsx'
    pbx.export_state(st, str(p))
    assert pbx.import_state(str(p))['support_mode'] == 'advanced'


def test_results_sheet_is_written_only_with_a_report(tmp_path):
    st = box_girder_state()
    beam = beam_from_state(st)
    inputs_only = tmp_path / 'inputs.xlsx'
    pbx.export_state(st, str(inputs_only))
    assert openpyxl.load_workbook(str(inputs_only)).sheetnames == ['Model']

    both = tmp_path / 'both.xlsx'
    pbx.export_state(st, str(both), report=pbm.analyze_beam(beam), beam=beam)
    assert 'Results' in openpyxl.load_workbook(str(both)).sheetnames


def test_results_sheet_carries_the_member_check_for_the_box_girder(tmp_path):
    """`member` used to be None for anything that was not a rolled I --
    the compound box girder included -- and this test pinned the sheet's
    "not evaluated" notice. `section_plastic` recovers Zx and the web
    geometry from the polygons, so the box girder now gets a real Chapter
    F/G check and the sheet must carry its numbers instead."""
    st = box_girder_state()
    beam = beam_from_state(st)
    rep = pbm.analyze_beam(beam)
    assert rep['member'] is not None, 'the box girder should now get a member check'

    p = tmp_path / 'member.xlsx'
    pbx.export_state(st, str(p), report=rep, beam=beam)
    rows = list(openpyxl.load_workbook(str(p))['Results'].iter_rows(values_only=True))
    labels = {str(r[0]): r[1] for r in rows if r and isinstance(r[0], str)}
    assert 'util_flexure' in labels and 'util_shear' in labels
    assert isinstance(labels['util'], float) and labels['util'] > 0
    text = '\n'.join(str(c) for row in rows for c in row if isinstance(c, str))
    assert 'not evaluated' not in text


def test_results_sheet_still_says_so_when_there_is_no_member_check(tmp_path):
    """The notice is kept for the case that still degrades: a report whose
    member check really is None."""
    st = box_girder_state()
    beam = beam_from_state(st)
    rep = pbm.analyze_beam(beam)
    rep['member'] = None
    p = tmp_path / 'nomember.xlsx'
    pbx.export_state(st, str(p), report=rep, beam=beam)
    text = '\n'.join(
        str(c) for row in openpyxl.load_workbook(str(p))['Results'].iter_rows(values_only=True)
        for c in row if isinstance(c, str))
    assert 'not evaluated' in text


# ── the load-bearing test ────────────────────────────────────────────────

def test_box_girder_survives_the_round_trip(tmp_path):
    """Export the 50.4 m box girder, import it, and analyse BOTH.

    Reactions, governing opening, governing web post and peak deflection
    must agree to floating-point noise. This is the test that would catch
    a 1350 mm circle coming back as 1349.98, which no state comparison
    would flag as wrong and which the drawings would render as a perfectly
    ordinary hole."""
    st0 = box_girder_state()
    p = tmp_path / 'box.xlsx'
    pbx.export_state(st0, str(p))
    st1 = pbx.import_state(str(p))

    assert len(st1['openings']) == 26
    for a, b in zip(st0['openings'], st1['openings']):
        assert a.label == b.label
        assert a.x_center == b.x_center
        assert a.vertices_local == b.vertices_local     # exact, not approx

    f0 = fingerprint(beam_from_state(st0))
    f1 = fingerprint(beam_from_state(st1))
    assert f1['reactions'] == f0['reactions']
    assert f1['gov_opening'] == f0['gov_opening']
    assert f1['gov_post'] == pytest.approx(f0['gov_post'], rel=1e-12)
    assert f1['defl_max'] == pytest.approx(f0['defl_max'], rel=1e-12)
    assert f1['A'] == pytest.approx(f0['A'], rel=1e-12)
    assert f1['I'] == pytest.approx(f0['I'], rel=1e-12)

    # and the absolute values, so a change that moves BOTH sides together
    # still shows up here rather than passing as "consistent"
    assert f1['reactions'][0] == pytest.approx(102974.68, abs=0.01)
    assert f1['reactions'][1] == pytest.approx(625436.58, abs=0.01)
    assert f1['reactions'][2] == pytest.approx(352388.75, abs=0.01)
    assert f1['defl_max'] == pytest.approx(39.97, abs=0.01)


# ── the wiring layer ─────────────────────────────────────────────────────

tk = pytest.importorskip('tkinter')
from apps.perforated_beam.perforated_beam_app import PerforatedBeamApp   # noqa: E402


@pytest.fixture(scope='session')
def tk_root():
    """Same retrying root as test_perforated_beam_app_features -- a bare
    skip on TclError turns a transient failure into a silently skipped
    test, which reads exactly like a passing one."""
    last = None
    for attempt in range(6):
        try:
            root = tk.Tk()
            break
        except tk.TclError as exc:
            last = exc
            time.sleep(0.5 * (attempt + 1))
    else:
        pytest.skip(f'no Tk display after 6 attempts: {last}')
    root.withdraw()
    yield root
    try:
        root.destroy()
    except tk.TclError:
        pass


@pytest.fixture
def app(tk_root):
    widget = PerforatedBeamApp(tk_root)
    yield widget
    widget.destroy()


def test_every_state_key_reaches_the_widgets(app, tmp_path):
    """Load a fully populated model into the app, export it, clear the tab
    to its defaults, import it back, and require the app's OWN state to
    match. This is what catches a field that the format handles correctly
    and `_apply_state` forgets to set."""
    app._apply_state(box_girder_state())
    p = tmp_path / 'wiring.xlsx'
    pbx.export_state(app._gather_state(), str(p))
    before = app._gather_state()

    app._apply_state(pbx.default_state())
    assert app.length == pytest.approx(8000.0)          # really was reset
    assert app.openings == []

    app._apply_state(pbx.import_state(str(p)))
    after = app._gather_state()

    for key in before:
        if key in ('openings', 'compound_welds'):
            continue                                     # objects, compared below
        assert after[key] == before[key], f'{key} did not survive the round trip'
    assert [(o.label, o.x_center, o.vertices_local) for o in after['openings']] == \
           [(o.label, o.x_center, o.vertices_local) for o in before['openings']]


def test_imported_model_analyses_through_the_app(app, tmp_path):
    """End to end: import, then press Analyze, and get the box girder's
    own numbers out of the real widget."""
    app._apply_state(box_girder_state())
    p = tmp_path / 'e2e.xlsx'
    pbx.export_state(app._gather_state(), str(p))

    app._apply_state(pbx.default_state())
    app._apply_state(pbx.import_state(str(p)))
    app._analyze()

    assert app.report is not None, app.result_text.get('1.0', 'end')[:400]
    R = [r for _x, r, _m in app.report['support_reactions']]
    assert R[0] == pytest.approx(102974.68, abs=0.01)
    assert R[1] == pytest.approx(625436.58, abs=0.01)
    assert R[2] == pytest.approx(352388.75, abs=0.01)
    assert app.report['governing_opening']['opening'].label == 'H9'


def test_a_failed_import_leaves_the_current_model_untouched(app, tmp_path, monkeypatch):
    """The tab holds work that is expensive to re-enter. A bad file must
    cost the user nothing -- import parses fully before any widget is
    written, so there is no half-applied state to recover from."""
    app._apply_state(box_girder_state())
    before = app._gather_state()

    bad = tmp_path / 'bad.xlsx'
    wb = openpyxl.Workbook()
    wb.active.title = 'NotModel'
    wb.save(str(bad))

    shown = []
    monkeypatch.setattr('apps.perforated_beam.perforated_beam_app.filedialog'
                        '.askopenfilename', lambda **kw: str(bad))
    monkeypatch.setattr('apps.perforated_beam.perforated_beam_app.messagebox'
                        '.showerror', lambda *a, **k: shown.append(a))
    app._import_excel()

    assert shown, 'a failed import must tell the user'
    after = app._gather_state()
    assert after['length'] == before['length']
    assert len(after['openings']) == len(before['openings']) == 26
    assert after['loads'] == before['loads']


def test_export_button_writes_a_file(app, tmp_path, monkeypatch):
    app._apply_state(box_girder_state())
    out = tmp_path / 'button.xlsx'
    monkeypatch.setattr('apps.perforated_beam.perforated_beam_app.filedialog'
                        '.asksaveasfilename', lambda **kw: str(out))
    monkeypatch.setattr('apps.perforated_beam.perforated_beam_app.messagebox'
                        '.showinfo', lambda *a, **k: None)
    app._export_excel()
    assert out.exists()
    assert pbx.import_state(str(out))['length'] == pytest.approx(50400.0)


def test_catalog_beam_round_trips_through_the_app(app, tmp_path):
    """The common case, which the box girder does not exercise: a catalog
    rolled section with the simple 2-support model and hexagonal openings."""
    st = pbx.default_state()
    st.update({
        'length': 9000.0,
        'section_mode': 'catalog',
        'catalog': {'section': 'IPE 400', 'rotate90': False, 'mirror': True},
        'openings': pbm.uniform_layout(pbm.opening_hexagon(300.0, 200.0), 6, 1200.0, 1500.0),
        'loads': [{'type': 'Distributed load', 'x1': 0.0, 'x2': 9000.0,
                   'v1': 12.0, 'v2': 12.0, 'e': 0.0, 'case': 'L'}],
    })
    app._apply_state(st)
    p = tmp_path / 'catalog.xlsx'
    pbx.export_state(app._gather_state(), str(p))
    app._apply_state(pbx.default_state())
    app._apply_state(pbx.import_state(str(p)))

    assert app.section_var.get() == 'IPE 400'
    assert bool(app.catalog_mirror_var.get()) is True
    assert len(app.openings) == 6
    app._analyze()
    assert app.report is not None
    assert app.report['governing_opening'] is not None

    # Unmirrored, the same catalog section DOES get a Cap. F/G member
    # check: `member_check` takes a bare RolledSection only, and the
    # mirror flag above wraps it in an OrientedSection. That is existing
    # behaviour, pinned here so the round-trip is not later blamed for it.
    app.catalog_mirror_var.set(False)
    app._analyze()
    assert app.report['member'] is not None


# ─────────────────────────────────────────────────────────────────────────
# The load case and the two new settings, added 2026-09-10.
#
# A workbook that loses the combination is worse than one that never had
# it: re-imported, it would silently fall back to unfactored loads and
# report utilisations 44% lower with nothing to show that anything had
# changed. The fixtures above gained a 'case' key for the same reason --
# it is now part of what a load IS, not an optional extra.
# ─────────────────────────────────────────────────────────────────────────

def test_the_load_case_round_trips(tmp_path):
    st = pbx.default_state()
    st['loads'] = [
        {'type': 'Distributed load', 'x1': 0.0, 'x2': 6000.0, 'v1': 5.0, 'v2': 5.0,
         'e': 0.0, 'case': 'D'},
        {'type': 'Point load', 'x1': 3000.0, 'v1': 40000.0, 'e': 0.0, 'case': 'L'},
    ]
    path = str(tmp_path / 'cases.xlsx')
    pbx.export_state(st, path)
    back = pbx.import_state(path)
    assert [ld['case'] for ld in back['loads']] == ['D', 'L']


def test_the_combination_and_method_round_trip(tmp_path):
    st = dict(pbx.default_state(), load_combination='asd_dl', vierendeel_method='hand')
    path = str(tmp_path / 'settings.xlsx')
    pbx.export_state(st, path)
    back = pbx.import_state(path)
    assert back['load_combination'] == 'asd_dl'
    assert back['vierendeel_method'] == 'hand'


def test_an_old_workbook_still_imports_with_the_old_behaviour(tmp_path):
    """Every file written before today lacks all three columns. Absent
    must mean exactly what those files meant: unfactored loads, the
    station scan, and no classification."""
    st = pbx.default_state()
    st['loads'] = [{'type': 'Point load', 'x1': 3000.0, 'v1': 40000.0, 'e': 0.0,
                    'case': 'D'}]
    path = str(tmp_path / 'old.xlsx')
    pbx.export_state(st, path)

    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb['Model']
    for row in ws.iter_rows():
        for cell in row:
            if cell.value in ('case', 'load_combination', 'vierendeel_method'):
                cell.value = None            # strip the header -> column gone
    wb.save(path)

    back = pbx.import_state(path)
    assert back['load_combination'] == 'as_entered'
    assert back['vierendeel_method'] == 'station'
    assert back['loads'][0]['case'] == 'L'   # the LARGER factor, never the smaller


def test_an_unknown_combination_key_falls_back_rather_than_raising(tmp_path):
    st = dict(pbx.default_state(), load_combination='asd_dl')
    path = str(tmp_path / 'stale.xlsx')
    pbx.export_state(st, path)
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb['Model']
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'asd_dl':
                cell.value = 'lrfd_from_the_future'
    wb.save(path)
    assert pbx.import_state(path)['load_combination'] == 'as_entered'
