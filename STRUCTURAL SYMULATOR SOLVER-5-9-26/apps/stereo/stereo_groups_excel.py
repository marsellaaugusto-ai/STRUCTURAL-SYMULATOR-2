"""Groups in the workbook: edit a whole branch's section in Excel.

Roadmap v2 4.2. The exported workbook carries a "Groups" sheet, one row per
group. Editing a row and importing the workbook again applies that row to
every rod of the group -- the bulk edit groups exist for, done where a
reader already keeps their numbers.

The sheet is also where the groups THEMSELVES travel: each row lists its
rods as ranges ("0-39, 45"), so a workbook exported from a grouped model
comes back grouped, and moving a rod from one row to another moves it
between groups. An import used to drop every group.

THE RULES, chosen so that importing an unedited workbook changes nothing:

  * A row sets its group's OWN rods, not its subgroups'. A subgroup has a
    row of its own; cascading a parent's value down would overwrite the
    child's rods whenever the child's own row was blank, and a blank cell
    means "leave it".
  * A blank cell leaves every rod as it is. A cell is exported blank when
    the group's rods do not agree on that value (the "mixed" column names
    which), so a mixed group survives the round trip untouched.
  * A profile is applied only to the rods that do not already carry it, so
    re-importing does not reset hand-edited values under an unchanged name.
    The numeric cells are then applied on top, which is how a catalog
    section with one value changed is written -- but only the cells that
    were EDITED: a cell still holding the value it was exported with is
    left alone, or changing just the profile would have the old A, I and c
    beside it written straight back over the new section.
  * A changed I without a c in the same row drops the rod's catalog depth:
    the depth belonged to the old section, and I/c with a stale c
    overstates bending capacity (see stereo_profiles.write_section).
  * A rod listed in two rows is an error, not a choice between them -- a
    rod belongs to exactly one group (stereo_groups, rule 1).
"""
from apps.stereo import stereo_groups as sgp

SHEET = 'Groups'

# (column header, member key, kind). The machine-read columns, in order.
FIELDS = (
    ('profile', 'profile', 'text'),
    ('conn', 'conn', 'conn'),
    ('E_GPa', 'E', 'num'),
    ('A_cm2', 'A', 'num'),
    ('I_cm4', 'I', 'num'),
    ('J_cm4', 'J', 'num'),
    ('Fy_MPa', 'Fy', 'num'),
    ('Fu_MPa', 'Fu', 'num'),
    ('K', 'K', 'num'),
    ('r_gyr_cm', 'r_gyr', 'num'),
    ('c_cm', 'c_cm', 'num'),
)
# Facts about the GROUP, not its rods: left out of the analysis, and where
# it sits among the iterations being compared (stereo_compare). Read back by
# name; a workbook without them reads as before.
GROUP_FLAGS = ('excluded', 'iteration', 'stage', 'position')
HEADERS = ('id', 'name', 'parent', 'rods') + tuple(f[0] for f in FIELDS) \
    + GROUP_FLAGS
# Written for the reader, never read back.
INFO_HEADERS = ('info: rods', 'info: worst util', 'info: mixed')

_REL_TOL = 1e-9


# ── rod lists as text ─────────────────────────────────────────────────────

def rods_to_ranges(idx):
    """[0,1,2,3,7,9,10] -> '0-3, 7, 9-10'."""
    idx = sorted({int(i) for i in idx})
    out = []
    k = 0
    while k < len(idx):
        j = k
        while j + 1 < len(idx) and idx[j + 1] == idx[j] + 1:
            j += 1
        out.append(str(idx[k]) if j == k else '%d-%d' % (idx[k], idx[j]))
        k = j + 1
    return ', '.join(out)


def parse_ranges(text):
    """'0-3, 7; 9 10' -> [0, 1, 2, 3, 7, 9, 10]. Raises ValueError on junk.

    Commas, semicolons and spaces all separate; a range is a-b with a <= b.
    A bare number typed into the cell arrives from Excel as a float, which
    is accepted when it is whole.
    """
    if text is None:
        return []
    if isinstance(text, (int, float)):
        if float(text) != int(text):
            raise ValueError('%r is not a rod number' % (text,))
        return [int(text)]
    out = []
    for tok in str(text).replace(';', ',').replace(' ', ',').split(','):
        tok = tok.strip()
        if not tok:
            continue
        if '-' in tok[1:]:
            a, b = tok.split('-', 1) if not tok.startswith('-') else (tok, '')
            try:
                lo, hi = int(float(a)), int(float(b))
            except ValueError:
                raise ValueError('%r is not a rod range' % (tok,))
            if lo > hi:
                raise ValueError('%r runs backwards' % (tok,))
            out.extend(range(lo, hi + 1))
        else:
            try:
                out.append(int(float(tok)))
            except ValueError:
                raise ValueError('%r is not a rod number' % (tok,))
    return sorted(set(out))


# ── what a group's row says ───────────────────────────────────────────────

def _same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= _REL_TOL * max(1.0, abs(a), abs(b))
    return a == b


def common_values(members, rods):
    """({member key: value}, [mixed keys]) over these rods.

    A key the rods agree on gets its value; one they do not is left out and
    named in the mixed list; one no rod has at all is left out silently.
    """
    values, mixed = {}, []
    for _hdr, key, _kind in FIELDS:
        seen = []
        for i in rods:
            if 0 <= i < len(members):
                v = members[i].get(key)
                if key == 'conn' and v is None:
                    v = 'pin'
                seen.append(v)
        present = [v for v in seen if v not in (None, '')]
        if not present:
            continue
        first = present[0]
        if len(present) == len(seen) and all(_same(first, v) for v in present):
            values[key] = first
        else:
            mixed.append(key)
    return values, mixed


def group_rows(groups, members, checks=None):
    """One dict per group, in tree order: what the sheet shows."""
    rows = []
    for g, lvl in sgp.walk(groups):
        own = sorted(g['members'])
        values, mixed = common_values(members, own)
        worst = None
        if checks:
            u = [checks[i].get('util') for i in own
                 if i < len(checks) and checks[i].get('checked')
                 and checks[i].get('util') is not None]
            worst = max(u) if u else None
        rows.append({'id': g['id'], 'name': g['name'], 'parent': g['parent'],
                     'flags': {k: g.get(k) for k in GROUP_FLAGS},
                     'level': lvl, 'rods': rods_to_ranges(own),
                     'values': values, 'mixed': mixed, 'n_rods': len(own),
                     'worst': worst})
    return rows


# ── writing ───────────────────────────────────────────────────────────────

def write_groups_sheet(wb, groups, members, checks=None):
    """Add the Groups sheet to an openpyxl workbook. Returns the sheet."""
    from openpyxl.styles import Font, PatternFill, Alignment
    ws = wb.create_sheet(SHEET)
    ws.cell(row=1, column=1, value='GROUPS — edit a whole branch here, then '
                                   'Import from Excel').font = Font(
        bold=True, size=12, color='1F4E79')
    notes = [
        'A filled cell sets that value on every rod the row lists; a BLANK '
        'cell leaves each rod as it is. A blank means the rods do not agree '
        '(see the mixed column).',
        'rods: this group\'s own rods, as ranges (0-39, 45). Move a rod by '
        'moving its number to another row; a rod may be in one row only. '
        'Rods in no row are Ungrouped.',
        'profile: a catalog name (IPE 200, CHS 76.1x3.6, …) or a name from '
        'the Model sheet\'s [PROFILES]. Numbers in the same row are applied '
        'on top of it.',
        'parent: the id of the group this one sits inside, or blank. A row '
        'sets its own rods only — a subgroup has its own row.',
        'Columns headed "info:" are for reading and are not imported.',
    ]
    for k, text in enumerate(notes):
        ws.cell(row=2 + k, column=1, value=text).font = Font(
            italic=True, size=9, color='5F6368')
    hdr_row = 2 + len(notes) + 1
    hdr_fill = PatternFill('solid', fgColor='404040')
    info_fill = PatternFill('solid', fgColor='8C8C8C')
    for col, h in enumerate(HEADERS + INFO_HEADERS, 1):
        c = ws.cell(row=hdr_row, column=col, value=h)
        c.font = Font(bold=True, color='FFFFFF', size=10)
        c.fill = info_fill if h.startswith('info:') else hdr_fill
        c.alignment = Alignment(horizontal='center')
    r = hdr_row + 1
    for row in group_rows(groups, members, checks):
        cells = [row['id'], row['name'],
                 row['parent'] if row['parent'] is not None else None,
                 row['rods']]
        cells += [row['values'].get(key) for _h, key, _k in FIELDS]
        flags = row['flags']
        cells += [1 if flags.get('excluded') else None] + [
            flags.get(k) or None for k in GROUP_FLAGS[1:]]
        cells += [row['n_rods'],
                  None if row['worst'] is None else round(row['worst'], 3),
                  ', '.join(row['mixed'])]
        for col, v in enumerate(cells, 1):
            c = ws.cell(row=r, column=col, value=v)
            # Shown to four places, STORED in full: a catalog area is
            # 8.19955682586936 cm², and rounding the stored value would make
            # an untouched cell read as an edit on the way back in.
            if isinstance(v, float) and col > 4:
                c.number_format = '0.0###'
        r += 1
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=3)
    return ws


# ── reading ───────────────────────────────────────────────────────────────

def read_groups_sheet(wb):
    """The sheet's rows as dicts keyed by header, or None if there is none."""
    if SHEET not in wb.sheetnames:
        return None
    rows = list(wb[SHEET].iter_rows(values_only=True))
    hdr_i = next((i for i, r in enumerate(rows)
                  if r and r[0] is not None and str(r[0]).strip() == 'id'),
                 None)
    if hdr_i is None:
        raise ValueError('The Groups sheet has no header row (a row starting '
                         'with "id").')
    headers = [str(h).strip() if h is not None else '' for h in rows[hdr_i]]
    out = []
    for r in rows[hdr_i + 1:]:
        if not r or all(v in (None, '') for v in r):
            continue
        out.append({headers[j]: r[j] for j in range(min(len(headers), len(r)))
                    if headers[j] and not headers[j].startswith('info:')})
    return out


def _num(v, where):
    try:
        return float(v)
    except (TypeError, ValueError):
        raise ValueError('%s: %r is not a number' % (where, v))


def build_groups(sheet_rows, n_members):
    """Groups from the sheet's rows -- membership and nesting only.

    Validates what would otherwise corrupt the model quietly: duplicate
    ids, a parent that does not exist or makes a cycle, a rod number the
    model does not have, a rod in two rows.
    """
    groups, seen_ids, owner = [], set(), {}
    next_id = 1 + max((int(_num(r['id'], 'id')) for r in sheet_rows
                       if r.get('id') not in (None, '')), default=0)
    ids = []
    for k, r in enumerate(sheet_rows):
        if r.get('id') in (None, ''):
            gid = next_id
            next_id += 1
        else:
            gid = int(_num(r['id'], 'Groups row %d, id' % (k + 1)))
        if gid in seen_ids:
            raise ValueError('Two rows of the Groups sheet have id %d.' % gid)
        seen_ids.add(gid)
        ids.append(gid)
    for k, (r, gid) in enumerate(zip(sheet_rows, ids)):
        name = str(r.get('name') or '').strip() or 'Group %d' % gid
        try:
            rods = parse_ranges(r.get('rods'))
        except ValueError as exc:
            raise ValueError('Groups row "%s", rods: %s' % (name, exc))
        bad = [i for i in rods if not (0 <= i < n_members)]
        if bad:
            raise ValueError('Groups row "%s" lists rod %d, but the model has '
                             'rods 0-%d.' % (name, bad[0], n_members - 1))
        for i in rods:
            if i in owner:
                raise ValueError('Rod %d is listed in two groups, "%s" and '
                                 '"%s". A rod belongs to one group only.'
                                 % (i, owner[i], name))
            owner[i] = name
        parent = r.get('parent')
        parent = None if parent in (None, '') else int(
            _num(parent, 'Groups row "%s", parent' % name))
        g = {'id': gid, 'name': name, 'parent': parent, 'members': set(rods)}
        if str(r.get('excluded') or '').strip().lower() not in (
                '', '0', 'no', 'false', 'n'):
            g['excluded'] = True
        for k in GROUP_FLAGS[1:]:
            v = r.get(k)
            if v not in (None, '') and str(v).strip():
                v = str(v).strip()
                g[k] = v[:-2] if v.endswith('.0') else v
        groups.append(g)
    for g in groups:
        if g['parent'] is not None and sgp.find(groups, g['parent']) is None:
            raise ValueError('Group "%s" names parent %d, which is not a row '
                             'of the sheet.' % (g['name'], g['parent']))
        if g['parent'] == g['id'] or g['id'] in sgp.descendant_ids(
                groups, g['id']) or _cycles(groups, g):
            raise ValueError('Group "%s" is nested inside itself.' % g['name'])
    return groups


def _cycles(groups, g):
    seen, cur = set(), g
    while cur is not None and cur['parent'] is not None:
        if cur['id'] in seen:
            return True
        seen.add(cur['id'])
        cur = sgp.find(groups, cur['parent'])
    return False


def apply_group_values(sheet_rows, groups, members, profiles=None):
    """Write each row's filled cells onto its group's own rods.

    Returns the report: one line per change actually made, so the reader
    sees what the workbook did. `profiles` is the model's named profiles,
    for a profile cell that is not a catalog name.
    """
    from apps.stereo import stereo_profiles as sp
    from apps.stereo import stereo_checks as sk
    report = []
    by_id = {}
    for r, g in zip(sheet_rows, groups):
        by_id[g['id']] = r
    for g, _lvl in sgp.walk(groups):
        r = by_id.get(g['id'])
        rods = sorted(g['members'])
        if r is None or not rods:
            continue
        name = g['name']
        # What the rods said BEFORE this row touches them -- which is what
        # the row said when it was exported. A cell still equal to it is a
        # cell nobody edited, and must not override a change made in
        # another cell of the same row: change only the profile, and the
        # old A, I and c still sitting beside it would otherwise be written
        # straight back over the new section.
        before, _mixed = common_values(members, rods)

        prof = r.get('profile')
        if prof not in (None, '') and str(prof).strip():
            prof = str(prof).strip()
            todo = [i for i in rods if members[i].get('profile') != prof]
            if todo:
                if prof in sp.CATALOG:
                    sk.apply_recommendation(members, todo, prof)
                    report.append('%s: section %s on %d rod(s)'
                                  % (name, prof, len(todo)))
                elif profiles and prof in profiles:
                    for i in todo:
                        sp.write_section(members[i], profiles[prof])
                        members[i]['profile'] = prof
                    report.append('%s: profile %s on %d rod(s)'
                                  % (name, prof, len(todo)))
                else:
                    report.append('%s: profile "%s" is not in the catalog or '
                                  'the model\'s profiles -- left as it was'
                                  % (name, prof))

        conn = r.get('conn')
        if conn not in (None, '') and str(conn).strip():
            conn = str(conn).strip().lower()
            if conn not in ('pin', 'rigid'):
                report.append('%s: conn "%s" is neither pin nor rigid -- '
                              'left as it was' % (name, conn))
            else:
                kept = 0
                todo = []
                for i in rods:
                    if members[i].get('conn', 'pin') == conn:
                        continue
                    if conn == 'pin' and members[i].get('rigid_required'):
                        kept += 1
                        continue
                    todo.append(i)
                for i in todo:
                    members[i]['conn'] = conn
                if todo:
                    report.append('%s: conn %s on %d rod(s)'
                                  % (name, conn, len(todo)))
                if kept:
                    report.append('%s: %d rod(s) must stay rigid and did'
                                  % (name, kept))

        c_cell = r.get('c_cm')
        c_given = c_cell not in (None, '') and not (
            'c_cm' in before and _same(before['c_cm'], _num(
                c_cell, 'Groups row "%s", c_cm' % name)))
        I_changed = set()
        for hdr, key, kind in FIELDS:
            if kind != 'num':
                continue
            v = r.get(hdr)
            if v in (None, ''):
                continue
            v = _num(v, 'Groups row "%s", %s' % (name, hdr))
            if v <= 0:
                raise ValueError('Groups row "%s", %s: %g is not a section '
                                 'value' % (name, hdr, v))
            if key in before and _same(before[key], v):
                continue            # as exported: not an edit
            todo = [i for i in rods if not _same(members[i].get(key), v)]
            if not todo:
                continue
            old = {members[i].get(key) for i in todo}
            for i in todo:
                members[i][key] = v
            if key == 'I':
                I_changed |= set(todo)
            was = ('%g' % next(iter(old))) if len(old) == 1 and \
                None not in old else 'mixed'
            report.append('%s: %s %s -> %g on %d rod(s)'
                          % (name, hdr, was, v, len(todo)))
        if I_changed and not c_given:
            dropped = [i for i in I_changed if members[i].pop('c_cm', None)]
            for i in I_changed:               # the weak axis was the old
                members[i].pop('Iw', None)    # section's too
                members[i].pop('cw_cm', None)
            if dropped:
                report.append('%s: I changed without a c, so %d rod(s) lost '
                              'their catalog depth (bending falls back to '
                              'the conservative estimate)' % (name, len(dropped)))
    return report


def import_groups(path, members, profiles=None):
    """(groups, report) from a workbook's Groups sheet, applied to members.

    (None, []) when the workbook has no Groups sheet -- one written before
    groups existed, or by a SketchUp model with nothing grouped in it -- so
    the caller can tell "no groups in this file" from "the file says there
    are no groups".

    The SketchUp extension writes this sheet too, from the groups and
    components the model was built with, and writes only its id/name/parent/
    rods columns: SketchUp knows nothing about steel sections, and every
    other cell arrives blank, which means "leave this rod as it is".
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = read_groups_sheet(wb)
    if rows is None:
        return None, []
    groups = build_groups(rows, len(members))
    report = apply_group_values(rows, groups, members, profiles)
    return groups, report
