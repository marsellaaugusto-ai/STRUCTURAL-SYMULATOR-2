"""
beam_app.py — Beam tab.

Hermitian (cubic-Hermite) FEM beam solver (BeamModel), exact-statics V/M
recovery + deflection integration (BeamResult), and the Beam schematic +
diagram UI (BeamApp), plus this tab's Excel report/import.
"""
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import math, os, sys, subprocess

import units

from common import (
    _ensure_openpyxl,
    PANEL_W, INIT_CW, INIT_CH, INIT_DH,
    ScrollPanel, WrapBar,
    _beam_gauss_solve, _GAUSS5_NODES, _GAUSS5_WEIGHTS,
    _nice_ticks, _find_diagram_maxima, make_shape_fn,
    draw_moment_arrow, LoadScale, declutter_text,
)

class BeamModel:
    def __init__(self, length):
        self.L = length
        self.supports = []     # [{'x':, 'type':}]
        self.point_loads = []  # [{'x':, 'P':}]   +P = downward (N)
        self.moments = []      # [{'x':, 'M':}]   +M = CCW (N*m)
        self.dloads = []       # [{'x1':,'x2':,'w1':,'w2':}]  +w = downward (N/m)
        self.nonuniform_loads = []  # [{'fn': q(x)->N/m (+down), 'x1':, 'x2':}]
        self.EI = None         # N*m^2

    def _on_beam(self, x, what):
        """Reject a station that is not on the beam.

        Nothing used to check these against the beam's own length, so a support
        or load at x = 99 on a 6 m beam was accepted and silently EXTENDED the
        mesh to 99 m while self.L stayed 6.0 -- the analysed structure was not
        the one the user described, and a 10 kN load came back as a 155 kN
        reaction with no warning (2026-09-05 finding B-5).
        """
        tol = 1e-9 * max(1.0, abs(self.L))
        if not (-tol <= x <= self.L + tol):
            raise ValueError(
                f"{what} at x = {x:g} m is not on the beam, which spans 0 to "
                f"{self.L:g} m. Move it onto the beam, or set the beam length "
                f"first.")
        return min(max(x, 0.0), self.L)

    def add_support(self, x, type_):
        self.supports.append({'x': self._on_beam(x, 'Support'), 'type': type_})

    def add_point_load(self, x, P):
        self.point_loads.append({'x': self._on_beam(x, 'Point load'), 'P': P})

    def add_moment(self, x, M):
        self.moments.append({'x': self._on_beam(x, 'Applied moment'), 'M': M})

    def add_dload(self, x1, x2, w1, w2):
        # Normalise a right-to-left entry. q() tests `x1 <= x <= x2`, which is
        # unsatisfiable when x1 > x2, so a reversed segment used to contribute
        # NOTHING at all -- zero reactions, zero moment, no warning (finding
        # B-4). The intensities travel with their own ends, or the ramp would
        # come out flipped. The Perforated Beam tab normalises the same way.
        x1 = self._on_beam(x1, 'Distributed load start')
        x2 = self._on_beam(x2, 'Distributed load end')
        if x2 < x1:
            x1, x2, w1, w2 = x2, x1, w2, w1
        self.dloads.append({'x1': x1, 'x2': x2, 'w1': w1, 'w2': w2})

    def add_nonuniform_load(self, fn, x1, x2):
        x1 = self._on_beam(x1, 'Non-uniform load start')
        x2 = self._on_beam(x2, 'Non-uniform load end')
        if x2 < x1:
            x1, x2 = x2, x1
        self.nonuniform_loads.append({'fn': fn, 'x1': x1, 'x2': x2})

    # Longest element allowed, as a fraction of the modelled span. A mesh made
    # only of the load/support discontinuities can be too coarse to solve at
    # all: a fixed-fixed beam under a full-span UDL has exactly two nodes, all
    # four of its DOF are restrained, and solve() rejected it as "fully
    # constrained" even though it is the most common indeterminate case in any
    # textbook (2026-09-05 diagnosis, finding B-1). Guaranteeing a minimum mesh
    # removes that whole class of failure.
    #
    # This changes no result. Cubic-Hermite elements with consistent load
    # vectors are nodally exact for the load types this model supports, so the
    # extra nodes only make the deflection integration marginally finer -- and
    # the cost is bounded: subdividing every gap to at most span/4 adds at most
    # four elements in total, whatever the load layout.
    MAX_ELEM_FRACTION = 0.25

    def _node_positions(self):
        xs = {0.0, round(self.L, 9)}
        for s in self.supports: xs.add(round(s['x'], 9))
        for p in self.point_loads: xs.add(round(p['x'], 9))
        for m in self.moments: xs.add(round(m['x'], 9))
        for d in self.dloads:
            xs.add(round(d['x1'], 9)); xs.add(round(d['x2'], 9))
        for d in self.nonuniform_loads:
            xs.add(round(d['x1'], 9)); xs.add(round(d['x2'], 9))
        return self._refine(sorted(xs))

    def _refine(self, stations):
        """Split any gap longer than MAX_ELEM_FRACTION of the modelled span.

        The cap is taken from the span the stations actually cover, not from
        self.L, so a coordinate lying outside the beam cannot make the mesh
        explode -- the added element count stays bounded either way.
        """
        if len(stations) < 2:
            return stations
        max_len = (stations[-1] - stations[0]) * self.MAX_ELEM_FRACTION
        if max_len <= 0:
            return stations
        out = [stations[0]]
        for a, b in zip(stations, stations[1:]):
            seg = b - a
            n_sub = max(1, int(math.ceil(seg / max_len - 1e-9)))
            for k in range(1, n_sub):
                out.append(round(a + seg * k / n_sub, 9))
            out.append(b)
        # A rounded interior point can land on the station that follows it when
        # a gap is degenerate; keep the mesh strictly increasing.
        uniq = [out[0]]
        for v in out[1:]:
            if v > uniq[-1] + 1e-12:
                uniq.append(v)
        return uniq

    def q(self, x):
        """Total distributed load (N/m, +down convention) at x, summing the
        linear-ramp segments (self.dloads) and any non-uniform function-based
        segments (self.nonuniform_loads). NOTE: this is evaluated only at
        interior Gauss points during solve() (see the endpoint-avoiding
        quadrature below) — never call it exactly at a load segment's own
        x1/x2, since a user-supplied non-uniform load may be singular there
        by construction (e.g. 1/sqrt(...) at its own domain edge)."""
        total = 0.0
        for d in self.dloads:
            if d['x1'] - 1e-9 <= x <= d['x2'] + 1e-9:
                span = d['x2'] - d['x1']
                frac = (x - d['x1']) / span if span > 1e-12 else 0.0
                w = d['w1'] + (d['w2'] - d['w1']) * frac
                total += -w
        for d in self.nonuniform_loads:
            if d['x1'] - 1e-9 <= x <= d['x2'] + 1e-9:
                total += -d['fn'](x)
        return total

    def solve(self):
        if self.EI is None:
            raise ValueError("EI must be set before solving")
        EI = self.EI
        xs = self._node_positions()
        n_nodes = len(xs)
        idx_of = {x: i for i, x in enumerate(xs)}
        dof = 2 * n_nodes
        K = [[0.0] * dof for _ in range(dof)]
        F = [0.0] * dof

        elements = []
        for e in range(n_nodes - 1):
            x1, x2 = xs[e], xs[e + 1]
            Le = x2 - x1
            if Le < 1e-9:
                continue
            elements.append((e, e + 1, x1, x2, Le))

        for ei, (n1, n2, x1, x2, Le) in enumerate(elements):
            k = EI / Le ** 3
            L = Le
            ke = [
                [12 * k, 6 * L * k, -12 * k, 6 * L * k],
                [6 * L * k, 4 * L ** 2 * k, -6 * L * k, 2 * L ** 2 * k],
                [-12 * k, -6 * L * k, 12 * k, -6 * L * k],
                [6 * L * k, 2 * L ** 2 * k, -6 * L * k, 4 * L ** 2 * k],
            ]
            gidx = [2 * n1, 2 * n1 + 1, 2 * n2, 2 * n2 + 1]
            for i in range(4):
                for j in range(4):
                    K[gidx[i]][gidx[j]] += ke[i][j]

            has_load = any(
                not (d['x2'] <= x1 + 1e-9 or d['x1'] >= x2 - 1e-9) for d in self.dloads
            ) or any(
                not (d['x2'] <= x1 + 1e-9 or d['x1'] >= x2 - 1e-9) for d in self.nonuniform_loads
            )
            if has_load:
                def vfn(xx, x1=x1, L=L):
                    n = (xx - x1) / L
                    qq = self.q(xx)
                    return (qq * (1 - 3*n**2 + 2*n**3),
                            qq * L * (n - 2*n**2 + n**3),
                            qq * (3*n**2 - 2*n**3),
                            qq * L * (-n**2 + n**3))
                f_eq = _adaptive_vector_integral(vfn, x1, x2, dim=4)
                for i in range(4):
                    F[gidx[i]] += f_eq[i]

        for p in self.point_loads:
            i = idx_of[round(p['x'], 9)]
            F[2 * i] += -p['P']
        for m in self.moments:
            i = idx_of[round(m['x'], 9)]
            F[2 * i + 1] += m['M']

        constrained = {}
        for s in self.supports:
            i = idx_of[round(s['x'], 9)]
            t = s['type']
            if t in ('pin', 'roller'):
                constrained[2 * i] = 0.0
            elif t == 'fixed':
                constrained[2 * i] = 0.0
                constrained[2 * i + 1] = 0.0
            elif t == 'guided':
                constrained[2 * i + 1] = 0.0
            else:
                raise ValueError(f"Unknown support type {t}")

        free = [i for i in range(dof) if i not in constrained]
        if not free:
            raise ValueError("Beam is fully constrained; nothing to solve")

        Kff = [[K[i][j] for j in free] for i in free]
        Ff = [F[i] for i in free]
        d_free = _beam_gauss_solve(Kff, Ff)
        if d_free is None:
            raise ValueError(
                "Singular stiffness matrix — beam is a mechanism "
                "(unstable / insufficient supports).")

        d = [0.0] * dof
        for k_, i in enumerate(free):
            d[i] = d_free[k_]

        Kd = [sum(K[i][j] * d[j] for j in range(dof)) for i in range(dof)]
        R = [Kd[i] - F[i] for i in range(dof)]

        return BeamResult(self, xs, elements, d, R, idx_of, EI)

class BeamResult:
    def __init__(self, model, xs, elements, d, R, idx_of, EI):
        self.model = model
        self.xs = xs
        self.elements = elements
        self.d = d
        self.R = R
        self.idx_of = idx_of
        self.EI = EI

    def reaction_at(self, x):
        """Returns (Fy, M) at the support/node located at x. Fy: vertical
        reaction (+up). M: raw rotational-DOF residual (K*d - F at that
        DOF) — used internally by _V_M_at below, which already accounts
        for its sign convention via the "- Mr" term. For a *display* value
        directly comparable to the M(x) diagram's sagging-positive
        convention, see BeamApp's presentation layer, which negates this
        value specifically at the beam's leftmost node (x=0) — verified
        against fixed supports at the left end, right end, and both ends
        simultaneously (fixed-fixed): negating only the x=0 case matches
        the diagram's M(x) in every case tested."""
        i = self.idx_of[round(x, 9)]
        return self.R[2 * i], self.R[2 * i + 1]

    def _V_M_at(self, x, side='right'):
        eps = 1e-7
        xx = x - eps if side == 'left' else x + eps
        xx = max(0.0, min(self.model.L, xx))

        V = 0.0
        M = 0.0
        for s in self.model.supports:
            if s['x'] <= xx + 1e-9:
                Fy, Mr = self.reaction_at(s['x'])
                V += Fy
                M += Fy * (x - s['x']) - Mr
        for p in self.model.point_loads:
            if p['x'] <= xx + 1e-9:
                Fy = -p['P']
                V += Fy
                M += Fy * (x - p['x'])
        for mm in self.model.moments:
            if mm['x'] <= xx + 1e-9:
                M -= mm['M']

        for d in self.model.dloads:
            x1, x2 = d['x1'], d['x2']
            if xx <= x1 + 1e-9:
                continue
            xend = min(xx, x2)
            if xend - x1 < 1e-12:
                continue
            span = x2 - x1
            def qfn(s, d=d, span=span):
                frac = (s - d['x1']) / span if span > 1e-12 else 0.0
                w = d['w1'] + (d['w2'] - d['w1']) * frac
                return -w
            Vq, Mq = _gauss_VM_integral(qfn, x1, xend, x)
            V += Vq
            M += Mq

        for d in self.model.nonuniform_loads:
            x1, x2 = d['x1'], d['x2']
            if xx <= x1 + 1e-9:
                continue
            xend = min(xx, x2)
            if xend - x1 < 1e-12:
                continue
            def qfn(s, d=d):
                return -d['fn'](s)
            Vq, Mq = _gauss_VM_integral(qfn, x1, xend, x)
            V += Vq
            M += Mq

        return V, M

    def shear_at(self, x, side='right'):
        return self._V_M_at(x, side=side)[0]

    def moment_at(self, x, side='right'):
        return self._V_M_at(x, side=side)[1]

    def deflection(self, x):
        return self._theta_v_at(x)[1]

    def rotation(self, x):
        return self._theta_v_at(x)[0]

    def _theta_v_at(self, x_target):
        i0 = self.idx_of[round(0.0, 9)]
        v = self.d[2 * i0]
        th = self.d[2 * i0 + 1]
        EI = self.EI

        x_prev = 0.0
        for (n1, n2, x1, x2, Le) in self.elements:
            if x2 <= x_prev + 1e-12:
                continue
            seg_end = min(x2, x_target)
            if seg_end <= x_prev + 1e-12:
                break
            n_sub = 30
            h = (seg_end - x_prev) / n_sub
            xs_local = [x_prev + k * h for k in range(n_sub + 1)]
            Ms = [self.moment_at(xx, side='right') for xx in xs_local]
            thetas = [th]
            for k_ in range(n_sub):
                xm = (xs_local[k_] + xs_local[k_ + 1]) / 2
                Mm = self.moment_at(xm, side='right')
                seg_h = xs_local[k_ + 1] - xs_local[k_]
                dtheta = seg_h / 6 * (Ms[k_] + 4 * Mm + Ms[k_ + 1]) / EI
                thetas.append(thetas[-1] + dtheta)
            vs = [v]
            for k_ in range(n_sub):
                seg_h = xs_local[k_ + 1] - xs_local[k_]
                vs.append(vs[-1] + seg_h / 2 * (thetas[k_] + thetas[k_ + 1]))
            th = thetas[-1]
            v = vs[-1]
            x_prev = seg_end
            if seg_end >= x_target - 1e-12:
                break
        return th, v

    def sample_diagram(self, n_per_element=30):
        """Samples V, M, and deflection along the beam for plotting.

        At any station coinciding with a discontinuity (a support reaction,
        point load, or applied moment), two points are emitted at that same
        x: the value just BEFORE the discontinuity, then the value just
        AFTER it — in that order. Plotted as connected line segments, this
        draws a clean vertical jump exactly at that point. Getting the order
        backwards (after-then-before) doesn't affect the vertical segment
        between the twin points themselves, but it DOES corrupt the sloped
        segments connecting to the neighboring stations on either side —
        the diagonal from the previous station ends up aiming at the wrong
        (post-jump) value, producing a visible zigzag/overshoot right at
        the discontinuity instead of a clean diagonal-jump-diagonal shape.
        """
        xs_out, Vs, Ms = [], [], []
        for (n1, n2, x1, x2, Le) in self.elements:
            for k in range(n_per_element + 1):
                x = x1 + Le * k / n_per_element
                # Start of an element (k=0): capture the value just AFTER
                # whatever discontinuity sits at this x (post-jump).
                # End of an element (k=n_per_element): capture the value
                # just BEFORE the discontinuity at the NEXT element's start
                # (pre-jump). Interior points have no discontinuity, so the
                # side choice there doesn't matter.
                side = 'left' if k == n_per_element else 'right'
                V = self.shear_at(x, side=side)
                M = self.moment_at(x, side=side)
                xs_out.append(x); Vs.append(V); Ms.append(M)

        # The right end of the beam (x = L) has no following element to
        # supply the closing side='right' "just after the final reaction"
        # point — the loop's own last point (side='left') already gives the
        # correct pre-final-reaction value, so just append the post-reaction
        # closing value (typically 0) after it.
        if xs_out and abs(xs_out[-1] - self.model.L) < 1e-9:
            xL = xs_out[-1]
            xs_out.append(xL)
            Vs.append(self.shear_at(xL, side='right'))
            Ms.append(self.moment_at(xL, side='right'))

        # Deflection along the SAME stations, integrated forward in one pass
        # rather than by calling deflection(x) per station. That method
        # (_theta_v_at) re-integrates the moment diagram from x = 0 every time
        # it is called, and each of its 30 sub-steps per element is an
        # O(loads) moment_at -- so sampling the curve station by station cost
        # O(stations * elements * loads), cubic in the model size (a 40-load
        # beam took 5.4 s, and this was the single biggest reason the test
        # suite ran over an hour). The moment values are already sampled just
        # above; integrating them forward once, with the same Simpson (for the
        # rotation) then trapezoid (for the deflection) rule _theta_v_at uses,
        # reproduces the same curve at a fraction of the cost. The point-query
        # deflection(x) is deliberately left untouched, so its closed-form
        # accuracy tests still bind and this fast path is checked against them.
        defl = self._integrate_deflection(xs_out, Ms)
        return {'x': xs_out, 'V': Vs, 'M': Ms, 'v': defl}

    def _integrate_deflection(self, xs, Ms):
        """Deflection at each station in `xs`, from the moment samples `Ms`.

        One cumulative sweep: rotation theta from integrating M/EI, deflection
        v from integrating theta, seeded at x = 0 from the left node's own
        solved DOF exactly as `_theta_v_at` seeds them. A zero-width interval
        (the twin before/after points that draw a shear or moment jump) leaves
        theta and v unchanged, which is correct -- slope and deflection are
        continuous across a jump in V or M. Only the sub-interval midpoint
        moment is evaluated fresh; the endpoints reuse the already-sampled
        `Ms`, whose side choice (pre-jump at an element's end, post-jump at the
        next element's start) is exactly the interior value each interval
        needs.
        """
        i0 = self.idx_of[round(0.0, 9)]
        v = self.d[2 * i0]
        th = self.d[2 * i0 + 1]
        EI = self.EI
        out = [v]
        for i in range(1, len(xs)):
            h = xs[i] - xs[i - 1]
            if h <= 1e-12:
                out.append(v)                 # zero-width twin point
                continue
            Mm = self.moment_at(0.5 * (xs[i - 1] + xs[i]), side='right')
            dth = h / 6.0 * (Ms[i - 1] + 4.0 * Mm + Ms[i]) / EI
            th_new = th + dth
            v = v + h / 2.0 * (th + th_new)
            th = th_new
            out.append(v)
        return out

def _adaptive_vector_integral(vfn, a, b, dim, tol=1e-7, max_depth=20):
    """Same adaptive-refinement scheme as `_adaptive_integral`, but for a
    vector-valued integrand `vfn(x)` that returns `dim` values sharing the
    same evaluation points (e.g. q(x) weighted by each of a beam element's
    4 Hermite shape functions in one pass, or a load's (force, moment)
    contribution together). Refinement is driven by the combined vector's
    magnitude of disagreement between the whole-panel and half-panel
    estimates. See `_adaptive_integral` for why a minimum panel width is
    also enforced."""
    min_width = max(1e-12, (b - a) * 1e-10) if b > a else 1e-12

    def gauss5(lo, hi):
        h = (hi - lo) / 2
        mid = (lo + hi) / 2
        acc = [0.0] * dim
        for node, wt in zip(_GAUSS5_NODES, _GAUSS5_WEIGHTS):
            try:
                vals = vfn(mid + node*h)
            except (ZeroDivisionError, ValueError, OverflowError):
                continue
            for i in range(dim):
                acc[i] += wt * vals[i]
        return [v * h for v in acc]

    def rec(lo, hi, whole, depth):
        if depth >= max_depth or (hi - lo) <= min_width:
            return whole
        mid = (lo + hi) / 2
        left = gauss5(lo, mid)
        right = gauss5(mid, hi)
        combined = [left[i] + right[i] for i in range(dim)]
        err = math.sqrt(sum((combined[i]-whole[i])**2 for i in range(dim)))
        scale = max(1.0, math.sqrt(sum(v*v for v in whole)))
        if err <= tol * scale:
            return combined
        lr = rec(lo, mid, left, depth+1)
        rr = rec(mid, hi, right, depth+1)
        return [lr[i] + rr[i] for i in range(dim)]

    if b <= a:
        return [0.0] * dim
    return rec(a, b, gauss5(a, b), 0)

def _gauss_VM_integral(qfn, a, b, x, n_panels=None):
    """Returns (∫ q(s) ds, ∫ q(s)*(x-s) ds) over [a,b] — the resultant force
    and its moment about station x — via adaptive Gauss-Legendre quadrature
    (see `_adaptive_vector_integral`). Accurate even where q is singular
    exactly at a or b, or varies sharply near either edge; `n_panels` is
    accepted for backward compatibility but no longer used."""
    def vfn(s):
        qs = qfn(s)
        return (qs, qs * (x - s))
    return tuple(_adaptive_vector_integral(vfn, a, b, dim=2))

def export_beam_excel(state, path, result=None, model=None):
    """
    Writes a workbook with:
      • 'Model'   — every input needed to rebuild this exact beam (length,
                    supports, loads, section) — paired with import_beam_excel().
      • 'Results' — reactions, extremes, and a station-by-station table of
                    V, M, and deflection (only if `result` given).
    `state` is a plain dict with keys: length, supports, point_loads,
    moments, dloads, nonuniform_loads, profile.
    """
    if not _ensure_openpyxl():
        raise RuntimeError('openpyxl is not available.')
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('Model')
    row = 1
    ws.cell(row=row, column=1, value='BEAM MODEL DATA — for Import from Excel (do not reorder columns)')
    ws.cell(row=row, column=1).font = Font(bold=True, size=11, color='1F4E79')
    row += 2

    ws.cell(row=row, column=1, value='[GEOMETRY]'); row += 1
    ws.cell(row=row, column=1, value='length_m'); row += 1
    ws.cell(row=row, column=1, value=state['length']); row += 2

    ws.cell(row=row, column=1, value='[SUPPORTS]'); row += 1
    for col, lbl in enumerate(['x_m', 'type'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for s in state['supports']:
        ws.cell(row=row, column=1, value=s['x'])
        ws.cell(row=row, column=2, value=s['type'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[POINT_LOADS]'); row += 1
    for col, lbl in enumerate(['x_m', 'P_kN'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for p in state['point_loads']:
        ws.cell(row=row, column=1, value=p['x'])
        ws.cell(row=row, column=2, value=p['P'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[MOMENTS]'); row += 1
    for col, lbl in enumerate(['x_m', 'M_kNm'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for mm in state['moments']:
        ws.cell(row=row, column=1, value=mm['x'])
        ws.cell(row=row, column=2, value=mm['M'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[DISTRIBUTED_LOADS]'); row += 1
    for col, lbl in enumerate(['x1_m', 'x2_m', 'w1_kNm', 'w2_kNm'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for d in state['dloads']:
        ws.cell(row=row, column=1, value=d['x1'])
        ws.cell(row=row, column=2, value=d['x2'])
        ws.cell(row=row, column=3, value=d['w1'])
        ws.cell(row=row, column=4, value=d['w2'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[NONUNIFORM_LOADS]'); row += 1
    for col, lbl in enumerate(['expr', 'x1_m', 'x2_m'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for d in state['nonuniform_loads']:
        ws.cell(row=row, column=1, value=d['expr'])
        ws.cell(row=row, column=2, value=d['x1'])
        ws.cell(row=row, column=3, value=d['x2'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[SECTION]'); row += 1
    keys = ['E', 'I', 'c', 'A', 'allow_bend', 'allow_shear']
    labels = ['E_GPa', 'I_cm4', 'c_cm', 'A_cm2', 'allow_bend_kNcm2', 'allow_shear_kNcm2']
    for col, lbl in enumerate(labels, 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for col, key in enumerate(keys, 1):
        ws.cell(row=row, column=col, value=state['profile'][key])
    row += 1

    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 14

    if result is not None and model is not None:
        wr = wb.create_sheet('Results')
        wr.cell(row=1, column=1, value='BEAM RESULTS').font = Font(bold=True, size=11, color='1F4E79')

        row_r = 3
        for s in model.supports:
            Fy, Mr = result.reaction_at(s['x'])
            if s['x'] < 1e-9:
                Mr = -Mr
            wr.cell(row=row_r, column=1, value=f"x={s['x']:.3f} ({s['type']})")
            wr.cell(row=row_r, column=2, value='Ry_kN'); wr.cell(row=row_r, column=3, value=Fy / 1e3)
            wr.cell(row=row_r, column=4, value='M_kNm'); wr.cell(row=row_r, column=5, value=Mr / 1e3)
            row_r += 1

        diag = result.sample_diagram(n_per_element=25)
        Vmax = max(diag['V'], key=abs) if diag['V'] else 0.0
        Mmax = max(diag['M'], key=abs) if diag['M'] else 0.0
        vmax = max(diag['v'], key=abs) if diag['v'] else 0.0
        row_r += 1
        wr.cell(row=row_r, column=1, value='Max |V| (kN)'); wr.cell(row=row_r, column=2, value=abs(Vmax) / 1e3)
        row_r += 1
        wr.cell(row=row_r, column=1, value='Max |M| (kN·m)'); wr.cell(row=row_r, column=2, value=abs(Mmax) / 1e3)
        row_r += 1
        wr.cell(row=row_r, column=1, value='Max |defl| (mm)'); wr.cell(row=row_r, column=2, value=abs(vmax) * 1000)
        row_r += 1

        hdr_row = row_r + 2
        for col, lbl in enumerate(['x_m', 'V_kN', 'M_kNm', 'defl_mm'], 1):
            wr.cell(row=hdr_row, column=col, value=lbl)
        for i in range(len(diag['x'])):
            rr = hdr_row + 1 + i
            wr.cell(row=rr, column=1, value=diag['x'][i])
            wr.cell(row=rr, column=2, value=diag['V'][i] / 1e3)
            wr.cell(row=rr, column=3, value=diag['M'][i] / 1e3)
            wr.cell(row=rr, column=4, value=diag['v'][i] * 1000)
        for col in range(1, 5):
            wr.column_dimensions[get_column_letter(col)].width = 13

    wb.save(path)


def import_beam_excel(path):
    """Reads a 'Model' sheet written by export_beam_excel() and returns a
    plain state dict (same shape as the `state` argument passed into
    export_beam_excel), or raises ValueError with a human-readable message."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if 'Model' not in wb.sheetnames:
        raise ValueError('This workbook has no "Model" sheet — it was not '
                          'exported by the Beam tab.')
    ws = wb['Model']
    rows = list(ws.iter_rows(values_only=True))

    def find_section(name):
        for i, r in enumerate(rows):
            if r and r[0] == name:
                return i
        return -1

    def read_table(start_idx):
        hdr_i = start_idx + 1
        if hdr_i >= len(rows) or not rows[hdr_i] or rows[hdr_i][0] is None:
            return []
        headers = [h for h in rows[hdr_i] if h is not None]
        out = []
        i = hdr_i + 1
        while i < len(rows) and rows[i] and rows[i][0] is not None and not str(rows[i][0]).startswith('['):
            vals = rows[i]
            out.append({headers[j]: vals[j] for j in range(len(headers))})
            i += 1
        return out

    gi = find_section('[GEOMETRY]')
    if gi < 0:
        raise ValueError('Missing [GEOMETRY] section.')
    grow = read_table(gi)
    if not grow:
        raise ValueError('Empty [GEOMETRY] section.')
    length = float(grow[0]['length_m'])

    supports = []
    si = find_section('[SUPPORTS]')
    if si >= 0:
        for r in read_table(si):
            supports.append({'x': float(r['x_m']), 'type': str(r['type'])})

    point_loads = []
    pi = find_section('[POINT_LOADS]')
    if pi >= 0:
        for r in read_table(pi):
            point_loads.append({'x': float(r['x_m']), 'P': float(r['P_kN'])})

    moments = []
    mi = find_section('[MOMENTS]')
    if mi >= 0:
        for r in read_table(mi):
            moments.append({'x': float(r['x_m']), 'M': float(r['M_kNm'])})

    dloads = []
    # Accept the pre-2026-09-07 name too, so workbooks already on disk
    # still import. New exports use [DISTRIBUTED_LOADS], matching Cable.
    di = find_section('[DISTRIBUTED_LOADS]')
    if di < 0:
        di = find_section('[DLOADS]')
    if di >= 0:
        for r in read_table(di):
            dloads.append({'x1': float(r['x1_m']), 'x2': float(r['x2_m']),
                            'w1': float(r['w1_kNm']), 'w2': float(r['w2_kNm'])})

    nonuniform_loads = []
    ni = find_section('[NONUNIFORM_LOADS]')
    if ni >= 0:
        for r in read_table(ni):
            nonuniform_loads.append({'expr': str(r['expr']), 'x1': float(r['x1_m']), 'x2': float(r['x2_m'])})

    profile = {'E': 200.0, 'I': 8000.0, 'c': 15.0, 'A': 80.0,
               'allow_bend': 16.0, 'allow_shear': 10.0}
    seci = find_section('[SECTION]')
    if seci >= 0:
        srow = read_table(seci)
        if srow:
            s0 = srow[0]
            profile['E'] = float(s0.get('E_GPa') or profile['E'])
            profile['I'] = float(s0.get('I_cm4') or profile['I'])
            profile['c'] = float(s0.get('c_cm') or profile['c'])
            profile['A'] = float(s0.get('A_cm2') or profile['A'])
            profile['allow_bend'] = float(s0.get('allow_bend_kNcm2') or profile['allow_bend'])
            profile['allow_shear'] = float(s0.get('allow_shear_kNcm2') or profile['allow_shear'])

    return {'length': length, 'supports': supports, 'point_loads': point_loads,
            'moments': moments, 'dloads': dloads, 'nonuniform_loads': nonuniform_loads,
            'profile': profile}

class BeamApp(tk.Frame):
    """
    Beam tab — isostatic and hyperstatic beams, single and double cantilevers
    (overhangs), via the BeamModel direct-stiffness solver above.
    Units: lengths in m, forces in kN, moments in kN*m, distributed loads in
    kN/m, E in GPa, section I in cm^4, c (extreme fiber distance) in cm, areas
    in cm^2. Stresses reported in kN/cm^2 (matches the rest of the app).
    """
    CBEAM, CSUP, CLOAD, CMOM, CDLOAD = '#333333', '#555555', '#D85A30', '#8e44ad', '#c0785a'
    CV_, CM_, CDEFL, CGRID = '#7F77DD', '#D85A30', '#1D9E75', '#e8e8e8'

    def __init__(self, master, **kw):
        super().__init__(master, bg='#f5f5f3', **kw)
        self.length = 6.0
        self.supports = []      # {'x','type'}
        self.point_loads = []   # {'x','P'}   kN, +down
        self.moments = []       # {'x','M'}   kN*m, +CCW
        self.dloads = []        # {'x1','x2','w1','w2'}  kN/m, +down
        self.nonuniform_loads = []  # {'expr','x1','x2'}  kN/m, +down, over [x1,x2] m
        self.profile = {'E': 200.0, 'I': 8000.0, 'c': 15.0, 'A': 80.0,
                        'allow_bend': 16.0, 'allow_shear': 10.0}
        self.result = None
        self.model = None
        self._build_ui()
        self._draw_schematic()
        # Nothing in the model changes when the convention does -- only how it
        # is written -- so the whole tab is simply repainted. Held as an
        # attribute so the listener can be removed if this tab is ever
        # destroyed and rebuilt, which would otherwise leave a dead callback
        # repainting a widget that no longer exists.
        self._units_listener = units.on_change(lambda _sys: self._on_units_changed())


    # ── UI ────────────────────────────────────────────────────────────────────
    def _build_ui(self):
        tb = tk.Frame(self, bg='#ebebea')
        tb.pack(fill='x', padx=6, pady=(6, 0))
        tk.Label(tb, text=f'Beam length ({units.label("length")}):', bg='#ebebea', font=('Helvetica', 11)).pack(side='left', padx=(4, 2))
        self.len_var = tk.DoubleVar(value=self._shown('x', self.length))
        tk.Entry(tb, textvariable=self.len_var, width=7, font=('Helvetica', 11)).pack(side='left')
        tk.Button(tb, text='Set length', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._set_length).pack(side='left', padx=4)
        tk.Frame(tb, width=1, bg='#ccc').pack(side='left', fill='y', padx=6, pady=3)
        tk.Button(tb, text='Example: cantilever', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._load_example_cantilever).pack(side='left', padx=2)
        tk.Button(tb, text='Example: double overhang', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._load_example_overhang).pack(side='left', padx=2)
        tk.Button(tb, text='Example: continuous (hyperstatic)', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._load_example_continuous).pack(side='left', padx=2)
        tk.Button(tb, text='Clear', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), command=self._clear_all).pack(side='left', padx=2)
        tk.Frame(tb, width=1, bg='#ccc').pack(side='left', fill='y', padx=6, pady=3)
        tk.Button(tb, text='Export Excel', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), fg='#1a6bbd', command=self._export_excel).pack(side='left', padx=2)
        tk.Button(tb, text='Import Excel', relief='flat', bd=0, padx=8, pady=4,
                  font=('Helvetica', 11), fg='#1a6bbd', command=self._import_excel).pack(side='left', padx=2)
        tk.Frame(tb, width=1, bg='#ccc').pack(side='left', fill='y', padx=6, pady=3)
        tk.Button(tb, text='▶ Analyze', relief='flat', bd=0, padx=10, pady=4,
                  bg='#1a6bbd', fg='white', font=('Helvetica', 11, 'bold'),
                  command=self._analyze).pack(side='left', padx=2)

        main = tk.Frame(self, bg='#f5f5f3')
        main.pack(fill='both', expand=True, padx=6, pady=6)

        # Right panel FIRST, expanding content SECOND. Tk's pack hands each
        # slave a parcel in packing order, so the previous order (content
        # first, expand=True) left the panel whatever the content did not
        # want -- which at narrow widths was nothing, and the panel was
        # unmapped entirely with no scrollbar and no error. ScrollPanel also
        # scrolls horizontally, so a row wider than the panel stays reachable
        # instead of being clipped mid-widget.
        self.panel_outer = ScrollPanel(main, width=PANEL_W + 105, bg='#f0f0ee',
                                        bd=1, relief='solid')
        self.panel_outer.pack(side='right', fill='y', padx=(6, 0))

        left = tk.Frame(main, bg='#f5f5f3')
        left.pack(side='left', fill='both', expand=True)

        tk.Label(left, text='Beam schematic', bg='#f5f5f3', font=('Helvetica', 9, 'bold'), fg='#777').pack(anchor='w')
        self.schem = tk.Canvas(left, bg='white', height=160, bd=1, relief='solid', highlightthickness=0)
        self.schem.pack(fill='x', pady=(0, 8))
        self.schem.bind('<Configure>', lambda e: self._draw_schematic())

        diag_hdr = tk.Frame(left, bg='#f5f5f3')
        diag_hdr.pack(fill='x')
        tk.Label(diag_hdr, text='Diagrams — shear V, moment M, deflection (run ▶ Analyze)', bg='#f5f5f3',
                 font=('Helvetica', 9, 'bold'), fg='#777').pack(side='left')
        self.reverse_bmd_var = tk.BooleanVar(value=False)
        tk.Checkbutton(diag_hdr, text='Reverse BMD (sagging down / hogging up)',
                       variable=self.reverse_bmd_var, bg='#f5f5f3', font=('Helvetica', 8),
                       command=self._draw_diagrams).pack(side='left', padx=(14, 0))
        self.diag_canvas = tk.Canvas(left, bg='#fafaf8', height=420, bd=1, relief='solid', highlightthickness=0)
        self.diag_canvas.pack(fill='both', expand=True)
        self.diag_canvas.bind('<Configure>', lambda e: self._draw_diagrams())

        self._build_panel(self.panel_outer.interior)
        # Adopt whatever width the panel's own content needs, so nothing
        # starts life behind the horizontal scrollbar.
        self.panel_outer.fit_to_content()

        # The toolbar was one long row of pack(side='left') calls, so its tail
        # ran off the right edge. WrapBar flows those same widgets across as
        # many rows as the width needs, without restructuring how they were
        # built. The same <Configure> drives the panel width, so the two can
        # never disagree about how wide the window currently is.
        self.toolbar_wrap = WrapBar(tb)
        self.toolbar_wrap.start()
        self.bind('<Configure>', self._on_root_configure, add='+')
        self.after_idle(lambda: self._on_root_configure(None))

    def _on_root_configure(self, _event=None):
        """Resize the right panel to match the window. Content that no longer
        fits stays reachable through ScrollPanel's horizontal scrollbar, so
        this can never hide a control -- unlike the previous fixed-width
        panel, which was simply dropped."""
        try:
            self.panel_outer.apply_responsive_width(self.winfo_width())
        except Exception:
            pass

    def _build_panel(self, panel):
        pad = dict(padx=8, pady=(8, 2))

        tk.Label(panel, text='SUPPORTS', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        self.sup_tree = ttk.Treeview(panel, columns=('x', 'type'), show='headings', height=4)
        self.sup_tree.heading('x', text=f'x ({self._u("x")})'); self.sup_tree.column('x', width=70)
        self.sup_tree.heading('type', text='Type'); self.sup_tree.column('type', width=100)
        self.sup_tree.pack(fill='x', padx=8)
        sf = tk.Frame(panel, bg='#f0f0ee'); sf.pack(fill='x', padx=8, pady=(2, 8))
        tk.Button(sf, text='Add', command=self._add_support).pack(side='left', padx=2)
        tk.Button(sf, text='Delete', command=lambda: self._del_row(self.sup_tree, self.supports)).pack(side='left', padx=2)

        tk.Label(panel, text='POINT LOADS (+down, kN)', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        self.pl_tree = ttk.Treeview(panel, columns=('x', 'P'), show='headings', height=3)
        self.pl_tree.heading('x', text=f'x ({self._u("x")})'); self.pl_tree.column('x', width=70)
        self.pl_tree.heading('P', text=f'P ({self._u("P")})'); self.pl_tree.column('P', width=100)
        self.pl_tree.pack(fill='x', padx=8)
        pf = tk.Frame(panel, bg='#f0f0ee'); pf.pack(fill='x', padx=8, pady=(2, 8))
        tk.Button(pf, text='Add', command=self._add_pointload).pack(side='left', padx=2)
        tk.Button(pf, text='Delete', command=lambda: self._del_row(self.pl_tree, self.point_loads)).pack(side='left', padx=2)

        tk.Label(panel, text='POINT MOMENTS (+CCW, kN·m)', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        self.mm_tree = ttk.Treeview(panel, columns=('x', 'M'), show='headings', height=2)
        self.mm_tree.heading('x', text=f'x ({self._u("x")})'); self.mm_tree.column('x', width=70)
        self.mm_tree.heading('M', text=f'M ({self._u("M")})'); self.mm_tree.column('M', width=100)
        self.mm_tree.pack(fill='x', padx=8)
        mf = tk.Frame(panel, bg='#f0f0ee'); mf.pack(fill='x', padx=8, pady=(2, 8))
        tk.Button(mf, text='Add', command=self._add_moment).pack(side='left', padx=2)
        tk.Button(mf, text='Delete', command=lambda: self._del_row(self.mm_tree, self.moments)).pack(side='left', padx=2)

        tk.Label(panel, text='DISTRIBUTED LOADS (+down, kN/m)', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        self.dl_tree = ttk.Treeview(panel, columns=('x1', 'x2', 'w1', 'w2'), show='headings', height=3)
        for c, w in [('x1', 55), ('x2', 55), ('w1', 60), ('w2', 60)]:
            self.dl_tree.heading(c, text=c); self.dl_tree.column(c, width=w)
        self.dl_tree.pack(fill='x', padx=8)
        df = tk.Frame(panel, bg='#f0f0ee'); df.pack(fill='x', padx=8, pady=(2, 8))
        tk.Button(df, text='Add', command=self._add_dload).pack(side='left', padx=2)
        tk.Button(df, text='Delete', command=lambda: self._del_row(self.dl_tree, self.dloads)).pack(side='left', padx=2)

        tk.Label(panel, text='NON-UNIFORM DISTRIBUTED LOADS', bg='#f0f0ee',
                 font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        tk.Label(panel, text='q(x) in kN/m (+down), over a sub-domain [x₁,x₂] (m)', bg='#f0f0ee',
                 font=('Helvetica', 8), fg='#777').pack(anchor='w', padx=8)
        self.ndl_tree = ttk.Treeview(panel, columns=('expr', 'x1', 'x2'), show='headings', height=3)
        for c, w, lbl in [('expr', 130, 'q(x)'), ('x1', 45, 'x₁'), ('x2', 45, 'x₂')]:
            self.ndl_tree.heading(c, text=lbl); self.ndl_tree.column(c, width=w)
        self.ndl_tree.pack(fill='x', padx=8)
        ndf = tk.Frame(panel, bg='#f0f0ee'); ndf.pack(fill='x', padx=8, pady=(2, 8))
        tk.Button(ndf, text='Add', command=self._add_nonuniform_load).pack(side='left', padx=2)
        tk.Button(ndf, text='Delete',
                  command=lambda: self._del_row(self.ndl_tree, self.nonuniform_loads)).pack(side='left', padx=2)

        tk.Label(panel, text='CROSS-SECTION / MATERIAL', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        sec = tk.Frame(panel, bg='#f0f0ee'); sec.pack(fill='x', padx=8)
        self.sec_vars = {}
        def _sec_fields():
            u = units.label
            return [('E', f'E ({u("modulus")})'), ('I', f'I ({u("inertia")})'),
                    ('c', f'c ({u("section_length")}, extreme fiber)'),
                    ('A', f'Shear area ({u("area")})'),
                    ('allow_bend', f'Allow. bending σ ({u("stress")})'),
                    ('allow_shear', f'Allow. shear τ ({u("stress")})')]

        self._sec_fields = _sec_fields
        self._sec_labels = {}
        for i, (key, label) in enumerate(_sec_fields()):
            lb = tk.Label(sec, text=label, bg='#f0f0ee', font=('Helvetica', 9))
            lb.grid(row=i, column=0, sticky='w', pady=1)
            self._sec_labels[key] = lb
            v = tk.DoubleVar(value=self._sec_shown(key, self.profile[key]))
            self.sec_vars[key] = v
            tk.Entry(sec, textvariable=v, width=8, font=('Helvetica', 9)).grid(row=i, column=1, pady=1, padx=4)

        tk.Label(panel, text='RESULTS', bg='#f0f0ee', font=('Helvetica', 10, 'bold')).pack(anchor='w', **pad)
        self.res_text = tk.Text(panel, height=16, bg='white', font=('Courier', 9), relief='solid', bd=1)
        self.res_text.pack(fill='both', expand=True, padx=8, pady=(2, 10))

    # ── row management ──────────────────────────────────────────────────────
    def _del_row(self, tree, data_list):
        sel = tree.selection()
        if not sel:
            return
        idx = tree.index(sel[0])
        del data_list[idx]
        self._refresh_tables()

    # Which physical quantity each model field is, so one table-refresh can
    # convert every column without a per-column special case. Fields that are
    # not numbers (a support type, a q(x) expression) map to None and are
    # passed through untouched.
    _FIELD_Q = {'x': 'length', 'x1': 'length', 'x2': 'length',
                'P': 'force', 'M': 'moment',
                'w1': 'line_load', 'w2': 'line_load',
                'type': None, 'expr': None}

    _SECTION_Q = {'E': 'modulus', 'I': 'inertia', 'c': 'section_length',
                  'A': 'area', 'allow_bend': 'stress', 'allow_shear': 'stress'}

    def _shown(self, field, stored):
        """A stored model value as the current unit convention writes it."""
        q = self._FIELD_Q.get(field)
        if q is None or not isinstance(stored, (int, float)):
            return stored
        return units.to_display(q, stored)

    def _stored(self, field, shown):
        """The inverse of `_shown`, for a number the user typed."""
        q = self._FIELD_Q.get(field)
        if q is None or not isinstance(shown, (int, float)):
            return shown
        return units.from_display(q, shown)

    def _on_units_changed(self):
        """Repaint every place a unit is written or a number is shown."""
        try:
            if not self.winfo_exists():
                return
        except Exception:
            return
        self.len_var.set(self._shown('x', self.length))
        for key, lb in getattr(self, '_sec_labels', {}).items():
            lb.config(text=dict(self._sec_fields())[key])
            self.sec_vars[key].set(self._sec_shown(key, self.profile[key]))
        self.sup_tree.heading('x', text=f'x ({self._u("x")})')
        self.pl_tree.heading('x', text=f'x ({self._u("x")})')
        self.pl_tree.heading('P', text=f'P ({self._u("P")})')
        self.mm_tree.heading('x', text=f'x ({self._u("x")})')
        self.mm_tree.heading('M', text=f'M ({self._u("M")})')
        self._refresh_tables()
        if self.result is not None:
            self._show_results()
            self._draw_diagrams()

    def _sec_shown(self, key, stored):
        return units.to_display(self._SECTION_Q[key], stored)

    def _sec_stored(self, key, shown):
        return units.from_display(self._SECTION_Q[key], shown)

    def _u(self, field):
        """Label for the unit a field is currently shown in."""
        q = self._FIELD_Q.get(field)
        return units.label(q) if q else ''

    def _refresh_tables(self):
        for tree, rows, cols in [
            (self.sup_tree, self.supports, ('x', 'type')),
            (self.pl_tree, self.point_loads, ('x', 'P')),
            (self.mm_tree, self.moments, ('x', 'M')),
            (self.dl_tree, self.dloads, ('x1', 'x2', 'w1', 'w2')),
            (self.ndl_tree, self.nonuniform_loads, ('expr', 'x1', 'x2')),
        ]:
            tree.delete(*tree.get_children())
            for r in rows:
                vals = []
                for c in cols:
                    v = self._shown(c, r[c])
                    vals.append(f'{v:g}' if isinstance(v, float) else v)
                tree.insert('', 'end', values=tuple(vals))
        self._draw_schematic()

    def _ask(self, title, fields):
        win = tk.Toplevel(self); win.title(title); win.grab_set()
        win.configure(bg='#f0f0ee')
        vars_ = {}
        for i, (key, label, default) in enumerate(fields):
            tk.Label(win, text=label, bg='#f0f0ee').grid(row=i, column=0, sticky='w', padx=8, pady=4)
            v = tk.DoubleVar(value=default) if not isinstance(default, str) else tk.StringVar(value=default)
            vars_[key] = v
            if isinstance(default, str):
                cb = ttk.Combobox(win, textvariable=v, values=['pin', 'roller', 'fixed', 'guided'],
                                   width=10, state='readonly')
                cb.grid(row=i, column=1, padx=8, pady=4)
            else:
                tk.Entry(win, textvariable=v, width=10).grid(row=i, column=1, padx=8, pady=4)
        result = {}

        def ok():
            for k, v in vars_.items():
                result[k] = v.get()
            win.destroy()
        tk.Button(win, text='OK', command=ok, bg='#1a6bbd', fg='white').grid(
            row=len(fields), column=0, columnspan=2, pady=8)
        win.wait_window()
        return result if result else None

    def _on_beam(self, r, *keys):
        """True if every named coordinate in a dialog result is on the beam.

        The dialogs took whatever was typed and appended it unchecked, so a
        mistyped station silently changed the structure being analysed
        (finding B-5). BeamModel refuses these too; catching it here means the
        user is told while the number is still in front of them, rather than
        at Analyze.
        """
        if not r:
            return False
        for k in keys:
            x = r.get(k)
            if x is None:
                continue
            if not (0.0 <= x <= self.length):
                messagebox.showwarning(
                    'Off the beam',
                    f'{k} = {x:g} m is not on the beam, which spans 0 to '
                    f'{self.length:g} m.\n\nNothing was added. Move it onto the '
                    f'beam, or set the beam length first.')
                return False
        return True

    def _add_support(self):
        r = self._ask('Add support',
                      [('x', f'x ({self._u("x")})',
                        self._shown('x', self.length / 2)), ('type', 'Type', 'pin')])
        if r:
            r['x'] = self._stored('x', r['x'])
        if self._on_beam(r, 'x'): self.supports.append(r); self._refresh_tables()

    def _add_pointload(self):
        r = self._ask('Add point load',
                      [('x', f'x ({self._u("x")})', self._shown('x', self.length / 2)),
                       ('P', f'P ({self._u("P")}, +down)', self._shown('P', 10.0))])
        if r:
            r['x'] = self._stored('x', r['x']); r['P'] = self._stored('P', r['P'])
        if self._on_beam(r, 'x'): self.point_loads.append(r); self._refresh_tables()

    def _add_moment(self):
        r = self._ask('Add point moment',
                      [('x', f'x ({self._u("x")})', self._shown('x', self.length / 2)),
                       ('M', f'M ({self._u("M")}, +CCW)', self._shown('M', 10.0))])
        if r:
            r['x'] = self._stored('x', r['x']); r['M'] = self._stored('M', r['M'])
        if self._on_beam(r, 'x'): self.moments.append(r); self._refresh_tables()

    def _add_dload(self):
        r = self._ask('Add distributed load', [
            ('x1', f'x1 ({self._u("x1")})', 0.0),
            ('x2', f'x2 ({self._u("x2")})', self._shown('x2', self.length)),
            ('w1', f'w1 ({self._u("w1")}, +down)', self._shown('w1', 5.0)),
            ('w2', f'w2 ({self._u("w2")}, +down)', self._shown('w2', 5.0))])
        if r:
            for k in ('x1', 'x2', 'w1', 'w2'):
                r[k] = self._stored(k, r[k])
        if self._on_beam(r, 'x1', 'x2'):
            if r['x2'] < r['x1']:      # entered right-to-left; the intensities
                r = dict(r, x1=r['x2'], x2=r['x1'], w1=r['w2'], w2=r['w1'])
            self.dloads.append(r); self._refresh_tables()

    def _add_nonuniform_load(self):
        win = tk.Toplevel(self); win.title('Add non-uniform distributed load'); win.grab_set()
        win.configure(bg='#f0f0ee')
        tk.Label(win, text='q(x) in kN/m, +down  (vars: x, L; ^ or ** = power):', bg='#f0f0ee',
                 font=('Helvetica', 9)).grid(row=0, column=0, columnspan=2, sticky='w', padx=8, pady=(8, 2))
        expr_var = tk.StringVar(value='10*sin(pi*x/L)')
        tk.Entry(win, textvariable=expr_var, width=28, font=('Helvetica', 9)).grid(
            row=1, column=0, columnspan=2, sticky='we', padx=8, pady=2)

        tk.Label(win, text='x₁ (m):', bg='#f0f0ee', font=('Helvetica', 9)).grid(
            row=2, column=0, sticky='w', padx=8, pady=4)
        x1_var = tk.DoubleVar(value=0.0)
        tk.Entry(win, textvariable=x1_var, width=10).grid(row=2, column=1, padx=8, pady=4)

        tk.Label(win, text='x₂ (m):', bg='#f0f0ee', font=('Helvetica', 9)).grid(
            row=3, column=0, sticky='w', padx=8, pady=4)
        x2_var = tk.DoubleVar(value=self.length)
        tk.Entry(win, textvariable=x2_var, width=10).grid(row=3, column=1, padx=8, pady=4)

        result = {}

        def ok():
            expr = expr_var.get().strip()
            try:
                ctx = {'L': self._stored('x', self.len_var.get())}
                x1_, x2_ = x1_var.get(), x2_var.get()
                x_mid = (min(x1_, x2_) + max(x1_, x2_)) / 2
                make_shape_fn(expr, ctx)(x_mid)   # validate it compiles & evaluates
            except Exception as e:
                messagebox.showerror('Invalid expression', str(e)); return
            result['expr'] = expr
            result['x1'] = x1_var.get()
            result['x2'] = x2_var.get()
            win.destroy()

        tk.Button(win, text='OK', command=ok, bg='#1a6bbd', fg='white').grid(
            row=4, column=0, columnspan=2, pady=8)
        win.wait_window()
        if result:
            self.nonuniform_loads.append(result)
            self._refresh_tables()

    def _set_length(self):
        self.length = self._stored('x', self.len_var.get())
        self._draw_schematic()

    def _clear_all(self):
        self.supports = []; self.point_loads = []; self.moments = []; self.dloads = []
        self.nonuniform_loads = []
        self.result = None; self.model = None
        self._refresh_tables()
        self.res_text.delete('1.0', 'end')
        self.diag_canvas.delete('all')

    # ── Excel export / import ────────────────────────────────────────────────
    def _current_state(self):
        return {
            'length': self._stored('x', self.len_var.get()),
            'supports': self.supports, 'point_loads': self.point_loads,
            'moments': self.moments, 'dloads': self.dloads,
            'nonuniform_loads': self.nonuniform_loads,
            'profile': {k: self._sec_stored(k, v.get())
                        for k, v in self.sec_vars.items()},
        }

    def _export_excel(self):
        if not _ensure_openpyxl():
            messagebox.showerror(
                'Missing library',
                'Could not install openpyxl automatically.\n\n'
                'Please open a terminal and run:\n'
                '    pip install openpyxl\n'
                'then try again.')
            return
        path = filedialog.asksaveasfilename(
            defaultextension='.xlsx',
            filetypes=[('Excel workbook', '*.xlsx')],
            initialfile='beam_report.xlsx',
            title='Save Excel report')
        if not path:
            return
        try:
            export_beam_excel(self._current_state(), path,
                               result=self.result, model=self.model)
            messagebox.showinfo(
                'Exported', f'Report saved to:\n{path}\n\n'
                'Includes a "Model" sheet — use "Import" to rebuild this '
                'exact beam from the file later.' +
                ('' if self.result else '\n\n(Run ▶ Analyze first to also '
                                         'include a "Results" sheet next time.)'))
        except Exception as e:
            messagebox.showerror('Export failed', str(e))

    def _import_excel(self):
        if not _ensure_openpyxl():
            messagebox.showerror(
                'Missing library',
                'Could not install openpyxl automatically.\n\n'
                'Please open a terminal and run:\n'
                '    pip install openpyxl\n'
                'then try again.')
            return
        path = filedialog.askopenfilename(
            filetypes=[('Excel workbook', '*.xlsx')],
            title='Import beam from Excel')
        if not path:
            return
        try:
            st = import_beam_excel(path)
        except Exception as e:
            messagebox.showerror('Import failed', str(e)); return

        self._clear_all()
        self.length = st['length']; self.len_var.set(self._shown('x', st['length']))
        self.supports = st['supports']
        self.point_loads = st['point_loads']
        self.moments = st['moments']
        self.dloads = st['dloads']
        self.nonuniform_loads = st['nonuniform_loads']
        for k, v in st['profile'].items():
            if k in self.sec_vars:
                self.sec_vars[k].set(self._sec_shown(k, v))
                self.profile[k] = v
        self._refresh_tables()
        self._draw_schematic()

    # ── examples ─────────────────────────────────────────────────────────────
    def _load_example_cantilever(self):
        self._clear_all()
        self.length = 4.0; self.len_var.set(self._shown('x', 4.0))
        self.supports = [{'x': 0.0, 'type': 'fixed'}]
        self.point_loads = [{'x': 4.0, 'P': 15.0}]
        self._refresh_tables()

    def _load_example_overhang(self):
        self._clear_all()
        self.length = 10.0; self.len_var.set(self._shown('x', 10.0))
        self.supports = [{'x': 2.0, 'type': 'pin'}, {'x': 8.0, 'type': 'roller'}]
        self.dloads = [{'x1': 0.0, 'x2': 10.0, 'w1': 6.0, 'w2': 6.0}]
        self._refresh_tables()

    def _load_example_continuous(self):
        self._clear_all()
        self.length = 10.0; self.len_var.set(self._shown('x', 10.0))
        self.supports = [{'x': 0.0, 'type': 'pin'}, {'x': 5.0, 'type': 'roller'},
                         {'x': 10.0, 'type': 'roller'}]
        self.dloads = [{'x1': 0.0, 'x2': 10.0, 'w1': 8.0, 'w2': 8.0}]
        self._refresh_tables()

    # ── analysis ─────────────────────────────────────────────────────────────
    def _analyze(self):
        try:
            self.length = self._stored('x', self.len_var.get())
            if not self.supports:
                messagebox.showwarning('Analyze', 'Add at least one support.'); return
            for k in self.sec_vars:
                self.profile[k] = self._sec_stored(k, self.sec_vars[k].get())

            E_Pa = self.profile['E'] * 1e9
            I_m4 = self.profile['I'] * 1e-8
            EI = E_Pa * I_m4

            m = BeamModel(self.length)
            m.EI = EI
            for s in self.supports:
                m.add_support(s['x'], s['type'])
            for p in self.point_loads:
                m.add_point_load(p['x'], p['P'] * 1e3)
            for mm in self.moments:
                m.add_moment(mm['x'], mm['M'] * 1e3)
            for d in self.dloads:
                m.add_dload(d['x1'], d['x2'], d['w1'] * 1e3, d['w2'] * 1e3)
            for d in self.nonuniform_loads:
                ctx = {'L': self.length}
                qfn = make_shape_fn(d['expr'], ctx)
                x1 = max(0.0, min(self.length, d['x1']))
                x2 = max(0.0, min(self.length, d['x2']))
                if x2 < x1:
                    x1, x2 = x2, x1

                def scaled_fn(x, _qfn=qfn):
                    return _qfn(x) * 1e3   # kN/m -> N/m

                m.add_nonuniform_load(scaled_fn, x1, x2)

            self.result = m.solve()
            self.model = m
            self._draw_diagrams()
            self._show_results()
            self._draw_schematic()
        except Exception as e:
            messagebox.showerror('Analysis failed', str(e))

    def _show_results(self):
        r = self.result; m = self.model
        diag = r.sample_diagram(n_per_element=25)
        Vmax = max(diag['V'], key=abs)
        Mmax = max(diag['M'], key=abs)
        vmax = max(diag['v'], key=abs)

        c_m = self.profile['c'] * 1e-2
        I_m4 = self.profile['I'] * 1e-8
        A_m2 = self.profile['A'] * 1e-4
        sigma_kncm2 = (abs(Mmax) * c_m / I_m4) * 1e-7
        tau_kncm2 = (abs(Vmax) / A_m2) * 1e-7

        lines = ['REACTIONS']
        for s in m.supports:
            Fy, Mr = r.reaction_at(s['x'])
            if s['x'] < 1e-9:
                # The leftmost node's rotational-DOF residual comes out with
                # the opposite sign from the sagging-positive convention used
                # by the M(x) diagram at that same station (a direct-
                # stiffness-method artifact of it always being the "a-side"
                # of its adjacent element) — negate here so this value is
                # directly comparable to the moment diagram.
                Mr = -Mr
            # r.reaction_at / the diagram are in SI; the model's own state is
            # in STORAGE units. Both are written in the convention the user
            # picked, which is the only thing the selector changes.
            lines.append(
                f"  x={self._shown('x', s['x']):.2f} {self._u('x')} "
                f"({s['type']:<7}): "
                f"Ry={units.from_si('force', Fy):+8.2f} {units.label('force')}   "
                f"M={units.from_si('moment', Mr):+8.2f} {units.label('moment')}")
        lines += ['',
                  f"Max |V|  = {units.from_si('force', abs(Vmax)):8.2f} {units.label('force')}",
                  f"Max |M|  = {units.from_si('moment', abs(Mmax)):8.2f} {units.label('moment')}",
                  f"Max |defl| = {units.from_si('deflection', abs(vmax)):8.3f} "
                  f"{units.label('deflection')}",
                  '', 'STRESS CHECK']
        bend_ratio = sigma_kncm2 / self.profile['allow_bend'] if self.profile['allow_bend'] else 0
        shear_ratio = tau_kncm2 / self.profile['allow_shear'] if self.profile['allow_shear'] else 0
        su = units.label('stress')
        lines.append(f'  Bending sigma = M*c/I = '
                      f'{self._sec_shown("allow_bend", sigma_kncm2):6.3f} {su}')
        lines.append(f'    allowable = '
                      f'{self._sec_shown("allow_bend", self.profile["allow_bend"]):.3f} {su} '
                      f'  ({"OK" if bend_ratio <= 1.0 else "FAIL"}, ratio {bend_ratio:.2f})')
        lines.append(f'  Shear tau = V/A   = '
                      f'{self._sec_shown("allow_shear", tau_kncm2):6.3f} {su}')
        lines.append(f'    allowable = '
                      f'{self._sec_shown("allow_shear", self.profile["allow_shear"]):.3f} {su} '
                      f'  ({"OK" if shear_ratio <= 1.0 else "FAIL"}, ratio {shear_ratio:.2f})')

        self.res_text.delete('1.0', 'end')
        self.res_text.insert('1.0', '\n'.join(lines))

    # ── drawing ──────────────────────────────────────────────────────────────
    def _draw_schematic(self):
        c = self.schem
        c.delete('all')
        w = c.winfo_width() or 600
        h = c.winfo_height() or 160
        margin = 40
        L = max(self.length, 0.001)
        scale = (w - 2 * margin) / L
        y0 = h * 0.55

        def X(x): return margin + x * scale

        c.create_line(X(0), y0, X(L), y0, width=3, fill=self.CBEAM)

        for s in self.supports:
            x = X(s['x']); t = s['type']
            if t in ('pin', 'roller'):
                c.create_polygon(x - 10, y0 + 18, x + 10, y0 + 18, x, y0, fill='', outline=self.CSUP, width=2)
                if t == 'roller':
                    c.create_oval(x - 10, y0 + 18, x - 4, y0 + 24, outline=self.CSUP)
                    c.create_oval(x + 4, y0 + 18, x + 10, y0 + 24, outline=self.CSUP)
            elif t == 'fixed':
                c.create_line(x, y0 - 16, x, y0 + 16, width=3, fill=self.CSUP)
                for k in range(-3, 4):
                    c.create_line(x, y0 + k * 5, x - 8, y0 + k * 5 + 8, fill=self.CSUP)
            elif t == 'guided':
                c.create_rectangle(x - 10, y0 + 2, x + 10, y0 + 14, outline=self.CSUP)
            c.create_text(x, y0 + 34,
                          text=f"{self._shown('x', s['x']):.2f} {self._u('x')}",
                          font=('Helvetica', 8), fill='#555')

        # Distributed loads are drawn TO SCALE against the largest intensity in
        # the model, and the chord above them follows the real q(x). Every
        # arrow used to be a fixed 30 px whatever the load, so a 0 -> 60 kN/m
        # triangular load was drawn exactly like a uniform one and the
        # non-uniform branch computed its own shape function and then threw it
        # away (2026-09-05 finding B-3). The load picture is the check an
        # engineer makes before pressing Analyze; it has to show the shape.
        DL_H = 34.0     # px at the largest intensity in the model
        DL_MIN = 5.0    # px floor, so a tiny ordinate still reads as a load
        N_SAMPLES = 20

        def _q_profile(d):
            """[(x_world, intensity_kNm), ...] for one distributed load, or
            None if a user expression will not evaluate."""
            a, b = min(d['x1'], d['x2']), max(d['x1'], d['x2'])
            if 'expr' in d:
                try:
                    fn = make_shape_fn(d['expr'], {'L': self.length})
                except Exception:
                    return None
                pts = []
                for k in range(N_SAMPLES + 1):
                    # sample strictly inside (a, b): the expression may be
                    # singular exactly at its own domain edge
                    t = (k + 0.5) / (N_SAMPLES + 1)
                    xv = a + (b - a) * t
                    try:
                        pts.append((xv, fn(xv)))
                    except Exception:
                        return None
                return pts
            return [(a + (b - a) * k / N_SAMPLES,
                     d['w1'] + (d['w2'] - d['w1']) * (k / N_SAMPLES))
                    for k in range(N_SAMPLES + 1)]

        profiles = []
        for d in list(self.dloads) + list(self.nonuniform_loads):
            pts = _q_profile(d)
            if pts:
                profiles.append((d, pts))
        # Same compressed relative scale as the point loads and moments above,
        # but normalised against an INTENSITY of its own: two distributed loads
        # of equal intensity must draw at equal height whatever length each
        # covers, which a scale shared with the point loads would break.
        q_scale = LoadScale.of((q for _, pts in profiles for _, q in pts),
                               DL_MIN, DL_H)

        def _height(q):
            # A genuinely zero ordinate sits on the beam line, so the start of a
            # triangular load reads as zero rather than as a small load.
            if abs(q) <= 1e-12:
                return 0.0
            return q_scale(q)

        load_labels = []          # decluttered together at the end of the paint

        # Tallest ordinate actually drawn above the beam. Point-load arrows and
        # moment arcs are laid over this band, so their LABELS are lifted clear
        # of it -- the arrows themselves still run to the beam at their true
        # scaled length, which a shifted arrow would falsify.
        dl_top = max((_height(q) for _, pts in profiles for _, q in pts if q > 0),
                     default=0.0)

        for d, pts in profiles:
            # +q is downward, so it is drawn above the beam; an uplift ordinate
            # hangs below it, and a load that changes sign crosses the beam line.
            tops = [(X(xv), y0 - _height(q) if q >= 0 else y0 + _height(q))
                    for xv, q in pts]
            kw = {'dash': (3, 2)} if 'expr' in d else {}
            c.create_line(*[v for pt in tops for v in pt],
                          fill=self.CDLOAD, width=2, **kw)
            step = max(1, len(tops) // 12)
            for i in range(0, len(tops), step):
                xx, yy = tops[i]
                c.create_line(xx, yy, xx, y0, arrow='last', fill=self.CDLOAD)
            label = (f"q(x) = {d['expr']}" if 'expr' in d
                     else f"{self._shown('w1', d['w1']):.1f}→"
                          f"{self._shown('w2', d['w2']):.1f} {self._u('w1')}")
            load_labels.append(c.create_text(
                (tops[0][0] + tops[-1][0]) / 2,
                min(y0 - DL_H, min(y for _, y in tops)) - 10,
                text=label, font=('Helvetica', 8, 'bold'), fill=self.CDLOAD))

        # Every load glyph is sized RELATIVE to the largest of its own kind on
        # the beam, square-root compressed. See common.LoadScale for why the
        # three families are scaled separately and why the mapping is not
        # linear. Point-load arrows used to be a flat 45 px and moment arcs a
        # flat 12 px radius, so a 500 kN load and a 5 kN load drew identically.
        p_scale = LoadScale.of((p['P'] for p in self.point_loads), 10.0, 48.0)
        m_scale = LoadScale.of((mm['M'] for mm in self.moments), 7.0, 17.0)

        for p in self.point_loads:
            x = X(p['x'])
            h = p_scale(p['P'])
            top = y0 - h if p['P'] >= 0 else y0 + h
            c.create_line(x, top, x, y0, arrow='last', fill=self.CLOAD, width=2)
            label_y = (min(top, y0 - dl_top) - 8 if p['P'] >= 0
                       else max(top, y0 + dl_top) + 8)
            load_labels.append(c.create_text(
                x, label_y,
                text=f"{self._shown('P', p['P']):.1f} {self._u('P')}",
                font=('Helvetica', 8, 'bold'), fill=self.CLOAD))

        for mm in self.moments:
            x = X(mm['x'])
            ccw = mm['M'] >= 0  # sign convention for this tab: +M = CCW
            r = m_scale(mm['M'])
            draw_moment_arrow(c, x, y0, r, ccw, self.CMOM, width=2)
            load_labels.append(c.create_text(
                x, min(y0 - r, y0 - dl_top) - 14,
                text=f"{self._shown('M', mm['M']):.1f} {self._u('M')}",
                font=('Helvetica', 8, 'bold'), fill=self.CMOM))

        # Two loads at the same station put their labels in the same place; lift
        # whichever was drawn later until nothing overlaps.
        declutter_text(c, load_labels)

    def _draw_diagrams(self):
        c = self.diag_canvas
        c.delete('all')
        if not self.result:
            return
        w = c.winfo_width() or 600
        h = c.winfo_height() or 420
        diag = self.result.sample_diagram(n_per_element=25)
        xs = diag['x']
        L = self.length
        margin = 45
        scale_x = (w - 2 * margin) / L
        band_h = (h - 40) / 3

        M_vals = [units.from_si('moment', mv) for mv in diag['M']]
        if self.reverse_bmd_var.get():
            M_vals = [-v for v in M_vals]
            M_label = (f'MOMENT M ({units.label("moment")}) — sagging plotted DOWN,'
                        ' hogging plotted UP')
        else:
            M_label = (f'MOMENT M ({units.label("moment")}) — sagging (+) plotted UP,'
                        ' hogging (−) plotted DOWN')

        bands = [(f'SHEAR V ({units.label("force")}) — positive plotted UP',
                  [units.from_si('force', v) for v in diag['V']], self.CV_, 0),
                 (M_label, M_vals, self.CM_, 1),
                 (f'DEFLECTION ({units.label("deflection")}) — negative = downward',
                  [units.from_si('deflection', vv) for vv in diag['v']], self.CDEFL, 2)]

        last_band_idx = len(bands) - 1
        for label, ys, color, bidx in bands:
            y0 = 20 + bidx * band_h + band_h / 2
            top = 20 + bidx * band_h
            bot = 20 + (bidx + 1) * band_h
            maxabs = max(1e-9, max(abs(v) for v in ys))
            avail = band_h / 2 - 14

            v_ticks = _nice_ticks(-maxabs, maxabs, 5)
            for vt in v_ticks:
                yy = y0 - (vt / maxabs) * avail
                if top + 2 <= yy <= bot - 2:
                    c.create_line(margin, yy, w - 10, yy, fill='#eee')
                    c.create_text(margin - 4, yy, text=f'{vt:g}', anchor='e',
                                  font=('Helvetica', 6), fill='#aaa')
            x_ticks = _nice_ticks(0, L, 8)
            for xt in x_ticks:
                xx = margin + xt * scale_x
                if margin - 1 <= xx <= w - 9:
                    c.create_line(xx, top, xx, bot, fill='#eee')
                    if bidx == last_band_idx:
                        c.create_text(xx, bot + 10, text=f'{xt:g}', anchor='n',
                                      font=('Helvetica', 7), fill='#888')

            c.create_line(margin, top, margin, bot, fill=self.CGRID)
            c.create_line(margin, y0, w - 10, y0, fill='#bbb')
            c.create_text(margin, top + 8, text=label, anchor='w', font=('Helvetica', 8, 'bold'), fill='#555')
            pts = []
            for x, v in zip(xs, ys):
                sx = margin + x * scale_x
                sy = y0 - (v / maxabs) * avail
                pts.append((sx, sy))
            poly = [(margin, y0)] + pts + [(margin + L * scale_x, y0)]
            flat = [c_ for pt in poly for c_ in pt]
            c.create_polygon(flat, fill=color, outline=color, stipple='gray50')
            for i in range(len(pts) - 1):
                c.create_line(*pts[i], *pts[i + 1], fill=color, width=2)
            c.create_text(w - 10, top + 8, text=f"max ±{maxabs:.2f}", anchor='e',
                          font=('Helvetica', 8), fill='#777')

            locs, _ = _find_diagram_maxima(xs, ys)
            for xv in locs:
                idx = min(range(len(xs)), key=lambda i: abs(xs[i] - xv))
                sx, sy = pts[idx]
                c.create_oval(sx - 4, sy - 4, sx + 4, sy + 4, fill='#c0392b', outline='white', width=1)
                label_y = sy - 11 if ys[idx] >= 0 else sy + 11
                c.create_text(sx, label_y, text=f"x={xv:.2f}", font=('Helvetica', 7), fill='#c0392b')
