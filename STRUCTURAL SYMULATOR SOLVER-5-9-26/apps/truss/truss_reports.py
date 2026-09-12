"""Excel reports and static report drawings for the Truss/Vierendeel tab."""
import math, os, sys, subprocess
from common import (PX_PER_M, CT, CC, CZ, CL, CR, _ensure_openpyxl, _ensure_matplotlib, render_math)
from .truss_math import compute_diagrams

def _get_ttf_font(size):
    """A truetype font for PIL text rendering — reuses the DejaVuSans font
    matplotlib already ships with, so no extra font file is needed."""
    try:
        import matplotlib, os
        from PIL import ImageFont
        path = os.path.join(matplotlib.get_data_path(), 'fonts', 'ttf', 'DejaVuSans.ttf')
        return ImageFont.truetype(path, size)
    except Exception:
        from PIL import ImageFont
        return ImageFont.load_default()

def pil_draw_rod_context(nodes, rods, rod_idx, rod_res=None, size=260):
    """
    Pure-PIL (no live tkinter widget needed) rendering of the whole truss in
    light gray with `rod_idx` highlighted — used to embed a "where is this
    member" picture directly into the Excel report.
    """
    from PIL import Image, ImageDraw
    img = Image.new('RGB', (size, size), (255,255,255))
    d = ImageDraw.Draw(img)
    if not nodes: return img
    xs = [n[0] for n in nodes]; ys = [n[1] for n in nodes]
    minx, maxx = min(xs), max(xs); miny, maxy = min(ys), max(ys)
    spanx = (maxx-minx) or 1; spany = (maxy-miny) or 1
    pad = 22
    scale = min((size-2*pad)/spanx, (size-2*pad)/spany)
    cx0, cy0 = (minx+maxx)/2, (miny+maxy)/2
    def tx(pt): return (size/2+(pt[0]-cx0)*scale, size/2+(pt[1]-cy0)*scale)

    for i, rod in enumerate(rods):
        if i == rod_idx: continue
        p0, p1 = tx(nodes[rod['a']]), tx(nodes[rod['b']])
        d.line([p0, p1], fill=(200,200,200), width=2)

    rod = rods[rod_idx]
    p0, p1 = tx(nodes[rod['a']]), tx(nodes[rod['b']])
    color = (136,136,136)
    if rod_res:
        f = rod_res[rod_idx]['force']
        if f > 0.01: color = (226,75,74)
        elif f < -0.01: color = (55,138,221)
    d.line([p0, p1], fill=color, width=4)

    font = _get_ttf_font(11)
    for i, n in enumerate(nodes):
        x, y = tx(n)
        if i in (rod['a'], rod['b']):
            d.ellipse([x-5,y-5,x+5,y+5], fill=(239,159,39), outline=(51,51,51))
            d.text((x, y-16), str(i), fill=(51,51,51), font=font, anchor='mm')
        else:
            d.ellipse([x-3,y-3,x+3,y+3], fill=(153,153,153))
    return img

def pil_draw_node_fbd(vecs, load=None, rxn=None, size=260):
    """
    Pure-PIL free-body diagram: every force vector acting on a node,
    originating at the canvas center. Mirrors the in-app tkinter version
    (TrussApp._draw_node_fbd) but rendered with PIL for Excel embedding.
    """
    from PIL import Image, ImageDraw
    img = Image.new('RGB', (size, size), (255,255,255))
    d = ImageDraw.Draw(img)
    cx = cy = size/2
    d.line([(6,cy),(size-6,cy)], fill=(238,238,238))
    d.line([(cx,6),(cx,size-6)], fill=(238,238,238))

    mags = [v['magnitude'] for v in vecs]
    if load: mags.append(math.hypot(load['fx'], load['fy']))
    if rxn:  mags.append(math.hypot(rxn.get('rx',0), rxn.get('ry',0)))
    maxmag = max(mags) if mags else 1.0
    if maxmag < 1e-9: maxmag = 1.0
    Rmax, Rmin = size*0.34, size*0.13
    font = _get_ttf_font(11)

    def arrow(Fx, Fy, color, label, dashed=False):
        mag = math.hypot(Fx, Fy)
        if mag < 1e-6: return
        length = Rmin + (Rmax-Rmin)*(mag/maxmag)
        ux, uy = Fx/mag, Fy/mag
        ex, ey = cx+ux*length, cy-uy*length
        if dashed:
            n_dash = max(2, int(length/6))
            for k in range(n_dash):
                if k % 2 == 0:
                    x0=cx+ux*length*k/n_dash; y0=cy-uy*length*k/n_dash
                    x1=cx+ux*length*(k+1)/n_dash; y1=cy-uy*length*(k+1)/n_dash
                    d.line([(x0,y0),(x1,y1)], fill=color, width=2)
        else:
            d.line([(cx,cy),(ex,ey)], fill=color, width=2)
        ang = math.atan2(-(ey-cy), ex-cx)
        ah = 8
        a1 = (ex-ah*math.cos(ang-0.4), ey+ah*math.sin(ang-0.4))
        a2 = (ex-ah*math.cos(ang+0.4), ey+ah*math.sin(ang+0.4))
        d.polygon([(ex,ey), a1, a2], fill=color)
        lx, ly = cx+ux*(length+22), cy-uy*(length+22)
        d.multiline_text((lx,ly), label, fill=color, font=font, anchor='mm', align='center')

    for v in vecs:
        color = (226,75,74) if v['kind']=='T' else ((55,138,221) if v['kind']=='C' else (136,136,136))
        arrow(v['Fx'], v['Fy'], color, f'R{v["rod"]}\n{v["magnitude"]:.1f}kN')
    if load and math.hypot(load['fx'], load['fy']) > 1e-6:
        arrow(load['fx'], -load['fy'], (216,90,48),
              f'Load\n{math.hypot(load["fx"],load["fy"]):.1f}kN', dashed=True)
    if rxn and math.hypot(rxn.get('rx',0), rxn.get('ry',0)) > 1e-6:
        arrow(rxn.get('rx',0), -rxn.get('ry',0), (46,204,113),
              f'Rxn\n{math.hypot(rxn.get("rx",0),rxn.get("ry",0)):.1f}kN', dashed=True)
    d.ellipse([cx-4,cy-4,cx+4,cy+4], fill=(51,51,51))
    return img

def _pil_to_xlsx_buf(pil_img):
    """
    openpyxl's Image wrapper expects a PIL Image that was opened from a file
    (it reads img.fp), but images built with Image.new()/ImageDraw don't have
    that. Round-tripping through a BytesIO PNG buffer gives openpyxl what it
    needs while keeping everything in memory (no temp files on disk).
    """
    import io
    buf = io.BytesIO()
    pil_img.save(buf, format='PNG')
    buf.seek(0)
    return buf

def export_excel(nodes, rods, loads, supports, results, path, profiles=None):
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import (Font, PatternFill, Alignment,
                                  Border, Side)
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    reactions = results.get('reactions', {})
    node_res  = results['node_res']
    rod_res   = results['rod_res']

    # ── palette ───────────────────────────────────────────────────────────────
    HDR_FILL  = PatternFill('solid', start_color='1F4E79')   # dark blue
    HDR_FONT  = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    SUB_FILL  = PatternFill('solid', start_color='D6E4F0')   # light blue
    SUB_FONT  = Font(name='Arial', bold=True, color='1F4E79', size=10)
    BODY_FONT = Font(name='Arial', size=10)
    TENS_FILL = PatternFill('solid', start_color='FDECEA')   # light red
    COMP_FILL = PatternFill('solid', start_color='EAF2FB')   # light blue
    ZERO_FILL = PatternFill('solid', start_color='F5F5F5')
    ALT_FILL  = PatternFill('solid', start_color='F7FBFF')

    thin = Side(style='thin', color='BFBFBF')
    med  = Side(style='medium', color='1F4E79')
    bord = Border(left=thin, right=thin, top=thin, bottom=thin)
    med_bord = Border(left=med, right=med, top=med, bottom=med)

    num2  = '0.00'
    num3  = '0.000'
    num1  = '0.0'

    def hdr(ws, row, col, text, w=None):
        c = ws.cell(row=row, column=col, value=text)
        c.font, c.fill, c.alignment, c.border = (
            HDR_FONT, HDR_FILL,
            Alignment(horizontal='center', vertical='center', wrap_text=True),
            med_bord)
        if w: ws.column_dimensions[get_column_letter(col)].width = w
        return c

    def sub(ws, row, col, text):
        c = ws.cell(row=row, column=col, value=text)
        c.font = SUB_FONT; c.fill = SUB_FILL
        c.alignment = Alignment(horizontal='center'); c.border = bord
        return c

    def cell(ws, row, col, val, fmt=None, fill=None, bold=False):
        c = ws.cell(row=row, column=col, value=val)
        c.font = Font(name='Arial', size=10, bold=bold)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = bord
        if fmt: c.number_format = fmt
        if fill: c.fill = fill
        return c

    # ══════════════════════════════════════════════════════════════════════════
    #  Sheet 1 – NODES
    # ══════════════════════════════════════════════════════════════════════════
    ws = wb.active
    ws.title = 'Nodes'
    ws.freeze_panes = 'A3'
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22

    # title
    ws.merge_cells('A1:P1')
    t = ws['A1']
    t.value = 'STEEL TRUSS SIMULATOR — NODE REPORT'
    t.font  = Font(name='Arial', bold=True, size=13, color='1F4E79')
    t.alignment = Alignment(horizontal='center', vertical='center')
    t.fill  = PatternFill('solid', start_color='D6E4F0')

    # column groups
    groups = [
        ('Node', 1, 1),
        ('Coordinates (m)', 2, 3),
        ('Connected rods', 4, 4),
        ('Applied load (kN)', 5, 6),
        ('Support', 7, 7),
        ('Reactions (kN)', 8, 9),
        ('Displacement (mm)', 10, 11),
    ]
    for label, c1, c2 in groups:
        if c1 == c2:
            hdr(ws, 2, c1, label)
        else:
            ws.merge_cells(start_row=2, start_column=c1,
                           end_row=2, end_column=c2)
            hdr(ws, 2, c1, label)

    sub_labels = ['#','X','Y','Rods','Fx','Fy',
                  'Type','Rx','Ry','ux','uy']
    widths     = [5, 7, 7, 18, 8, 8, 10, 8, 8, 9, 9]
    for col, (lbl, w) in enumerate(zip(sub_labels, widths), 1):
        sub(ws, 3, col, lbl)
        ws.column_dimensions[get_column_letter(col)].width = w

    for ni, nd in enumerate(nodes):
        row = 4 + ni
        fill = ALT_FILL if ni % 2 else None
        xm   = nd[0]/PX_PER_M
        ym   = -nd[1]/PX_PER_M     # flip y (canvas y grows down)

        connected = [str(ri) for ri,r in enumerate(rods) if r['a']==ni or r['b']==ni]
        ld  = next((l for l in loads    if l['node']==ni), None)
        sup = next((s for s in supports if s['node']==ni), None)
        rxn = reactions.get(ni, {})
        res = node_res[ni]

        cell(ws, row, 1,  ni,          fill=fill, bold=True)
        cell(ws, row, 2,  xm,   num2,  fill=fill)
        cell(ws, row, 3,  ym,   num2,  fill=fill)
        cell(ws, row, 4,  ', '.join(connected) if connected else '—', fill=fill)
        cell(ws, row, 5,  ld['fx'] if ld else 0.0, num2, fill=fill)
        cell(ws, row, 6,  ld['fy'] if ld else 0.0, num2, fill=fill)
        cell(ws, row, 7,  sup['type'] if sup else '—', fill=fill)
        cell(ws, row, 8,  rxn.get('rx', '—'), num2 if rxn else None, fill=fill)
        cell(ws, row, 9,  rxn.get('ry', '—'), num2 if rxn else None, fill=fill)
        cell(ws, row, 10, res['ux'], num3, fill=fill)
        cell(ws, row, 11, res['uy'], num3, fill=fill)

    # ══════════════════════════════════════════════════════════════════════════
    #  Sheet 2 – RODS
    # ══════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet('Rods')
    ws2.freeze_panes = 'A3'
    ws2.row_dimensions[1].height = 28
    ws2.row_dimensions[2].height = 22

    ws2.merge_cells('A1:L1')
    t2 = ws2['A1']
    t2.value = 'STEEL TRUSS SIMULATOR — ROD REPORT (pure axial forces)'
    t2.font  = Font(name='Arial', bold=True, size=13, color='1F4E79')
    t2.alignment = Alignment(horizontal='center', vertical='center')
    t2.fill  = PatternFill('solid', start_color='D6E4F0')

    rod_groups = [
        ('Rod',          1, 1),
        ('End nodes',    2, 3),
        ('Geometry',     4, 6),
        ('Material (assumed)', 7, 8),
        ('Results',      9, 15),
    ]
    for label, c1, c2 in rod_groups:
        if c1 == c2:
            hdr(ws2, 2, c1, label)
        else:
            ws2.merge_cells(start_row=2, start_column=c1,
                            end_row=2, end_column=c2)
            hdr(ws2, 2, c1, label)

    rod_subs = ['#','Node A','Node B','Length (m)','Δx (m)','Δy (m)',
                'E (GPa)','A (cm²)','Axial Force (kN)','Connection',
                'Shear V (kN)','Moment Ma (kN·m)','Moment Mb (kN·m)',
                'Load @ A (kN)','Load @ B (kN)']
    rod_widths = [5,7,7,10,8,8,8,8,14,10,11,14,14,14,14]
    for col,(lbl,w) in enumerate(zip(rod_subs,rod_widths),1):
        sub(ws2, 3, col, lbl)
        ws2.column_dimensions[get_column_letter(col)].width = w

    for ri, rod in enumerate(rods):
        row  = 4 + ri
        na, nb = nodes[rod['a']], nodes[rod['b']]
        ddx  = (nb[0]-na[0])/PX_PER_M
        ddy  = -(nb[1]-na[1])/PX_PER_M
        Lm   = math.hypot(ddx, ddy)
        rr   = rod_res[ri]
        f    = rr['force']
        kind = 'Tension' if f>0.01 else ('Compression' if f<-0.01 else 'Zero')
        rfill = TENS_FILL if f>0.01 else (COMP_FILL if f<-0.01 else ZERO_FILL)
        is_rigid = rr.get('conn') == 'rigid'

        ld_a = next((l for l in loads if l['node']==rod['a']), None)
        ld_b = next((l for l in loads if l['node']==rod['b']), None)
        la_str = (f"Fx={ld_a['fx']} Fy={ld_a['fy']}" if ld_a else '—')
        lb_str = (f"Fx={ld_b['fx']} Fy={ld_b['fy']}" if ld_b else '—')

        cell(ws2, row, 1,  ri,           fill=rfill, bold=True)
        cell(ws2, row, 2,  rod['a'],     fill=rfill)
        cell(ws2, row, 3,  rod['b'],     fill=rfill)
        cell(ws2, row, 4,  Lm,  num2,   fill=rfill)
        cell(ws2, row, 5,  ddx, num2,   fill=rfill)
        cell(ws2, row, 6,  ddy, num2,   fill=rfill)
        cell(ws2, row, 7,  rod['E'],     fill=rfill)
        cell(ws2, row, 8,  rod['A'],     fill=rfill)
        cell(ws2, row, 9,  f,   num2,   fill=rfill)
        cell(ws2, row, 10, 'Rigid' if is_rigid else 'Pinned', fill=rfill, bold=True)
        cell(ws2, row, 11, rr.get('V', 0.0) if is_rigid else '—', num2, fill=rfill)
        cell(ws2, row, 12, rr.get('Ma', 0.0) if is_rigid else '—', num2, fill=rfill)
        cell(ws2, row, 13, rr.get('Mb', 0.0) if is_rigid else '—', num2, fill=rfill)
        cell(ws2, row, 14, la_str,       fill=rfill)
        cell(ws2, row, 15, lb_str,       fill=rfill)


    # ══════════════════════════════════════════════════════════════════════════
    #  Sheet 3 – SUMMARY
    # ══════════════════════════════════════════════════════════════════════════
    ws3 = wb.create_sheet('Summary')
    ws3.column_dimensions['A'].width = 32
    ws3.column_dimensions['B'].width = 18
    ws3.column_dimensions['C'].width = 12

    def srow(r, label, val, fmt=None, bold=False, fill=None):
        c1 = ws3.cell(row=r, column=1, value=label)
        c1.font = Font(name='Arial', size=10, bold=bold)
        c1.border = bord
        if fill: c1.fill = fill
        c2 = ws3.cell(row=r, column=2, value=val)
        c2.font = Font(name='Arial', size=10, bold=bold)
        c2.border = bord; c2.alignment = Alignment(horizontal='right')
        if fmt: c2.number_format = fmt
        if fill: c2.fill = fill

    ws3.merge_cells('A1:C1')
    t3 = ws3['A1']
    t3.value = 'STEEL TRUSS SIMULATOR — SUMMARY'
    t3.font  = Font(name='Arial', bold=True, size=13, color='1F4E79')
    t3.alignment = Alignment(horizontal='center', vertical='center')
    t3.fill  = PatternFill('solid', start_color='D6E4F0')
    ws3.row_dimensions[1].height = 28

    r = 2
    hdr(ws3, r, 1, 'Parameter'); hdr(ws3, r, 2, 'Value'); hdr(ws3, r, 3, 'Unit')

    def srow3(row, label, val, unit, fmt=None, bold=False, fill=None):
        for c, v in [(1,label),(2,val),(3,unit)]:
            cc = ws3.cell(row=row, column=c, value=v)
            cc.font = Font(name='Arial', size=10, bold=bold)
            cc.border = bord; cc.alignment = Alignment(horizontal='center')
            if fmt and c==2: cc.number_format = fmt
            if fill: cc.fill = fill

    forces = [r['force'] for r in rod_res]
    disps  = [math.hypot(r['ux'],r['uy']) for r in node_res]
    max_t  = max(forces, default=0)
    max_c  = min(forces, default=0)
    max_d  = max(disps,  default=0)
    # Include interior extrema from member UDLs and point loads, not merely
    # the two element ends.  `compute_diagrams` inserts exact V=0 stations.
    member_diagrams = compute_diagrams(nodes, rods, loads, results)
    rigid_diagrams = [d for rod, d in zip(rods, member_diagrams)
                      if rod.get('conn', 'pin') == 'rigid']
    max_V = max((abs(v) for d in rigid_diagrams for v in d['V']), default=0)
    max_M = max((abs(m) for d in rigid_diagrams for m in d['M']), default=0)
    tot_rx = sum(v.get('rx',0) for v in reactions.values() if isinstance(v.get('rx'),float))
    tot_ry = sum(v.get('ry',0) for v in reactions.values() if isinstance(v.get('ry'),float))
    tot_fl = sum(ld['fy'] for ld in loads)

    data = [
        ('Nodes',                 len(nodes),  '—',       None,  True,  SUB_FILL),
        ('Rods',                  len(rods),   '—',       None,  True,  SUB_FILL),
        ('Supports',              len(supports),'—',      None,  False, None),
        ('Applied loads',         len(loads),  '—',       None,  False, None),
        ('──────────────────────','',           '',        None,  False, None),
        ('Max axial tension',     max_t,       'kN',     num2,  True,  TENS_FILL),
        ('Max axial compression', abs(max_c),  'kN',     num2,  True,  COMP_FILL),
        ('Max shear |V| (rigid rods)', max_V,  'kN',     num2,  True,  None),
        ('Max moment |M| (rigid rods)',max_M,  'kN·m',   num2,  True,  None),
        ('Max nodal displacement',max_d,       'mm',     num3,  False, None),
        ('──────────────────────','',           '',        None,  False, None),
        ('ΣRx (horizontal equil.)',tot_rx,     'kN',     num2,  False, None),
        ('ΣRy (vertical equil.)', tot_ry,      'kN',     num2,  False, None),
        ('Total applied Fy',      tot_fl,      'kN',     num2,  False, None),
        ('Check Σy=0 (Ry+Fy)',   tot_ry+tot_fl,'kN',    num3,  True,
         PatternFill('solid', start_color='E2EFDA') if abs(tot_ry+tot_fl)<0.01
         else PatternFill('solid', start_color='FDECEA')),
    ]
    for i, (lbl, val, unit, fmt, bold, fill) in enumerate(data, 3):
        srow3(i, lbl, val, unit, fmt, bold, fill)
        ws3.row_dimensions[i].height = 18

    # reaction detail block
    r_start = 3+len(data)+1
    ws3.merge_cells(f'A{r_start}:C{r_start}')
    hc = ws3.cell(row=r_start, column=1, value='Reaction Details per Support Node')
    hc.font = HDR_FONT; hc.fill = HDR_FILL
    hc.alignment = Alignment(horizontal='center')
    ws3.row_dimensions[r_start].height = 20
    r_start += 1
    for col, lbl in [(1,'Node'),(2,'Rx (kN)'),(3,'Ry (kN)')]:
        sub(ws3, r_start, col, lbl)
    r_start += 1
    for ni, rxn in reactions.items():
        for col, val in [(1,ni),(2,rxn.get('rx','—')),(3,rxn.get('ry','—'))]:
            c = ws3.cell(row=r_start, column=col, value=val)
            c.font = BODY_FONT; c.border = bord
            c.alignment = Alignment(horizontal='center')
            if isinstance(val, float): c.number_format = num2
        r_start += 1

    # ══════════════════════════════════════════════════════════════════════════
    #  Sheet 4 – NODE FORCE VECTORS  (for gusset / node-plate design)
    # ══════════════════════════════════════════════════════════════════════════
    node_vectors = results.get('node_vectors', {})
    ws4 = wb.create_sheet('Node Force Vectors')
    ws4.freeze_panes = 'A3'
    ws4.row_dimensions[1].height = 28
    ws4.row_dimensions[2].height = 22

    ws4.merge_cells('A1:H1')
    t4 = ws4['A1']
    t4.value = 'STEEL TRUSS SIMULATOR — NODE FORCE VECTORS (gusset plate design)'
    t4.font  = Font(name='Arial', bold=True, size=13, color='1F4E79')
    t4.alignment = Alignment(horizontal='center', vertical='center')
    t4.fill  = PatternFill('solid', start_color='D6E4F0')

    note = ws4.cell(row=2, column=1,
        value='Convention: global X-> right, Y-up positive, angle CCW from +X. '
              'Vector = force the rod exerts ON the node (tension pulls, compression pushes).')
    ws4.merge_cells('A2:H2')
    note.font = Font(name='Arial', italic=True, size=9, color='555555')
    note.alignment = Alignment(horizontal='left', vertical='center')

    nv_subs   = ['Node','Rod','-> Node','Type','|F| (kN)','Fx (kN)','Fy (kN)','Angle (deg)']
    nv_widths = [7, 6, 8, 7, 10, 10, 10, 12]
    for col, (lbl, w) in enumerate(zip(nv_subs, nv_widths), 1):
        sub(ws4, 3, col, lbl)
        ws4.column_dimensions[get_column_letter(col)].width = w

    row = 4
    for ni in range(len(nodes)):
        vecs = node_vectors.get(ni, [])
        if not vecs:
            continue
        first = True
        for v in vecs:
            rfill = TENS_FILL if v['kind']=='T' else (COMP_FILL if v['kind']=='C' else ZERO_FILL)
            cell(ws4, row, 1, ni if first else '', fill=rfill, bold=first)
            cell(ws4, row, 2, v['rod'], fill=rfill)
            cell(ws4, row, 3, v['other'], fill=rfill)
            cell(ws4, row, 4, v['kind'], fill=rfill)
            cell(ws4, row, 5, v['magnitude'], num2, fill=rfill)
            cell(ws4, row, 6, v['Fx'], num2, fill=rfill)
            cell(ws4, row, 7, v['Fy'], num2, fill=rfill)
            cell(ws4, row, 8, v['angle_deg'], num1, fill=rfill)
            row += 1
            first = False
        # equilibrium check row: sum(rod vectors) + load + reaction ~= 0
        sum_fx = sum(v['Fx'] for v in vecs)
        sum_fy = sum(v['Fy'] for v in vecs)
        ld  = next((l for l in loads if l['node']==ni), None)
        rxn = reactions.get(ni)
        chk_fx = sum_fx + (ld['fx'] if ld else 0.0) + (rxn.get('rx',0.0) if rxn else 0.0)
        chk_fy = sum_fy + (-ld['fy'] if ld else 0.0) + (-rxn.get('ry',0.0) if rxn else 0.0)
        cell(ws4, row, 1, '', bold=True)
        cell(ws4, row, 4, 'Sum F+load+rxn', bold=True)
        cell(ws4, row, 6, chk_fx, num3, fill=SUB_FILL, bold=True)
        cell(ws4, row, 7, chk_fy, num3, fill=SUB_FILL, bold=True)
        row += 1

    # ══════════════════════════════════════════════════════════════════════════
    #  Sheet 5 – ROD CALCULATIONS  (picture-in-context + FBD at each end)
    # ══════════════════════════════════════════════════════════════════════════
    try:
        from openpyxl.drawing.image import Image as XLImage
        ws5 = wb.create_sheet('Rod Calculations')
        ws5.merge_cells('A1:N1')
        t5 = ws5['A1']
        t5.value = 'STEEL TRUSS SIMULATOR — ROD CALCULATIONS (picture + force vectors)'
        t5.font = Font(name='Arial', bold=True, size=13, color='1F4E79')
        t5.alignment = Alignment(horizontal='center', vertical='center')
        t5.fill = PatternFill('solid', start_color='D6E4F0')
        note5 = ws5.cell(row=2, column=1,
            value='Left: this rod\'s true location in the truss. Middle/Right: the free-body '
                  'diagram at each end node (every rod, load, and reaction converging there).')
        ws5.merge_cells('A2:N2')
        note5.font = Font(name='Arial', italic=True, size=9, color='555555')

        IMG_PX = 230           # square image size in px
        ROWS_PER_BLOCK = 13    # excel rows spanned by one image block (~ IMG_PX tall)
        COLS_PER_IMG   = 5     # excel columns spanned by one image (~ IMG_PX wide)
        for c in range(1, 3*COLS_PER_IMG+3):
            ws5.column_dimensions[get_column_letter(c)].width = 9

        row0 = 4
        for ri, rod in enumerate(rods):
            a, b = rod['a'], rod['b']
            na, nb = nodes[a], nodes[b]
            dxp, dyp = nb[0]-na[0], nb[1]-na[1]
            Lm = math.hypot(dxp, dyp) / PX_PER_M
            f = rod_res[ri]['force']
            kind = 'Tension' if f>0.01 else ('Compression' if f<-0.01 else 'Zero')
            hdr_row = row0
            ws5.merge_cells(start_row=hdr_row, start_column=1, end_row=hdr_row, end_column=14)
            hc = ws5.cell(row=hdr_row, column=1,
                value=f'Rod {ri}  (Node {a} → Node {b})   L={Lm:.3f} m   '
                      f'N={f:+.2f} kN   [{kind}]   profile="{rod.get("profile","Default")}"')
            hc.font = Font(name='Arial', bold=True, size=11, color='1F4E79')
            hc.fill = PatternFill('solid', start_color='EFF4FA')

            img_row = hdr_row + 1
            ctx_img = pil_draw_rod_context(nodes, rods, ri, rod_res, size=IMG_PX)
            fbd_a   = pil_draw_node_fbd(node_vectors.get(a, []),
                          next((l for l in loads if l['node']==a), None),
                          reactions.get(a), size=IMG_PX)
            fbd_b   = pil_draw_node_fbd(node_vectors.get(b, []),
                          next((l for l in loads if l['node']==b), None),
                          reactions.get(b), size=IMG_PX)

            ws5.cell(row=img_row, column=1, value='Location in truss').font = Font(size=9, italic=True, color='777777')
            ws5.cell(row=img_row, column=1+COLS_PER_IMG, value=f'End A — Node {a}').font = Font(size=9, italic=True, color='777777')
            ws5.cell(row=img_row, column=1+2*COLS_PER_IMG, value=f'End B — Node {b}').font = Font(size=9, italic=True, color='777777')

            ws5.add_image(XLImage(_pil_to_xlsx_buf(ctx_img)), f'{get_column_letter(1)}{img_row+1}')
            ws5.add_image(XLImage(_pil_to_xlsx_buf(fbd_a)),   f'{get_column_letter(1+COLS_PER_IMG)}{img_row+1}')
            ws5.add_image(XLImage(_pil_to_xlsx_buf(fbd_b)),   f'{get_column_letter(1+2*COLS_PER_IMG)}{img_row+1}')

            row0 = hdr_row + ROWS_PER_BLOCK
    except Exception:
        pass   # if Pillow/matplotlib aren't available, skip the picture sheet gracefully

    # ══════════════════════════════════════════════════════════════════════════
    #  Sheet 6 – MODEL  (machine-parseable — lets "Import from Excel" rebuild
    #  this exact truss, loads included)
    # ══════════════════════════════════════════════════════════════════════════
    _write_model_sheet(wb, nodes, rods, loads, supports, profiles or {'Default':{'E':200.0,'A':10.0}})

    wb.save(path)


def _write_model_sheet(wb, nodes, rods, loads, supports, profiles):
    """
    Writes a plain, machine-parseable 'Model' sheet containing everything
    needed to rebuild this exact truss (nodes, rods incl. profile, supports,
    loads, and the profile/family definitions). Paired with
    import_excel_model() for the round trip.
    """
    ws = wb.create_sheet('Model')
    ws.sheet_state = 'visible'
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    row = 1
    ws.cell(row=row, column=1, value='TRUSS MODEL DATA — for Import from Excel (do not reorder columns)')
    ws.cell(row=row, column=1).font = Font(bold=True, size=11, color='1F4E79')
    row += 2

    ws.cell(row=row, column=1, value='[PROFILES]'); row += 1
    ws.cell(row=row, column=1, value='name'); ws.cell(row=row, column=2, value='E_GPa')
    ws.cell(row=row, column=3, value='A_cm2'); ws.cell(row=row, column=4, value='I_cm4'); row += 1
    for name, p in profiles.items():
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=p['E'])
        ws.cell(row=row, column=3, value=p['A'])
        ws.cell(row=row, column=4, value=p.get('I', 8000.0))
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[NODES]'); row += 1
    ws.cell(row=row, column=1, value='idx'); ws.cell(row=row, column=2, value='x_m')
    ws.cell(row=row, column=3, value='y_m'); row += 1
    for i, n in enumerate(nodes):
        mx, my = n[0]/PX_PER_M, -n[1]/PX_PER_M
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=mx)
        ws.cell(row=row, column=3, value=my)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[RODS]'); row += 1
    for col, lbl in enumerate(['idx','a','b','E_GPa','A_cm2','I_cm4','conn','profile',
                                'udl_kNm','udl_rotation_deg'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for i, r in enumerate(rods):
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=r['a'])
        ws.cell(row=row, column=3, value=r['b'])
        ws.cell(row=row, column=4, value=r['E'])
        ws.cell(row=row, column=5, value=r['A'])
        ws.cell(row=row, column=6, value=r.get('I', 8000.0))
        ws.cell(row=row, column=7, value=r.get('conn', 'pin'))
        ws.cell(row=row, column=8, value=r.get('profile','Default'))
        ws.cell(row=row, column=9, value=r.get('udl', 0.0))
        ws.cell(row=row, column=10, value=r.get('udl_rotation_deg', 0.0))
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[POINT_LOADS]'); row += 1
    for col, lbl in enumerate(['rod','P_kN','t_from_A','angle_deg'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for ri, r in enumerate(rods):
        for pl in r.get('point_loads', []):
            ws.cell(row=row, column=1, value=ri)
            ws.cell(row=row, column=2, value=pl.get('P', 0.0))
            ws.cell(row=row, column=3, value=pl.get('t', 0.5))
            ws.cell(row=row, column=4, value=pl.get('angle_deg', 90.0))
            row += 1
    row += 1

    ws.cell(row=row, column=1, value='[SUPPORTS]'); row += 1
    ws.cell(row=row, column=1, value='node'); ws.cell(row=row, column=2, value='type'); row += 1
    for s in supports:
        ws.cell(row=row, column=1, value=s['node'])
        ws.cell(row=row, column=2, value=s['type'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[LOADS]'); row += 1
    for col, lbl in enumerate(['node','fx_kN','fy_kN'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for l in loads:
        ws.cell(row=row, column=1, value=l['node'])
        ws.cell(row=row, column=2, value=l['fx'])
        ws.cell(row=row, column=3, value=l['fy'])
        row += 1

    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 12


def import_excel_model(path):
    """
    Reads a 'Model' sheet written by _write_model_sheet() and rebuilds
    nodes, rods, supports, loads, and profiles from it. Coordinates are
    converted back from metres to this app's internal pixel/world units.
    Returns (nodes, rods, loads, supports, profiles) or raises ValueError
    with a human-readable message if the sheet isn't found / is malformed.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if 'Model' not in wb.sheetnames:
        raise ValueError('This workbook has no "Model" sheet — it was not '
                          'exported by this app (or is an older report-only export).')
    ws = wb['Model']
    rows = list(ws.iter_rows(values_only=True))

    def find_section(name):
        for i, r in enumerate(rows):
            if r and r[0] == name:
                return i
        return -1

    def read_table(start_idx):
        """start_idx points at the [SECTION] row; returns list of dict rows."""
        hdr_i = start_idx + 1
        if hdr_i >= len(rows) or not rows[hdr_i] or rows[hdr_i][0] is None:
            return []
        headers = [h for h in rows[hdr_i] if h is not None]
        out = []
        i = hdr_i + 1
        while i < len(rows) and rows[i] and rows[i][0] is not None and not str(rows[i][0]).startswith('['):
            vals = rows[i]
            out.append({headers[j]: vals[j] for j in range(len(headers))})
            i += 1
        return out

    profiles = {}
    pi = find_section('[PROFILES]')
    if pi >= 0:
        for r in read_table(pi):
            profiles[str(r['name'])] = {'E': float(r['E_GPa']), 'A': float(r['A_cm2']),
                                        'I': float(r.get('I_cm4') or 8000.0)}
    if not profiles:
        profiles = {'Default': {'E': 200.0, 'A': 10.0, 'I': 8000.0}}

    nodes = []
    ni = find_section('[NODES]')
    if ni >= 0:
        for r in read_table(ni):
            mx, my = float(r['x_m']), float(r['y_m'])
            nodes.append((mx*PX_PER_M, -my*PX_PER_M))

    rods = []
    ri = find_section('[RODS]')
    if ri >= 0:
        for r in read_table(ri):
            rods.append({'a': int(r['a']), 'b': int(r['b']),
                         'E': float(r['E_GPa']), 'A': float(r['A_cm2']),
                         'I': float(r.get('I_cm4') or 8000.0),
                         'conn': str(r.get('conn') or 'pin'),
                         'profile': str(r.get('profile') or 'Default'),
                         'udl': float(r.get('udl_kNm') or 0.0),
                         'udl_rotation_deg': float(r.get('udl_rotation_deg') or 0.0),
                         'point_loads': []})

    pli = find_section('[POINT_LOADS]')
    if pli >= 0:
        for r in read_table(pli):
            ri0 = int(r['rod'])
            if 0 <= ri0 < len(rods):
                rods[ri0].setdefault('point_loads', []).append({
                    'P': float(r.get('P_kN') or 0.0),
                    't': float(r.get('t_from_A') or 0.5),
                    'angle_deg': float(r.get('angle_deg') or 90.0)
                })

    supports = []
    si = find_section('[SUPPORTS]')
    if si >= 0:
        for r in read_table(si):
            supports.append({'node': int(r['node']), 'type': str(r['type'])})

    loads = []
    li = find_section('[LOADS]')
    if li >= 0:
        for r in read_table(li):
            loads.append({'node': int(r['node']), 'fx': float(r['fx_kN']), 'fy': float(r['fy_kN'])})

    return nodes, rods, loads, supports, profiles

