"""Tests for gusset-plate verification in 3D space frames
(apps/stereo/stereo_plates.py).

Covers coplanarity detection, Whitmore width, block shear areas,
gusset buckling, and the full check_all_gussets pipeline.
"""
import math
import pytest

from apps.stereo import stereo_plates as sp


# ── geometry helpers ─────────────────────────────────────────────────────

class TestVectorHelpers:
    def test_cross_product(self):
        c = sp._cross((1, 0, 0), (0, 1, 0))
        assert c == pytest.approx((0, 0, 1))

    def test_dot_product(self):
        assert sp._dot((1, 2, 3), (4, 5, 6)) == pytest.approx(32.0)

    def test_normalise(self):
        n = sp._normalise((3, 4, 0))
        assert n == pytest.approx((0.6, 0.8, 0.0))

    def test_normalise_zero_vector(self):
        assert sp._normalise((0, 0, 0)) == (0.0, 0.0, 0.0)


# ── coplanarity detection ───────────────────────────────────────────────

class TestCoplanarGroups:
    def test_triangle_in_xy_plane_yields_one_group(self):
        nodes = [(0, 0, 0), (1, 0, 0), (0, 1, 0)]
        members = [{'a': 0, 'b': 1, 'conn': 'pin'},
                   {'a': 0, 'b': 2, 'conn': 'pin'},
                   {'a': 1, 'b': 2, 'conn': 'pin'}]
        groups = sp.find_coplanar_groups(nodes, members, 0)
        assert len(groups) == 1
        mis = {d['mi'] for d in groups[0]}
        assert mis == {0, 1}

    def test_tetrahedron_apex_has_three_groups(self):
        nodes = [(0, 0, 1),    # apex
                 (1, 0, 0), (0, 1, 0), (-1, 0, 0)]
        members = [{'a': 0, 'b': 1, 'conn': 'pin'},
                   {'a': 0, 'b': 2, 'conn': 'pin'},
                   {'a': 0, 'b': 3, 'conn': 'pin'}]
        groups = sp.find_coplanar_groups(nodes, members, 0)
        assert len(groups) >= 3

    def test_node_with_one_member_returns_no_groups(self):
        nodes = [(0, 0, 0), (1, 0, 0)]
        members = [{'a': 0, 'b': 1, 'conn': 'pin'}]
        groups = sp.find_coplanar_groups(nodes, members, 0)
        assert groups == []

    def test_four_coplanar_members_yield_one_group(self):
        nodes = [(0, 0, 0), (1, 0, 0), (0, 1, 0), (-1, 0, 0), (0, -1, 0)]
        members = [{'a': 0, 'b': i, 'conn': 'pin'} for i in range(1, 5)]
        groups = sp.find_coplanar_groups(nodes, members, 0)
        coplanar_group = [g for g in groups if len(g) >= 4]
        assert len(coplanar_group) >= 1


# ── Whitmore width ──────────────────────────────────────────────────────

class TestWhitmoreWidth:
    def test_default_width(self):
        w = sp.whitmore_width_m()
        expected = 2.0 * 0.30 * math.tan(math.radians(30))
        assert w == pytest.approx(expected)

    def test_custom_landing(self):
        w = sp.whitmore_width_m(landing_m=0.50)
        expected = 2.0 * 0.50 * math.tan(math.radians(30))
        assert w == pytest.approx(expected)

    def test_width_is_positive(self):
        assert sp.whitmore_width_m() > 0.0


# ── block shear areas ──────────────────────────────────────────────────

class TestBlockShearAreas:
    def test_welded_net_equals_gross(self):
        Agv, Anv, Ant, Agt = sp.block_shear_areas_welded(10.0, 300.0, 346.0)
        assert Agv == Anv
        assert Ant == Agt

    def test_shear_area_formula(self):
        t, landing, whitmore = 10.0, 300.0, 346.0
        Agv, Anv, Ant, Agt = sp.block_shear_areas_welded(t, landing, whitmore)
        assert Agv == pytest.approx(2.0 * t * landing)
        assert Ant == pytest.approx(t * whitmore)

    def test_zero_thickness(self):
        Agv, _, _, _ = sp.block_shear_areas_welded(0.0, 300.0, 346.0)
        assert Agv == 0.0


# ── gusset buckling ─────────────────────────────────────────────────────

class TestGussetBuckling:
    def test_tension_does_not_apply(self):
        r = sp.gusset_buckling_check(50.0, 346.0, 10.0, 300.0, 235.0)
        assert r['applies'] is False
        assert r['util'] == 0.0

    def test_compression_applies(self):
        r = sp.gusset_buckling_check(-50.0, 346.0, 10.0, 300.0, 235.0)
        assert r['applies'] is True
        assert r['util'] > 0.0
        assert 'Fcr_MPa' in r
        assert 'Pd_kN' in r

    def test_zero_force_does_not_apply(self):
        r = sp.gusset_buckling_check(0.0, 346.0, 10.0, 300.0, 235.0)
        assert r['applies'] is False


# ── check_gusset_at_node ────────────────────────────────────────────────

def _simple_model():
    """A triangle in the XY plane: 3 nodes, 3 members."""
    nodes = [(0, 0, 0), (3, 0, 0), (1.5, 2.598, 0)]
    members = [
        {'a': 0, 'b': 1, 'conn': 'pin', 'A': 20.0, 'Fy': 235.0, 'Fu': 360.0},
        {'a': 0, 'b': 2, 'conn': 'pin', 'A': 20.0, 'Fy': 235.0, 'Fu': 360.0},
        {'a': 1, 'b': 2, 'conn': 'pin', 'A': 20.0, 'Fy': 235.0, 'Fu': 360.0},
    ]
    member_res = [
        {'N': 50.0, 'conn': 'pin', 'length_m': 3.0},
        {'N': -30.0, 'conn': 'pin', 'length_m': 3.0},
        {'N': 10.0, 'conn': 'pin', 'length_m': 3.0},
    ]
    return nodes, members, member_res


class TestCheckGussetAtNode:
    def test_returns_one_entry_per_member_in_group(self):
        nodes, members, member_res = _simple_model()
        groups = sp.find_coplanar_groups(nodes, members, 0)
        assert len(groups) >= 1
        group = groups[0]
        results = sp.check_gusset_at_node(0, group, member_res)
        assert len(results) == len(group)

    def test_each_result_has_required_keys(self):
        nodes, members, member_res = _simple_model()
        groups = sp.find_coplanar_groups(nodes, members, 0)
        results = sp.check_gusset_at_node(0, groups[0], member_res)
        required_keys = {'node', 'member', 'other', 'N_kN', 'mode',
                        'whitmore_mm', 'sigma_MPa', 'whitmore_util',
                        'block_shear_util', 'buckling_applies', 'buckling_util',
                        'util', 'governing', 'ok'}
        for r in results:
            assert required_keys.issubset(r.keys()), \
                f'missing keys: {required_keys - r.keys()}'

    def test_tension_member_mode(self):
        nodes, members, member_res = _simple_model()
        groups = sp.find_coplanar_groups(nodes, members, 0)
        results = sp.check_gusset_at_node(0, groups[0], member_res)
        for r in results:
            if r['N_kN'] > 0:
                assert r['mode'] == 'tension'
            elif r['N_kN'] < 0:
                assert r['mode'] == 'compression'

    def test_compression_member_has_buckling(self):
        nodes, members, member_res = _simple_model()
        groups = sp.find_coplanar_groups(nodes, members, 0)
        results = sp.check_gusset_at_node(0, groups[0], member_res)
        comp_results = [r for r in results if r['N_kN'] < -1e-9]
        for r in comp_results:
            assert r['buckling_applies'] is True


# ── check_all_gussets ───────────────────────────────────────────────────

class TestCheckAllGussets:
    def test_returns_flat_list(self):
        nodes, members, member_res = _simple_model()
        all_checks = sp.check_all_gussets(nodes, members, member_res)
        assert isinstance(all_checks, list)
        assert len(all_checks) > 0

    def test_every_entry_has_group_key(self):
        nodes, members, member_res = _simple_model()
        all_checks = sp.check_all_gussets(nodes, members, member_res)
        for c in all_checks:
            assert 'group' in c

    def test_utilization_is_positive(self):
        nodes, members, member_res = _simple_model()
        all_checks = sp.check_all_gussets(nodes, members, member_res)
        for c in all_checks:
            assert c['util'] >= 0.0


# ── Excel round-trip with enhanced sheets ──────────────────────────────

class TestEnhancedExcelExport:
    def _export_and_open(self, tmp_path):
        """Export a solved model and return the openpyxl workbook."""
        import tempfile, os
        from apps.stereo import stereo_geometry as sg
        from apps.stereo import stereo_math as sm
        from apps.stereo import stereo_reports as sr
        from apps.stereo import stereo_checks as sc

        mesh = sg.flat_grid(span_x=6.0, span_y=6.0, depth=1.0, module=3.0,
                            offset=True)
        nodes, members = mesh['nodes'], mesh['members']
        for m in members:
            m.update(E=200.0, A=15.0, I=400.0, J=400.0,
                     Fy=250.0, Fu=400.0, r_gyr=2.5, K=1.0)
        supports = [{'node': n, 'type': 'pin'}
                    for n in mesh['support_candidates']]
        loads = sm.self_weight_loads(nodes, members) + [
            {'node': 0, 'fz': -5.0, 'fx': 0, 'fy': 0, 'mx': 0, 'my': 0, 'mz': 0}]
        res, err = sm.analyze(nodes, members, loads, supports)
        assert err is None
        checks = sc.check_all_members(nodes, members, res['member_res'])

        path = str(tmp_path / 'enhanced.xlsx')
        sr.export_excel(nodes, members, loads, supports, res, path,
                        checks=checks)

        from common import _ensure_openpyxl
        assert _ensure_openpyxl()
        import openpyxl
        return openpyxl.load_workbook(path), nodes, members, res, checks

    def test_node_properties_sheet_exists(self, tmp_path):
        wb, nodes, members, res, checks = self._export_and_open(tmp_path)
        assert 'Node Properties' in wb.sheetnames

    def test_node_properties_has_displacement_and_reaction_columns(self, tmp_path):
        wb, nodes, members, res, checks = self._export_and_open(tmp_path)
        ws = wb['Node Properties']
        headers = [ws.cell(row=1, column=c).value for c in range(1, 14)]
        assert 'ux_mm' in headers
        assert 'Fx_kN' in headers
        assert 'Mz_kNm' in headers

    def test_node_properties_row_count(self, tmp_path):
        wb, nodes, members, res, checks = self._export_and_open(tmp_path)
        ws = wb['Node Properties']
        assert ws.max_row == 1 + len(nodes)

    def test_member_checks_has_detailed_columns(self, tmp_path):
        wb, nodes, members, res, checks = self._export_and_open(tmp_path)
        ws = wb['Member Checks']
        headers = [ws.cell(row=1, column=c).value for c in range(1, 21)]
        assert 'slenderness' in headers
        assert 'sigma_MPa' in headers
        assert 'Fcr_MPa' in headers
        assert 'Pd_kN' in headers
        assert 'A_cm2' in headers
        assert 'r_gyr_cm' in headers

    def test_member_checks_row_count(self, tmp_path):
        wb, nodes, members, res, checks = self._export_and_open(tmp_path)
        ws = wb['Member Checks']
        assert ws.max_row == 1 + len(members)

    def test_gusset_plates_sheet_exists(self, tmp_path):
        wb, nodes, members, res, checks = self._export_and_open(tmp_path)
        assert 'Gusset Plates' in wb.sheetnames

    def test_gusset_plates_has_required_columns(self, tmp_path):
        wb, nodes, members, res, checks = self._export_and_open(tmp_path)
        ws = wb['Gusset Plates']
        headers = [ws.cell(row=1, column=c).value for c in range(1, 22)]
        assert 'whitmore_mm' in headers
        assert 'block_shear_util' in headers
        assert 'buckling_applies' in headers
        assert 'utilization' in headers
        assert 'governing' in headers

    def test_gusset_plates_has_data(self, tmp_path):
        wb, nodes, members, res, checks = self._export_and_open(tmp_path)
        ws = wb['Gusset Plates']
        assert ws.max_row > 1
