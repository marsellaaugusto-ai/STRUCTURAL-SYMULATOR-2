"""The Stereo tab's use of the scene graph, driven through the real tab.

Two halves, and the split is deliberate.

  stereo_scene_bridge is pure, so most of this file runs with no window at
  all: it converts REAL generated models -- the geometry the tab's own
  generators produce, grouped by the tab's own auto-grouper -- and holds the
  trip through the graph to changing nothing.

  The wiring itself (export, import, the panel's actions) is driven through
  a real StereoApp, the way tests/test_excel_roundtrip.py argues for: a
  hand-written state dict encodes what the model shape was the day the test
  was written and drifts silently, while the real tab's own methods are the
  ones the buttons call.

What is being protected here is not "it does not raise". It is that a
conversion which quietly moved a rod, lost a section, re-attached a group to
different rods, or dropped a support would produce a model that loads
cleanly, analyses cleanly, and is NOT the model the user drew.
"""
import os
import time

import pytest

from apps.stereo import stereo_autogroup as ag
from apps.stereo import stereo_geometry as sg
from apps.stereo import stereo_groups as sgp
from apps.stereo import stereo_scene_bridge as sbr


# ── the conversion, on real generated models ──────────────────────────────

def _sectioned(mesh):
    nodes, members = mesh['nodes'], mesh['members']
    for i, m in enumerate(members):
        m.setdefault('E', 200.0)
        m.setdefault('A', 20.0 + i % 3)
        m.setdefault('conn', 'pin')
    return nodes, members


MESHES = {
    'flat square': lambda: sg.flat_grid(span_x=12.0, span_y=9.0, depth=1.0,
                                        module=3.0, offset=True),
    'flat diagonal': lambda: sg.flat_grid(span_x=9.0, span_y=9.0, depth=1.0,
                                          module=3.0, pattern='diagonal'),
    'barrel vault': lambda: sg.barrel_vault(span=10.0, rise=2.5, length=15.0,
                                            n_arch=6, n_bays=5, depth=1.0),
}


@pytest.mark.parametrize('name', sorted(MESHES))
@pytest.mark.parametrize('localize', [True, False])
def test_a_real_model_survives_the_round_trip_exactly(name, localize):
    nodes, members = _sectioned(MESHES[name]())
    groups = ag.auto_groups(nodes, members)
    supports = [{'node': n, 'type': 'pin'}
                for n in (MESHES[name]()).get('support_candidates', [])[:6]]
    loads = [{'node': 0, 'fz': -5.0}, {'node': len(nodes) - 1, 'fz': -3.0}]
    problems = sbr.compare_round_trip(nodes, members, groups, supports, loads,
                                      localize=localize)
    assert problems == [], problems[:4]


def test_the_auto_grouper_really_does_nest(request):
    """Guards the test above: a round trip through FLAT groups would prove
    much less, and the nesting is what the frames have to compose through."""
    nodes, members = _sectioned(MESHES['flat square']())
    groups = ag.auto_groups(nodes, members)
    assert max(sgp.depth(groups, g['id']) for g in groups) >= 1


def test_a_localized_group_gets_a_frame_at_its_own_centroid():
    """The frames the flat model never had. Without them a group has no
    pivot of its own, which is what makes "move the bay" a graph operation
    rather than a loop over whichever nodes happen to be in it."""
    nodes, members = _sectioned(MESHES['flat square']())
    groups = ag.auto_groups(nodes, members)
    doc = sbr.graph_from_model(nodes, members, groups, localize=True)
    framed = [o for o in doc.root.walk()
              if o.is_group() and not o.local.is_identity()]
    assert framed, 'every group sat at the origin'
    for obj in framed:
        rods = [r for r in obj.walk() if r.KIND == 'rod']
        pts = [p for r in rods for p in (r.a, r.b)]
        # The rods' own coordinates now straddle their group's origin.
        assert min(p[0] for p in pts) <= 0.0 <= max(p[0] for p in pts)

    flat = sbr.graph_from_model(nodes, members, groups, localize=False)
    assert all(o.local.is_identity() for o in flat.root.walk() if o.is_group())


def test_the_member_order_change_is_reported_not_hidden():
    """A bake orders members by tree position, so a conversion renumbers
    them. Reports and saved comparisons name members by index, so a
    renumbering that is not reported points them at the wrong rods."""
    nodes, members = _sectioned(MESHES['flat square']())
    groups = ag.auto_groups(nodes, members)
    doc = sbr.graph_from_model(nodes, members, groups)
    out_nodes, out_members, _groups, remap = sbr.model_from_graph(doc)
    assert set(remap['members']) == set(range(len(members)))
    assert sorted(remap['members'].values()) == list(range(len(out_members)))
    assert remap['members'] != {i: i for i in range(len(members))}, \
        'this model really is reordered, so the map is doing work'
    # And following it lands on the same rod.
    for old, new in remap['members'].items():
        assert out_members[new]['A'] == members[old]['A']


def test_the_provenance_key_does_not_reach_the_solver():
    nodes, members = _sectioned(MESHES['flat square']())
    doc = sbr.graph_from_model(nodes, members, ag.auto_groups(nodes, members))
    _n, out_members, _g, _r = sbr.model_from_graph(doc)
    assert all(sbr.MEMBER_INDEX_KEY not in m for m in out_members)


def test_a_support_follows_its_joint_through_a_renumber():
    """The point of carrying supports as named joints: a model that gained
    a rod at the front keeps its restraints on the right joints."""
    nodes, members = _sectioned(MESHES['flat square']())
    groups = ag.auto_groups(nodes, members)
    supports = [{'node': 0, 'type': 'pin'}, {'node': 5, 'type': 'rollerX'}]
    doc = sbr.graph_from_model(nodes, members, groups, supports=supports)
    out_nodes, _m, _g, remap = sbr.model_from_graph(doc, supports=supports)
    assert len(remap['supports']) == 2
    for was, now in zip(supports, remap['supports']):
        assert out_nodes[now['node']] == pytest.approx(nodes[was['node']])
        assert now['type'] == was['type'], 'the restraint itself is unchanged'


def test_a_support_whose_joint_is_gone_is_reported_not_moved():
    """Dropped and SAID so. A support silently landing on node 0 is a
    different structure that still analyses."""
    nodes, members = _sectioned(MESHES['flat square']())
    doc = sbr.graph_from_model(nodes, members, supports=[{'node': 0}])
    # Take the joint back out, as deleting the rod it sat on would.
    for obj in list(doc.root.children()):
        if obj.KIND == 'joint':
            doc.root.remove(obj)
    _n, _m, _g, remap = sbr.model_from_graph(doc, supports=[{'node': 0}])
    assert remap['supports'] == []
    assert any('dropped' in w for w in remap['warnings'])


def test_a_saved_scene_keeps_each_support_s_own_kind():
    """A roller read back as a pin removes a degree of freedom nobody
    removed, and the model still analyses -- so the wrong answer arrives
    with nothing on screen to say so."""
    nodes, members = _sectioned(MESHES['flat square']())
    supports = [{'node': 0, 'type': 'pin'},
                {'node': 1, 'type': 'rollerX'},
                {'node': 2, 'type': 'fixed'},
                {'node': 3, 'type': 'free', 'dofs': {'uz': True}}]
    loads = [{'node': 7, 'fz': -9.0, 'fx': 1.5}]
    doc = sbr.graph_from_model(nodes, members, supports=supports, loads=loads)
    back_sup, back_loads = sbr.anchors_from_graph(doc)

    assert sorted(s['type'] for s in back_sup) == \
        sorted(s['type'] for s in supports)
    by_type = {s['type']: s for s in back_sup}
    assert by_type['free']['dofs'] == {'uz': True}
    assert not any(s.get('assumed') for s in back_sup)
    assert len(back_loads) == 1
    assert back_loads[0]['fz'] == -9.0 and back_loads[0]['fx'] == 1.5


def test_a_support_anchor_with_no_record_is_flagged_as_assumed():
    """A hand-edited file, or one from before the record travelled. Named,
    rather than a restraint being invented in silence."""
    from apps.stereo.scene import JointPrimitive
    nodes, members = _sectioned(MESHES['flat square']())
    doc = sbr.graph_from_model(nodes, members)
    doc.root.add(JointPrimitive(nodes[0], 'support_0'))
    back_sup, _ = sbr.anchors_from_graph(doc)
    assert len(back_sup) == 1 and back_sup[0]['assumed'] is True


def test_the_support_kinds_survive_a_file_on_disk(tmp_path):
    """Through the real codec, not just in memory."""
    from apps.stereo.scene import load_json, save_json
    nodes, members = _sectioned(MESHES['flat square']())
    supports = [{'node': 0, 'type': 'pin'}, {'node': 4, 'type': 'rollerY'}]
    doc = sbr.graph_from_model(nodes, members,
                               groups=ag.auto_groups(nodes, members),
                               supports=supports)
    path = str(tmp_path / 'anchored.scene.json')
    save_json(doc, path)
    reopened, warnings, _ = load_json(path)
    assert warnings == []
    back_sup, _ = sbr.anchors_from_graph(reopened)
    assert sorted(s['type'] for s in back_sup) == ['pin', 'rollerY']


def test_repeated_parts_finds_what_the_flat_model_cannot_say():
    """Two identical trusses are two sets of rods and nothing says they are
    one part. This is the answer the graph makes possible."""
    nodes, members, groups = [], [], []
    for copy in range(3):
        y = copy * 4.0
        base = len(nodes)
        nodes += [(0, y, 0), (6, y, 0), (3, y, 2)]
        start = len(members)
        members += [{'a': base, 'b': base + 1, 'conn': 'pin', 'A': 20.0},
                    {'a': base, 'b': base + 2, 'conn': 'pin', 'A': 20.0},
                    {'a': base + 2, 'b': base + 1, 'conn': 'pin', 'A': 20.0}]
        sgp.new_group(groups, 'Truss %d' % (copy + 1),
                      members=range(start, len(members)))
    # ...and one that is genuinely a different shape.
    base = len(nodes)
    nodes += [(0, 20, 0), (9, 20, 0)]
    members.append({'a': base, 'b': base + 1, 'conn': 'pin', 'A': 20.0})
    sgp.new_group(groups, 'Odd one', members=[len(members) - 1])

    doc = sbr.graph_from_model(nodes, members, groups)
    found = sbr.instanceable_groups(doc)
    assert len(found) == 1, 'one part, repeated'
    names = sorted(o.name for o in next(iter(found.values())))
    assert names == ['Truss 1', 'Truss 2', 'Truss 3']


def test_repeated_parts_sees_through_position_and_name():
    """Compared by shape alone. A copy moved and renamed is the same part;
    a copy with one rod longer is not."""
    def model(dx, length, name):
        base = len(nodes)
        nodes.extend([(dx, 0, 0), (dx + length, 0, 0), (dx + 3, 0, 2)])
        start = len(members)
        members.extend([{'a': base, 'b': base + 1}, {'a': base, 'b': base + 2},
                        {'a': base + 2, 'b': base + 1}])
        sgp.new_group(groups, name, members=range(start, len(members)))

    nodes, members, groups = [], [], []
    model(0.0, 6.0, 'North truss')
    model(40.0, 6.0, 'T-02')
    model(80.0, 6.5, 'Nearly the same')
    found = sbr.instanceable_groups(sbr.graph_from_model(nodes, members, groups))
    assert len(found) == 1
    assert sorted(o.name for o in next(iter(found.values()))) == \
        ['North truss', 'T-02']


# ── the tab itself ────────────────────────────────────────────────────────

pytest.importorskip('tkinter')
import tkinter as tk                                            # noqa: E402

from apps.stereo.stereo_app import StereoApp                    # noqa: E402


@pytest.fixture(autouse=True)
def dialogs(monkeypatch):
    """Every popup into a list. A modal box with no mainloop behind it does
    not fail a test -- it HANGS, waiting for a click that never comes."""
    seen = []
    for kind in ('showinfo', 'showerror', 'showwarning'):
        for where in ('apps.stereo.stereo_app_scene.messagebox',
                      'apps.stereo.stereo_app.messagebox'):
            monkeypatch.setattr('%s.%s' % (where, kind),
                                lambda *a, _k=kind, **kw: seen.append((_k,) + a))
    return seen


@pytest.fixture(scope='module')
def tk_root():
    last = None
    for attempt in range(6):
        try:
            root = tk.Tk()
            break
        except tk.TclError as exc:
            last = exc
            time.sleep(0.5 * (attempt + 1))
    else:
        pytest.skip('no Tk display after 6 attempts: %s' % last)
    root.geometry('1400x900')
    yield root
    try:
        root.destroy()
    except tk.TclError:
        pass


@pytest.fixture
def app(tk_root):
    tab = tk.Frame(tk_root)
    a = StereoApp(tab)
    a._generate(push_undo=False)
    tab.pack(fill='both', expand=True)
    tk_root.update_idletasks()
    tk_root.update()
    yield a
    tab.destroy()


def test_the_new_mixin_overrides_nothing_the_tab_already_had():
    """A mixin that redefines an existing method is a silent swap: whichever
    comes first in the MRO wins, and the loser's callers keep working until
    the order changes. This module had exactly that -- its own _set_status,
    shadowed by the shell's -- and it was harmless right up until it would
    not have been."""
    from apps.stereo.stereo_app_scene import StereoSceneMixin
    mine = {n for n in vars(StereoSceneMixin) if not n.startswith('__')}
    clashes = {}
    for base in StereoApp.__mro__:
        if base in (StereoSceneMixin, StereoApp, object):
            continue
        hit = mine & {n for n in vars(base) if not n.startswith('__')}
        if hit:
            clashes[base.__name__] = sorted(hit)
    assert clashes == {}, clashes


def test_the_tab_builds_a_document_from_its_own_model(app):
    app.groups = ag.auto_groups(app.nodes, app.members)
    doc = app._scene_document()
    baked = doc.bake()
    assert len(baked.members) == len(app.members)
    assert not baked.warnings


def test_the_document_is_derived_so_it_cannot_go_stale(app):
    """The reason it is not mirrored: self.groups is assigned in eleven
    places, and a stored document missing one would still bake."""
    app.groups = ag.auto_groups(app.nodes, app.members)
    before = len(app._scene_document().root.children())
    app.groups = []                       # as any of those eleven would
    after = app._scene_document()
    assert len(after.root.children()) != before
    assert all(not o.is_group() or o is after.root
               for o in after.root.children())


def test_the_tab_round_trips_its_own_model_through_the_graph(app):
    app.groups = ag.auto_groups(app.nodes, app.members)
    assert app._scene_self_check() == []


def test_an_exported_workbook_carries_the_scene_and_brings_it_back(app, tmp_path):
    """The wiring, end to end, through the tab's own export and import."""
    pytest.importorskip('openpyxl')
    import openpyxl
    from apps.stereo import stereo_reports as sr

    app.groups = ag.auto_groups(app.nodes, app.members)
    names_before = sorted(g['name'] for g in app.groups)
    rods_before = len(app.members)
    path = str(tmp_path / 'with_scene.xlsx')
    sr.export_excel(app.nodes, app.members, app._all_loads(), app.supports,
                    None, path, groups=app.groups,
                    scene=app._scene_document())

    wb = openpyxl.load_workbook(path)
    assert 'Scene' in wb.sheetnames, 'the sheet is in the workbook'
    assert 'Model' in wb.sheetnames and 'Groups' in wb.sheetnames, \
        'and it did not displace the sheets that were already there'

    groups = app._groups_from_scene_sheet(path, app.members)
    assert groups is not None
    assert sorted(g['name'] for g in groups) == names_before
    # Every rod accounted for exactly once, against the ORIGINAL numbering.
    assigned = [i for g in groups for i in g['members']]
    assert len(assigned) == len(set(assigned))
    assert set(assigned) <= set(range(rods_before))
    by_rods = lambda gs: sorted(tuple(sorted(g['members'])) for g in gs
                                if g['members'])
    assert by_rods(groups) == by_rods(app.groups), \
        'the groups came back on the same rods they went out on'


def test_a_workbook_with_no_scene_sheet_changes_nothing(app, tmp_path):
    """Every workbook written before this existed must import as it did."""
    pytest.importorskip('openpyxl')
    from apps.stereo import stereo_reports as sr
    path = str(tmp_path / 'old_style.xlsx')
    sr.export_excel(app.nodes, app.members, app._all_loads(), app.supports,
                    None, path, groups=app.groups)       # no scene=
    assert app._groups_from_scene_sheet(path, app.members) is None


def test_a_scene_sheet_for_a_different_model_is_refused(app, tmp_path):
    """A hand-edited workbook whose sheets disagree about how many rods
    there are: refused, rather than attaching groups to the wrong rods."""
    pytest.importorskip('openpyxl')
    from apps.stereo import stereo_reports as sr
    app.groups = ag.auto_groups(app.nodes, app.members)
    path = str(tmp_path / 'mismatch.xlsx')
    sr.export_excel(app.nodes, app.members, app._all_loads(), app.supports,
                    None, path, groups=app.groups,
                    scene=app._scene_document())
    assert app._groups_from_scene_sheet(path, app.members[:-5]) is None


def test_the_tab_saves_and_opens_a_scene_file(app, tmp_path, monkeypatch):
    from apps.stereo.scene import SUFFIX
    app.groups = ag.auto_groups(app.nodes, app.members)
    names_before = sorted(g['name'] for g in app.groups)
    rods_before, nodes_before = len(app.members), len(app.nodes)
    # A mixture, so a round trip that flattened them all to pins would show.
    for i, kind in enumerate(('pin', 'rollerX', 'fixed')):
        if i < len(app.supports):
            app.supports[i]['type'] = kind
    supports_before = [dict(s) for s in app.supports]
    assert len({s['type'] for s in supports_before}) > 1, 'a real mixture'
    path = str(tmp_path / ('model' + SUFFIX))

    monkeypatch.setattr('apps.stereo.stereo_app_scene.filedialog'
                        '.asksaveasfilename', lambda **kw: path)
    app._save_scene_file()
    assert os.path.exists(path)

    app._clear_model()
    assert not app.members
    monkeypatch.setattr('apps.stereo.stereo_app_scene.filedialog'
                        '.askopenfilename', lambda **kw: path)
    app._open_scene_file()
    assert len(app.members) == rods_before
    assert len(app.nodes) == nodes_before
    assert sorted(g['name'] for g in app.groups) == names_before
    assert app._model_label == os.path.basename(path)
    # And the supports came back as the kinds they were, not as pins.
    assert sorted(s.get('type') for s in app.supports) == \
        sorted(s.get('type') for s in supports_before)
    assert not any(s.get('assumed') for s in app.supports)


def test_an_opened_scene_still_solves(app, tmp_path, monkeypatch):
    from apps.stereo import stereo_math as sm
    from apps.stereo.scene import SUFFIX
    app.groups = ag.auto_groups(app.nodes, app.members)
    path = str(tmp_path / ('solvable' + SUFFIX))
    monkeypatch.setattr('apps.stereo.stereo_app_scene.filedialog'
                        '.asksaveasfilename', lambda **kw: path)
    app._save_scene_file()
    monkeypatch.setattr('apps.stereo.stereo_app_scene.filedialog'
                        '.askopenfilename', lambda **kw: path)
    app._open_scene_file()
    res, err = sm.analyze(app.nodes, app.members,
                          [{'node': 0, 'fz': -5.0}], app.supports)
    assert err is None, err
    assert len(res['member_res']) == len(app.members)


def test_the_repeated_parts_window_opens_and_says_something(app, dialogs):
    app.groups = ag.auto_groups(app.nodes, app.members)
    app._group_repeated_parts()
    kids = [w for w in app.root.winfo_children()
            if isinstance(w, tk.Toplevel)]
    assert kids, 'a window opened'
    win = kids[-1]
    texts = [w for w in win.winfo_children() if isinstance(w, tk.Text)]
    assert texts
    body = texts[0].get('1.0', 'end')
    assert body.strip(), 'and it has something in it'
    win.destroy()


def test_repeated_parts_with_no_groups_says_so_rather_than_opening(app, dialogs):
    app.groups = []
    app._group_repeated_parts()
    assert any(kind == 'showinfo' for kind, *_ in dialogs)
    assert not [w for w in app.root.winfo_children()
                if isinstance(w, tk.Toplevel)]


def test_an_export_survives_a_scene_that_cannot_be_built(app, tmp_path):
    """The workbook is the deliverable; the Scene sheet is an addition to
    it. A conversion that raises must not cost the user their export."""
    pytest.importorskip('openpyxl')
    import openpyxl
    from apps.stereo import stereo_reports as sr

    class Exploding:
        def bake(self, **kw):
            raise RuntimeError('no')

    path = str(tmp_path / 'still_written.xlsx')
    sr.export_excel(app.nodes, app.members, app._all_loads(), app.supports,
                    None, path, groups=app.groups, scene=Exploding())
    wb = openpyxl.load_workbook(path)
    assert 'Model' in wb.sheetnames
    assert 'Scene' not in wb.sheetnames


# ── what a flat group record holds besides its rods ──────────────────────

def test_a_groups_own_facts_survive_the_trip_through_the_graph():
    """A group record is not only rods and a parent. It also says whether
    it is left out of the analysis, which iteration it belongs to, and
    which fabricated part it is a copy of -- and the graph used to drop
    every one of them.

    "excluded" is the one that matters: a group left out of the analysis
    coming back INSIDE it changes the structure that gets solved, and
    nothing on screen says so.
    """
    from apps.stereo.stereo_groups_excel import GROUP_FLAGS
    nodes, members = _sectioned(MESHES['flat square']())
    groups = ag.auto_groups(nodes, members)
    assert groups, 'the auto-grouper found something to group'
    marked = groups[0]
    marked['excluded'] = True
    marked['iteration'] = '2'
    marked['stage'] = 'lift'
    marked['position'] = 'north'
    marked['component'] = 'Gable truss'

    doc = sbr.graph_from_model(nodes, members, groups)
    back = doc.bake().legacy_groups(doc.root)
    same = [g for g in back if g['name'] == marked['name']]
    assert same, 'the group came back'
    meta = same[0].get('meta') or {}
    for key in GROUP_FLAGS:
        assert meta.get(key) == marked[key], key


@pytest.mark.parametrize('flag,value', [
    ('excluded', True), ('iteration', '3'), ('stage', 'erect'),
    ('position', 'south'), ('component', 'Bay A'),
])
def test_the_scene_sheet_brings_a_groups_facts_back(app, tmp_path, flag,
                                                     value, monkeypatch):
    """The Scene sheet is preferred over the Groups sheet on import
    because it is the richer record. It was coming back poorer: the
    nesting and the frames survived and everything else did not."""
    pytest.importorskip('openpyxl')
    app.groups = ag.auto_groups(app.nodes, app.members)
    assert app.groups
    name = app.groups[0]['name']
    app.groups[0][flag] = value

    path = str(tmp_path / 'facts.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.askopenfilename',
                        lambda **kw: path)
    app._import_excel()

    same = [g for g in app.groups if g['name'] == name]
    assert same, 'the group came back'
    assert same[0].get(flag) == value


def test_no_group_comes_back_carrying_the_raw_metadata(app, tmp_path,
                                                        monkeypatch):
    """`meta` is the envelope the facts travelled in, not a field of a
    group record. Leaving it on would put a dict into every snapshot and
    every sheet that walks a group's keys."""
    pytest.importorskip('openpyxl')
    app.groups = ag.auto_groups(app.nodes, app.members)
    path = str(tmp_path / 'clean.xlsx')
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.asksaveasfilename',
                        lambda **kw: path)
    app._export_excel()
    monkeypatch.setattr('apps.stereo.stereo_app.filedialog.askopenfilename',
                        lambda **kw: path)
    app._import_excel()
    assert all('meta' not in g for g in app.groups)
