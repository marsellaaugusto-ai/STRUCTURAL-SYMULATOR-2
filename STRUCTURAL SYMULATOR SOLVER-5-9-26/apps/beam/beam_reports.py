"""beam_reports.py — the Beam tab's Excel report and model import.

Split out of `beam_app.py` with the solver on 2026-10-04 (R-8), mirroring
`apps/truss/truss_reports.py`: the workbook layer owns openpyxl and nothing
else owns the workbook layer. Unlike `beam_math.py` this is NOT on the
math/UI boundary -- it reaches `common` for the openpyxl bootstrap, so it
needs the same environment the tab does, exactly as truss_reports does.

Both functions are a matched pair: `export_beam_excel` writes a 'Model' sheet
holding every input needed to rebuild the beam, and `import_beam_excel` reads
that sheet back. tests/test_excel_roundtrip.py holds them to it field by
field.

UNITS. The workbook is written in the tab's STORAGE units (m, kN, kN*m, cm^4,
GPa) whatever the Units selector shows, so a saved file never changes meaning
because someone picked a different convention -- see `units.py`.
"""
from common import _ensure_openpyxl
from .beam_math import BeamModel


def export_beam_excel(state, path, result=None, model=None):
    """
    Writes a workbook with:
      • 'Model'   — every input needed to rebuild this exact beam (length,
                    supports, loads, section) — paired with import_beam_excel().
      • 'Results' — reactions, extremes, and a station-by-station table of
                    V, M, and deflection (only if `result` given).
    `state` is a plain dict with keys: length, supports, point_loads,
    moments, dloads, nonuniform_loads, profile.
    """
    if not _ensure_openpyxl():
        raise RuntimeError('openpyxl is not available.')
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('Model')
    row = 1
    ws.cell(row=row, column=1, value='BEAM MODEL DATA — for Import from Excel (do not reorder columns)')
    ws.cell(row=row, column=1).font = Font(bold=True, size=11, color='1F4E79')
    row += 1
    # Which units these numbers are in. The sheet is ALWAYS written in the
    # tab's storage units, deliberately: a saved file must not change meaning
    # because someone picked a different convention in the selector (see
    # units.py). That is only safe if the file says so (R-17).
    ws.cell(row=row, column=1,
            value='Written in the app\'s storage units (m, kN, kN/m, kN·m, '
                  'GPa, cm², cm⁴, kN/cm²) — NOT in whatever the Units '
                  'selector shows.')
    row += 2

    ws.cell(row=row, column=1, value='[GEOMETRY]'); row += 1
    ws.cell(row=row, column=1, value='length_m'); row += 1
    ws.cell(row=row, column=1, value=state['length']); row += 2

    ws.cell(row=row, column=1, value='[SUPPORTS]'); row += 1
    for col, lbl in enumerate(['x_m', 'type'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for s in state['supports']:
        ws.cell(row=row, column=1, value=s['x'])
        ws.cell(row=row, column=2, value=s['type'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[POINT_LOADS]'); row += 1
    for col, lbl in enumerate(['x_m', 'P_kN'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for p in state['point_loads']:
        ws.cell(row=row, column=1, value=p['x'])
        ws.cell(row=row, column=2, value=p['P'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[MOMENTS]'); row += 1
    for col, lbl in enumerate(['x_m', 'M_kNm'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for mm in state['moments']:
        ws.cell(row=row, column=1, value=mm['x'])
        ws.cell(row=row, column=2, value=mm['M'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[DISTRIBUTED_LOADS]'); row += 1
    # 'w1_kNm' read as kN·m to anyone filling this in by hand; it is a load
    # per unit length. import_beam_excel accepts the old spelling, exactly as
    # it does for the [DLOADS] -> [DISTRIBUTED_LOADS] rename (R-17).
    for col, lbl in enumerate(['x1_m', 'x2_m', 'w1_kN_per_m', 'w2_kN_per_m'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for d in state['dloads']:
        ws.cell(row=row, column=1, value=d['x1'])
        ws.cell(row=row, column=2, value=d['x2'])
        ws.cell(row=row, column=3, value=d['w1'])
        ws.cell(row=row, column=4, value=d['w2'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[NONUNIFORM_LOADS]'); row += 1
    for col, lbl in enumerate(['expr', 'x1_m', 'x2_m'], 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for d in state['nonuniform_loads']:
        ws.cell(row=row, column=1, value=d['expr'])
        ws.cell(row=row, column=2, value=d['x1'])
        ws.cell(row=row, column=3, value=d['x2'])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='[SECTION]'); row += 1
    keys = ['E', 'I', 'c', 'A', 'allow_bend', 'allow_shear', 'defl_ratio']
    labels = ['E_GPa', 'I_cm4', 'c_cm', 'A_cm2', 'allow_bend_kNcm2',
              'allow_shear_kNcm2', 'defl_limit_L_over_n']
    for col, lbl in enumerate(labels, 1):
        ws.cell(row=row, column=col, value=lbl)
    row += 1
    for col, key in enumerate(keys, 1):
        ws.cell(row=row, column=col, value=state['profile'][key])
    row += 1

    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 16

    if result is not None and model is not None:
        wr = wb.create_sheet('Results')
        wr.cell(row=1, column=1, value='BEAM RESULTS').font = Font(bold=True, size=11, color='1F4E79')

        row_r = 3
        wr.cell(row=row_r, column=1, value='Support')
        wr.cell(row=row_r, column=2, value='Ry_kN (+up)')
        wr.cell(row=row_r, column=3, value='M_kNm (couple, +CCW)')
        wr.cell(row=row_r, column=4, value='M_left_kNm (internal)')
        wr.cell(row=row_r, column=5, value='M_right_kNm (internal)')
        row_r += 1
        for s in model.supports:
            # No sign heuristic here either: the couple is the couple at every
            # station, and the internal moments are reported beside it (R-3).
            react = result.support_reaction(s['x'])
            wr.cell(row=row_r, column=1, value=f"x={s['x']:.3f} ({s['type']})")
            wr.cell(row=row_r, column=2, value=react['Fy'] / 1e3)
            wr.cell(row=row_r, column=3, value=react['M'] / 1e3)
            wr.cell(row=row_r, column=4, value=react['M_left'] / 1e3)
            wr.cell(row=row_r, column=5, value=react['M_right'] / 1e3)
            row_r += 1

        diag = result.sample_diagram(n_per_element=25)
        Vmax = max(diag['V'], key=abs) if diag['V'] else 0.0
        Mmax = max(diag['M'], key=abs) if diag['M'] else 0.0
        vmax = max(diag['v'], key=abs) if diag['v'] else 0.0
        row_r += 1
        wr.cell(row=row_r, column=1, value='Max |V| (kN)'); wr.cell(row=row_r, column=2, value=abs(Vmax) / 1e3)
        row_r += 1
        wr.cell(row=row_r, column=1, value='Max |M| (kN·m)'); wr.cell(row=row_r, column=2, value=abs(Mmax) / 1e3)
        row_r += 1
        wr.cell(row=row_r, column=1, value='Max |defl| (mm)'); wr.cell(row=row_r, column=2, value=abs(vmax) * 1000)
        row_r += 1

        # Does the answer balance? The reviewer who opens this workbook is the
        # reader most likely to want it, and the sheet said nothing about it
        # until 2026-10-04 -- see BeamResult.equilibrium() and R-2.
        eq = result.equilibrium()
        row_r += 1
        wr.cell(row=row_r, column=1, value='EQUILIBRIUM').font = Font(bold=True)
        row_r += 1
        for label, value in [
            ('Total load (kN, down)', eq['applied_down'] / 1e3),
            ('Sum of reactions (kN, up)', eq['reactions_up'] / 1e3),
            ('Residual sum Fy (kN)', eq['residual_Fy'] / 1e3),
            ('Residual sum M about x=0 (kN*m)', eq['residual_M0'] / 1e3),
            ('V beyond x=L (kN)', eq['shear_beyond_end'] / 1e3),
            ('M beyond x=L (kN*m)', eq['moment_beyond_end'] / 1e3),
            ('Restrained DOF', eq['constrained_dof']),
            ('Degree of indeterminacy', max(0, eq['indeterminacy'])),
            ('Balanced', 'yes' if eq['ok'] else 'NO -- DO NOT USE'),
        ]:
            wr.cell(row=row_r, column=1, value=label)
            wr.cell(row=row_r, column=2, value=value)
            row_r += 1

        for mg in getattr(model, 'support_merges', []):
            wr.cell(row=row_r, column=1, value='Note')
            wr.cell(row=row_r, column=2,
                    value=f"two supports at x={mg['x']:.3f} m merged "
                          f"({mg['kept']} + {mg['added']} -> {mg['result']})")
            row_r += 1

        hdr_row = row_r + 2
        for col, lbl in enumerate(['x_m', 'V_kN', 'M_kNm', 'defl_mm'], 1):
            wr.cell(row=hdr_row, column=col, value=lbl)
        for i in range(len(diag['x'])):
            rr = hdr_row + 1 + i
            wr.cell(row=rr, column=1, value=diag['x'][i])
            wr.cell(row=rr, column=2, value=diag['V'][i] / 1e3)
            wr.cell(row=rr, column=3, value=diag['M'][i] / 1e3)
            wr.cell(row=rr, column=4, value=diag['v'][i] * 1000)
        for col in range(1, 5):
            wr.column_dimensions[get_column_letter(col)].width = 13

    wb.save(path)


def import_beam_excel(path):
    """Reads a 'Model' sheet written by export_beam_excel() and returns a
    plain state dict (same shape as the `state` argument passed into
    export_beam_excel), or raises ValueError with a human-readable message."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if 'Model' not in wb.sheetnames:
        raise ValueError('This workbook has no "Model" sheet — it was not '
                          'exported by the Beam tab.')
    ws = wb['Model']
    rows = list(ws.iter_rows(values_only=True))

    def find_section(name):
        for i, r in enumerate(rows):
            if r and r[0] == name:
                return i
        return -1

    def read_table(start_idx):
        hdr_i = start_idx + 1
        if hdr_i >= len(rows) or not rows[hdr_i] or rows[hdr_i][0] is None:
            return []
        headers = [h for h in rows[hdr_i] if h is not None]
        out = []
        i = hdr_i + 1
        while i < len(rows) and rows[i] and rows[i][0] is not None and not str(rows[i][0]).startswith('['):
            vals = rows[i]
            # A row with fewer cells than its header used to raise IndexError
            # with no row number at all. Missing cells read as None and are
            # reported by name and row below (R-17). `row` is the sheet's own
            # 1-based number, which is what the user sees in Excel.
            out.append(dict({headers[j]: (vals[j] if j < len(vals) else None)
                             for j in range(len(headers))},
                            row=i + 1))
            i += 1
        return out

    def number(r, *names, what=None):
        """One cell as a number, or a ValueError naming the column and row.

        Nothing validated these before 2026-10-04: a hand-edited workbook
        failed at Analyze instead, far from the file that caused it, and a
        blank cell raised `float(None)` with no row number (R-17). `names`
        takes more than one spelling so a renamed column can keep reading old
        files.
        """
        for name in names:
            if name in r:
                value = r[name]
                break
        else:
            raise ValueError(
                f"Row {r['row']}: no '{names[0]}' column in this section.")
        if value is None or value == '':
            raise ValueError(f"Row {r['row']}: '{names[0]}' is empty.")
        try:
            return float(value)
        except (TypeError, ValueError):
            raise ValueError(
                f"Row {r['row']}: '{names[0]}' is not a number "
                f"({value!r}).") from None

    def on_beam(r, value, what):
        tol = 1e-9 * max(1.0, abs(length))
        if not (-tol <= value <= length + tol):
            raise ValueError(
                f"Row {r['row']}: {what} at x = {value:g} m is not on the "
                f"beam, which spans 0 to {length:g} m.")
        return min(max(value, 0.0), length)

    gi = find_section('[GEOMETRY]')
    if gi < 0:
        raise ValueError('Missing [GEOMETRY] section.')
    grow = read_table(gi)
    if not grow:
        raise ValueError('Empty [GEOMETRY] section.')
    length = number(grow[0], 'length_m')
    if not (length > 0) or length != length:        # NaN-safe
        raise ValueError(
            f"Row {grow[0]['row']}: the beam length must be greater than "
            f"zero; this workbook says {length:g}.")

    supports = []
    si = find_section('[SUPPORTS]')
    if si >= 0:
        for r in read_table(si):
            kind = str(r.get('type') or '').strip().lower()
            if kind not in BeamModel.SUPPORT_DOF:
                raise ValueError(
                    f"Row {r['row']}: {r.get('type')!r} is not a support "
                    f"type. Use one of: "
                    f"{', '.join(sorted(BeamModel.SUPPORT_DOF))}.")
            supports.append({'x': on_beam(r, number(r, 'x_m'), 'Support'),
                             'type': kind})

    point_loads = []
    pi = find_section('[POINT_LOADS]')
    if pi >= 0:
        for r in read_table(pi):
            point_loads.append({'x': on_beam(r, number(r, 'x_m'), 'Point load'),
                                'P': number(r, 'P_kN')})

    moments = []
    mi = find_section('[MOMENTS]')
    if mi >= 0:
        for r in read_table(mi):
            moments.append({'x': on_beam(r, number(r, 'x_m'), 'Applied moment'),
                            'M': number(r, 'M_kNm')})

    dloads = []
    # Accept the pre-2026-09-07 name too, so workbooks already on disk
    # still import. New exports use [DISTRIBUTED_LOADS], matching Cable.
    di = find_section('[DISTRIBUTED_LOADS]')
    if di < 0:
        di = find_section('[DLOADS]')
    if di >= 0:
        for r in read_table(di):
            # w1_kN_per_m since 2026-10-04; w1_kNm is the old, misleading
            # spelling of the same column.
            dloads.append({
                'x1': on_beam(r, number(r, 'x1_m'), 'Distributed load start'),
                'x2': on_beam(r, number(r, 'x2_m'), 'Distributed load end'),
                'w1': number(r, 'w1_kN_per_m', 'w1_kNm'),
                'w2': number(r, 'w2_kN_per_m', 'w2_kNm')})

    nonuniform_loads = []
    ni = find_section('[NONUNIFORM_LOADS]')
    if ni >= 0:
        for r in read_table(ni):
            expr = str(r.get('expr') or '').strip()
            if not expr:
                raise ValueError(f"Row {r['row']}: the q(x) expression is "
                                 f"empty.")
            nonuniform_loads.append({
                'expr': expr,
                'x1': on_beam(r, number(r, 'x1_m'), 'Non-uniform load start'),
                'x2': on_beam(r, number(r, 'x2_m'), 'Non-uniform load end')})

    profile = {'E': 200.0, 'I': 8000.0, 'c': 15.0, 'A': 80.0,
               'allow_bend': 16.0, 'allow_shear': 10.0, 'defl_ratio': 360.0}
    seci = find_section('[SECTION]')
    if seci >= 0:
        srow = read_table(seci)
        if srow:
            s0 = srow[0]
            # `or default` keeps a blank cell at its default rather than
            # reading it as zero -- which for an allowable would have meant
            # "not checked" and for E a beam with no stiffness.
            for key, column in (('E', 'E_GPa'), ('I', 'I_cm4'),
                                ('c', 'c_cm'), ('A', 'A_cm2'),
                                ('allow_bend', 'allow_bend_kNcm2'),
                                ('allow_shear', 'allow_shear_kNcm2'),
                                ('defl_ratio', 'defl_limit_L_over_n')):
                if s0.get(column) not in (None, ''):
                    profile[key] = number(s0, column)

    return {'length': length, 'supports': supports, 'point_loads': point_loads,
            'moments': moments, 'dloads': dloads, 'nonuniform_loads': nonuniform_loads,
            'profile': profile}
