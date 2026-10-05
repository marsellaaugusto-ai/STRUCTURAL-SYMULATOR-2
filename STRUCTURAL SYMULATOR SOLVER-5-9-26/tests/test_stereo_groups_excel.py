"""Editing groups in the workbook (roadmap 4.2): export, change a row, import.

Driven through real .xlsx files: the workbook is written by export_excel,
edited cell by cell the way a reader edits it in Excel, and read back by the
same two calls Import from Excel makes.
"""
import copy

import openpyxl
import pytest

from apps.stereo import stereo_groups as sgp
from apps.stereo import stereo_groups_excel as sge
from apps.stereo import stereo_profiles as sp
from apps.stereo import stereo_reports as sr


def _bar(a, b, **kw):
    m = dict(a=a, b=b, conn='pin', E=200.0, A=20.0, I=400.0, J=400.0,
             Fy=235.0, Fu=360.0, K=1.0, r_gyr=3.0, profile='Default web')
    m.update(kw)
    return m


def _model():
    """A chain of 12 rods. Roof = 0-5 (with Bay = 2-3 inside it), Edge = 8-9,
    and 6, 7, 10, 11 Ungrouped. Rod 5 carries a different A, so Roof's own
    rods disagree on A."""
    nodes = [(float(i), 0.0, 0.0) for i in range(13)]
    members = [_bar(i, i + 1) for i in range(12)]
    members[5]['A'] = 30.0
    groups = []
    roof = sgp.new_group(groups, 'Roof', members=[0, 1, 2, 3, 4, 5])
    sgp.new_group(groups, 'Bay', parent=roof['id'], members=[2, 3])
    sgp.new_group(groups, 'Edge', members=[8, 9])
    supports = [{'node': 0, 'type': 'pin'}, {'node': 12, 'type': 'pin'}]
    return nodes, members, groups, supports


def _export(tmp_path, nodes, members, groups, supports):
    path = str(tmp_path / 'model.xlsx')
    sr.export_excel(nodes, members, [], supports, None, path, groups=groups)
    return path


def _import(path):
    nodes, members, loads, supports, profiles = sr.import_excel_model(path)
    groups, report = sge.import_groups(path, members, profiles)
    return members, groups, report


def _edit(path, name, **cells):
    """Set cells of the row whose name is `name`, as a reader would."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sge.SHEET]
    rows = list(ws.iter_rows())
    hdr = next(r for r in rows if r[0].value == 'id')
    col = {c.value: c.column for c in hdr}
    for r in rows:
        if r[1].value == name:
            for k, v in cells.items():
                ws.cell(row=r[0].row, column=col[k], value=v)
            break
    else:
        raise AssertionError('no row %r' % name)
    wb.save(path)


def _add_row(path, **cells):
    wb = openpyxl.load_workbook(path)
    ws = wb[sge.SHEET]
    hdr = next(r for r in ws.iter_rows() if r[0].value == 'id')
    col = {c.value: c.column for c in hdr}
    row = ws.max_row + 1
    for k, v in cells.items():
        ws.cell(row=row, column=col[k], value=v)
    wb.save(path)


# ── rod lists as text ─────────────────────────────────────────────────────

def test_rod_lists_are_written_as_ranges_and_read_back():
    assert sge.rods_to_ranges([0, 1, 2, 3, 7, 9, 10]) == '0-3, 7, 9-10'
    assert sge.parse_ranges('0-3, 7, 9-10') == [0, 1, 2, 3, 7, 9, 10]
    assert sge.rods_to_ranges([]) == ''


def test_rod_lists_accept_what_a_person_types():
    assert sge.parse_ranges('4; 2  3') == [2, 3, 4]
    assert sge.parse_ranges(7.0) == [7], 'Excel hands a lone number as float'
    assert sge.parse_ranges(None) == []
    for bad in ('5-2', 'abc', '1-x', 7.5):
        with pytest.raises(ValueError):
            sge.parse_ranges(bad)


# ── the round trip ────────────────────────────────────────────────────────

def test_an_unedited_workbook_changes_nothing(tmp_path):
    nodes, members, groups, supports = _model()
    before = copy.deepcopy(members)
    path = _export(tmp_path, nodes, members, groups, supports)
    got, g2, report = _import(path)
    assert report == []
    for a, b in zip(before, got):
        for k in ('E', 'A', 'I', 'J', 'Fy', 'Fu', 'K', 'r_gyr', 'conn',
                  'profile'):
            assert a.get(k) == pytest.approx(b.get(k)) if isinstance(
                a.get(k), float) else a.get(k) == b.get(k), k


def test_the_groups_come_back_with_their_rods_and_nesting(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _got, g2, _r = _import(path)
    assert [(g['name'], sorted(g['members'])) for g in g2] == \
        [('Roof', [0, 1, 4, 5]), ('Bay', [2, 3]), ('Edge', [8, 9])]
    bay = next(g for g in g2 if g['name'] == 'Bay')
    roof = next(g for g in g2 if g['name'] == 'Roof')
    assert bay['parent'] == roof['id']
    assert sgp.ungrouped_rods(g2, 12) == [6, 7, 10, 11]


def test_a_mixed_value_is_exported_blank_and_named(tmp_path):
    nodes, members, groups, supports = _model()
    rows = {r['name']: r for r in sge.group_rows(groups, members)}
    assert 'A' in rows['Roof']['mixed']
    assert 'A' not in rows['Roof']['values']
    assert rows['Bay']['values']['A'] == 20.0


def test_a_workbook_without_a_groups_sheet_is_not_a_workbook_without_groups(
        tmp_path):
    nodes, members, _groups, supports = _model()
    path = _export(tmp_path, nodes, members, [], supports)
    assert _import(path)[1:] == (None, [])


# ── editing a row ─────────────────────────────────────────────────────────

def test_a_value_typed_in_a_row_goes_on_every_rod_of_that_group(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, 'Roof', A_cm2=12.5)
    got, _g, report = _import(path)
    assert [got[i]['A'] for i in (0, 1, 4, 5)] == [12.5] * 4
    assert got[2]['A'] == 20.0, 'the subgroup keeps its own row'
    assert any('Roof: A_cm2 mixed -> 12.5 on 4 rod(s)' == ln for ln in report)


def test_a_catalog_profile_sets_the_whole_section_and_numbers_go_on_top(
        tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, 'Edge', profile='IPE 200', K=0.8)
    got, _g, report = _import(path)
    props = sp.section_to_props(sp.CATALOG['IPE 200'])
    for i in (8, 9):
        assert got[i]['profile'] == 'IPE 200'
        assert got[i]['A'] == pytest.approx(props['A'])
        assert got[i]['r_gyr'] == pytest.approx(props['r_gyr'])
        assert got[i]['c_cm'] == pytest.approx(props['c_cm'])
        assert got[i]['K'] == 0.8
    assert got[7]['profile'] == 'Default web'


def test_changing_I_without_a_depth_drops_the_stale_depth(tmp_path):
    nodes, members, groups, supports = _model()
    for i in (8, 9):
        members[i]['c_cm'] = 5.0
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, 'Edge', I_cm4=900.0)
    got, _g, report = _import(path)
    assert all('c_cm' not in got[i] for i in (8, 9))
    assert any('lost their catalog depth' in ln for ln in report)


def test_changing_I_with_its_depth_keeps_the_new_depth(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, 'Edge', I_cm4=900.0, c_cm=7.5)
    got, _g, _r = _import(path)
    assert [got[i]['c_cm'] for i in (8, 9)] == [7.5, 7.5]


def test_moving_a_rod_number_to_another_row_moves_it(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, 'Edge', rods='8-9, 11')
    _got, g2, _r = _import(path)
    assert next(g for g in g2 if g['name'] == 'Edge')['members'] == {8, 9, 11}


def test_a_new_row_makes_a_new_group(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _add_row(path, name='Tie', rods='6-7', conn='rigid')
    got, g2, report = _import(path)
    tie = next(g for g in g2 if g['name'] == 'Tie')
    assert tie['members'] == {6, 7}
    assert tie['id'] not in {g['id'] for g in g2 if g is not tie}
    assert got[6]['conn'] == 'rigid' and got[7]['conn'] == 'rigid'


def test_deleting_every_row_ungroups_everything(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    wb = openpyxl.load_workbook(path)
    ws = wb[sge.SHEET]
    hdr = next(r for r in ws.iter_rows() if r[0].value == 'id')[0].row
    ws.delete_rows(hdr + 1, ws.max_row)
    wb.save(path)
    _got, g2, _r = _import(path)
    assert g2 == []


def test_an_unknown_profile_is_reported_and_changes_nothing(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, 'Edge', profile='IPE 9999')
    got, _g, report = _import(path)
    assert got[8]['profile'] == 'Default web'
    assert any('IPE 9999' in ln and 'left as it was' in ln for ln in report)


def test_a_rod_that_must_stay_rigid_stays_rigid(tmp_path):
    nodes, members, groups, supports = _model()
    for i in (8, 9):
        members[i]['conn'] = 'rigid'
    members[9]['rigid_required'] = True
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, 'Edge', conn='pin')
    got = sr.import_excel_model(path)[1]
    got[9]['rigid_required'] = True          # not a workbook column
    g, report = sge.import_groups(path, got)
    assert got[8]['conn'] == 'pin' and got[9]['conn'] == 'rigid'
    assert any('must stay rigid' in ln for ln in report)


# ── what the sheet must not be allowed to do ──────────────────────────────

@pytest.mark.parametrize('name, cells, says', [
    ('Edge', {'rods': '8-9, 0'}, 'two groups'),
    ('Edge', {'rods': '8-9, 99'}, 'the model has rods 0-11'),
    ('Edge', {'parent': 77}, 'not a row'),
    ('Edge', {'A_cm2': -3}, 'not a section value'),
    ('Edge', {'A_cm2': 'twelve'}, 'not a number'),
    ('Edge', {'rods': '9-8'}, 'backwards'),
])
def test_a_sheet_that_does_not_make_sense_is_refused_with_a_reason(
        tmp_path, name, cells, says):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, name, **cells)
    with pytest.raises(ValueError) as exc:
        _import(path)
    assert says in str(exc.value)


def test_a_group_nested_inside_itself_is_refused(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    roof_id = next(g['id'] for g in groups if g['name'] == 'Roof')
    bay_id = next(g['id'] for g in groups if g['name'] == 'Bay')
    _edit(path, 'Roof', parent=bay_id)
    with pytest.raises(ValueError) as exc:
        _import(path)
    assert 'inside itself' in str(exc.value)
    assert roof_id != bay_id


def test_a_new_profile_and_an_edited_number_in_one_row_both_land(tmp_path):
    nodes, members, groups, supports = _model()
    path = _export(tmp_path, nodes, members, groups, supports)
    _edit(path, 'Edge', profile='HEA 200', A_cm2=50.0)
    got, _g, _r = _import(path)
    props = sp.section_to_props(sp.CATALOG['HEA 200'])
    assert [got[i]['A'] for i in (8, 9)] == [50.0, 50.0], 'the edit wins'
    assert got[8]['I'] == pytest.approx(props['I']), 'the rest is catalog'


def test_catalog_values_are_shown_rounded_but_stored_whole(tmp_path):
    """Rounding the stored number would turn every untouched catalog cell
    into an edit on re-import."""
    nodes, members, groups, supports = _model()
    from apps.stereo import stereo_checks as sk
    sk.apply_recommendation(members, [8, 9], 'CHS 76.1x3.6')
    path = _export(tmp_path, nodes, members, groups, supports)
    ws = openpyxl.load_workbook(path)[sge.SHEET]
    hdr = next(r for r in ws.iter_rows() if r[0].value == 'id')
    col = {c.value: c.column for c in hdr}
    row = next(r for r in ws.iter_rows() if r[1].value == 'Edge')
    cell = row[col['A_cm2'] - 1]
    assert cell.value == members[8]['A'], 'the full value is kept'
    assert cell.number_format == '0.0###'
    assert _import(path)[2] == [], 'and an untouched row is still a no-op'
